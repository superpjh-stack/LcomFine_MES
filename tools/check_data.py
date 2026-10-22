#!/usr/bin/env python
"""QA2 흐름 · 데이터 검사기 — G-05 ~ G-15 판정 (`make check-data` · `tools/gate.py` 가 부른다). 소유 QA2.

W1 = G-05 ~ G-11 · W2 = G-11 `erp_upload_error` + G-12 ~ G-14(수금) · W3 = G-15(대시보드 8 독립 재계산) + G-05 RPT 8 화면 diff 0.
W4 = G-05 DOC(doc.py · sso.py 기능 단위 정적 · DOC GET 2 diff 0) + G-07 완료 기준 사슬(출하 1건 → 원자재 LOT · 독립 사슬 SQL = LOT-020 = trace). `LATER` 없음 — 전부 잰다.

하는 일 — **코드는 읽기만 한다. 만든 행은 끝에 지운다(접두 `QA2-`).**
  G-05  쓰기 경계 — 라우터 파일별 쓰기 SQL 정적 스캔(+ 그 라우터가 부르는 공용 모듈 · COL 은 `erp_upload.py` 포함) ↔ `function-list.md`
        「쓰는 테이블」 + `db-schema.md` §2.2 예외 · TRK/LOT/PRP · trace · stats 쓰기 0 · 조회 화면(COL 7 포함) GET 전후 업무 테이블 행 수 diff 0
  G-06  E2E 1건을 **HTTP 만으로** — 부서 계정을 바꿔 가며 기준정보(BAS) → 의뢰 → 지시 → 입고 → 입고검사 → (분할) → 불출 → 일보
        → 검사일보 → 제품 입고 → 판정 → 출하 지시 → 확정. 단계마다 계약이 말하는 DB 상태를 확인한다.
  G-07  역추적 — LOT-020 화면 · `trace.backward` = QA2 독립 SQL(원 입고 LOT 집합) · 출하 후 불합격 표시(D-515 · D-702)
  G-08  정방향 — LOT-010 화면 · `trace.forward` = 독립 SQL(제품 LOT · 출하 집합)
  G-09  차단 — 검사대기 · 보류 LOT 불출 422 · 보류 LOT 출하 확정 422 · 검사대기 LOT 할당 422 · 중지 품목 입고 · 출하 지시 422
  G-10  가용 — 할당 뒤 가용 = qty − allocated · 확정 뒤 가용 불변 · 동시 할당(스레드 2) 초과 없음 · 전 LOT 할당 = 열린 출하 행 합 · INV-010 화면 값
        · F-SHP-012 라인 전부 제거 = 지시 취소(Ruling 21) → 할당 0 · 가용 복원 · 독립 SQL 대조
  G-11  알림 — W1 7종 + due_d3_unshipped(배치) + erp_upload_error(W2 · AC 1행 · dedupe · 교체 · 재해석으로 닫힘) 수신 부서마다 1행
        · dedupe · 처리 화면에서 닫힘 · HOME-020 처리 · 남의 부서 404
  G-12  회차 교체 — 같은 기준 기간 재업로드(HTTP · openpyxl 로 만든 (예시) xlsx) → 새 회차 `반영` · 이전 회차 `교체` + replaced_by · 행 보존 ·
        기간마다 `반영` 하나 · `v_dq_collections` 는 반영 회차만 · COL-020 구분 · col.py 의 POST 는 업로드 · 매핑뿐(정적)
  G-13  오류 행 · 미매핑 — 날짜 · 금액 · 필수 · 미매핑 행이 error 와 함께 보존 · 반영(뷰 · COL-040 · COL-050)에서 제외 · COL-030 미매핑 목록 ·
        매핑 뒤 반영 회차만 재해석(error is null) · 교체 회차 그대로 · COL-020 회차별 오류 행 · erp_upload_row 쓰기 = parsed · error 만(정적)
  G-14  출하 대비 수금 — E2E 의 확정 출하 + 같은 월 · 거래처의 ERP 매출 · 수금 → COL-070 화면 값 = QA2 독립 SQL(shipment_order 확정 confirmed_at 월
        × erp_upload_row 반영 · error 없음 · 사용 중 매핑 · txn_date 월) · 출하만 = MES 쪽만 · ERP 만 = ERP 쪽만 · 미매핑 코드 제외
  G-15  대시보드 8 수치 — RPT-010~080 마다 `dashboard.<fn>(period, customer_id)` 반환(KPI · series · tables) = **QA2 가 따로 쓴 SQL**(D-805 정의 ·
        group by · filter · 뷰 · stats 를 거치지 않음) = 화면 HTML(`data-kpi` 글자 · `data-series` SVG <title> 「라벨: 값」 · 표 둘째 열 · `data-table` 행).
        세계 = 이 흐름의 달(S1 확정 월) · 거래처 C1(흐름) · C2(ERP 만). 020 · 030 · 040 은 전 DB 집계라 함수 · 독립 SQL · 화면 GET 을 **한 REPEATABLE READ
        스냅샷**(`Flow.snapshot` — `conn.q/q1` 만 그 트랜잭션으로) 에서 읽는다(다른 QA 가 같은 DB 에 동시에 쓴다). 0건 → `미수집` · 로스율 `미확정 (D-519)` · 병목 `미확정 (D-520)` · RPT-060 asof 명시 ·
        RPT-080 일일/주간/월간 구성(D-512) · 거래처 필터 010 · 050 · 060 · 070 · 080 만(D-805) · RPT 8 화면 GET 전후 업무 테이블 diff 0(G-05)

    uv run python tools/check_data.py                # 판정표 (종료코드 = FAIL 수)
    uv run python tools/check_data.py --transcript   # + E2E 기록(마크다운 표) — outputs/qa2-흐름데이터.md 에 붙인다
    uv run python tools/check_data.py --keep         # 정리하지 않는다(디버그)
"""

from __future__ import annotations

import warnings

warnings.filterwarnings("ignore", message="Using `httpx` with `starlette.testclient`")

import html as _html
import re
import subprocess
import sys
import threading
import time
from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

PASS, FAIL, UNVERIFIED = "PASS", "FAIL", "미검증"

GATE_NAMES = {
    "G-05": "쓰기 경계", "G-06": "E2E 흐름 1건", "G-07": "역추적", "G-08": "정방향 추적", "G-09": "차단 규칙",
    "G-10": "가용 재계산", "G-11": "알림 9종", "G-12": "회차 교체", "G-13": "오류 행 · 미매핑",
    "G-14": "출하 대비 수금", "G-15": "대시보드 8 수치",
}
LATER: dict[str, str] = {}      # W3 까지 전부 잰다 — 잴 수 없는 것은 이유와 함께 FAIL
#: 흐름이 앞에서 끊기면 행이 없어 PASS 로 보일 수 있는 것 — 행이 없으면 FAIL 을 넣는다 (게이트를 낮추지 않는다)
MUST_MEASURE = {("G-11", "erp_upload_error"), ("G-12", "같은 기간 재업로드"), ("G-13", "오류 행 보존"), ("G-14", "COL-070 화면")}
#: W3 대시보드 단계가 돌지 못하면 행이 없어 PASS 로 보일 수 있는 것
MUST_MEASURE_W3 = {("G-15", "RPT-010"), ("G-15", "RPT-080"), ("G-05", "RPT 대시보드 8 화면")}
#: W4 — 완료 기준 사슬(CLAUDE.md) 은 추적 단계에서 반드시 잰다 · DOC 정적 행은 static_doc 이 늘 남긴다(파일 없으면 FAIL)
MUST_MEASURE_W4 = {("G-07", "완료 기준 사슬")}
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ════════════════════════════════════════════════════════════════════════════
# 기록
# ════════════════════════════════════════════════════════════════════════════
@dataclass
class Check:
    gate: str
    item: str
    status: str
    measured: str


@dataclass
class Step:
    step: str
    account: str
    request: str
    status: int | str
    numbers: str = ""
    db: str = ""


@dataclass
class Report:
    checks: list[Check] = field(default_factory=list)
    steps: list[Step] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    def add(self, gate: str, item: str, ok: bool | str, measured: str) -> bool:
        st = ok if isinstance(ok, str) else (PASS if ok else FAIL)
        self.checks.append(Check(gate, item, st, " ".join(str(measured).split())))
        return st == PASS

    def get(self, gate: str, item_prefix: str) -> Check | None:
        return next((c for c in self.checks if c.gate == gate and c.item.startswith(item_prefix)), None)


class StepFailed(RuntimeError):
    """흐름의 앞 단계가 실패해 뒤 단계를 이어 갈 수 없다."""


# ════════════════════════════════════════════════════════════════════════════
# G-05 쓰기 경계 — 정적 스캔
# ════════════════════════════════════════════════════════════════════════════
ROUTERS = ROOT / "src/lcomfine/app/routers"
APP = ROOT / "src/lcomfine/app"
WRITE_RE = re.compile(r"\b(?:insert\s+into|delete\s+from)\s+([A-Za-z_{}]+)|\bupdate\s+([A-Za-z_{}]+)\s+set\b", re.I)
#: 공용 모듈 → 그 모듈이 쓰는 테이블(정적 스캔으로 다시 확인한다)
HELPERS = {"notify": "notify.py", "numbering": "numbering.py", "audit": "util/audit.py", "auth": "auth.py"}
#: 모든 라우터에 허용 — db-schema §2.2 「모든 쓰기」 · 「모든 화면 조회」 줄
ALWAYS = {"notification", "sys_change_log", "sys_access_log"}
#: 읽기만 하는 모듈 (db-schema §2.1 규칙 4 · CLAUDE.md)
READ_ONLY = {"trk.py": "TRK", "lot.py": "LOT", "prp.py": "PRP"}
READ_ONLY_HELPERS = ("trace.py", "stats.py", "dashboard.py")


def _sql_writes(text: str) -> list[str]:
    out = []
    for m in WRITE_RE.finditer(text):
        t = (m.group(1) or m.group(2) or "").lower()
        if t and t not in ("set",):
            out.append(t)
    return out


def contract_writes() -> dict[str, set[str]]:
    """function-list.md §2 → {대메뉴: 쓰는 테이블 합}."""
    out: dict[str, set[str]] = {}
    for ln in (ROOT / "contracts/function-list.md").read_text(encoding="utf-8").splitlines():
        if not ln.startswith("| F-"):
            continue
        cells = [c.strip() for c in ln.strip("|").split("|")]
        menu, tables = cells[1], cells[5]
        if tables and tables != "-":
            out.setdefault(menu, set()).update(t.strip() for t in tables.split(","))
        else:
            out.setdefault(menu, set())
    return out


def schema_exceptions() -> set[str]:
    """db-schema.md §2.2 예외 표에 나온 테이블 이름 전부."""
    text = (ROOT / "contracts/db-schema.md").read_text(encoding="utf-8")
    sec = text.split("### 2.2", 1)[1].split("### 2.3", 1)[0]
    return set(re.findall(r"`([a-z_]+)`", sec)) - {"status", "allocated_qty", "qty", "parent_lot_id"}


def helper_writes() -> dict[str, set[str]]:
    return {name: set(_sql_writes((APP / rel).read_text(encoding="utf-8"))) for name, rel in HELPERS.items()}


def _resolve_dynamic(fname: str, text: str, tables: list[str]) -> tuple[set[str], list[str]]:
    """f-string 테이블(`{table}` · `{insp}`)을 그 파일 안의 근거로 푼다. 못 풀면 목록으로 돌려준다."""
    real, unresolved = set(), []
    for t in tables:
        if "{" not in t:
            real.add(t)
            continue
        if t == "{table}" and "def remove_or_disable" in text:
            # bas.remove_or_disable 은 호출한 라우터의 테이블에 쓴다 — 호출 쪽에서 센다(아래 callers)
            continue
        if t == "{insp}" and fname == "qc.py":
            real.update(re.findall(r'"((?:incoming|product)_inspection)"', text))
            continue
        unresolved.append(f"{fname}:{t}")
    return real, unresolved


def static_scan(report: Report) -> None:
    contract = contract_writes()
    exc = schema_exceptions()
    helpers = helper_writes()
    menu_of = {p.stem.upper(): p for p in ROUTERS.glob("*.py") if p.stem != "__init__"}
    violations, unresolved, detail = [], [], []
    infra = {}   # 공용 모듈의 쓰기 중 §2.2 에 없는 것 → 그 모듈을 부르는 라우터
    # COL 영역 모듈 erp_upload.py — 라우터가 부르면 그 쓰기(+ 안에서 부르는 notify · numbering)는 부른 라우터의 쓰기다
    erp_text = (APP / "erp_upload.py").read_text(encoding="utf-8")
    erp_tables = set(_sql_writes(erp_text))
    erp_callers = []
    for menu, path in sorted(menu_of.items()):
        text = path.read_text(encoding="utf-8")
        tables, unr = _resolve_dynamic(path.name, text, _sql_writes(text))
        unresolved += unr
        tables |= set(re.findall(r'remove_or_disable\(cur,\s*"([a-z_]+)"', text))
        for h in ("notify", "numbering"):
            if re.search(rf"\b{h}\.(emit|handle|handle_by_link|next)\(", text):
                tables |= helpers[h]
        if re.search(r"\berp_upload\.(apply|remap|save_column_map)\(", text):
            erp_callers.append(menu)
            tables |= erp_tables | helpers["notify"] | helpers["numbering"]
        if re.search(r"\baudit\.log_change\(", text) or "templating.render(" in text:
            tables |= helpers["audit"]
        # 공용 모듈(numbering · notify · audit)이 쓰는 테이블은 db-schema §2.2 「이 표가 전부다」 에 있으면 허용 — 라우터가 부른 결과다
        helper_tables = set().union(*helpers.values())
        allowed = contract.get(menu, set()) | ALWAYS | (helper_tables & exc)
        bad = sorted(tables - allowed)
        for t in bad:
            if t in helper_tables:
                infra.setdefault(t, []).append(menu)
            else:
                violations.append(f"{menu}→{t}")
        detail.append(f"{menu} {len(tables)}")
    # 공용 모듈 자체 — 라우터가 부르지 않는 auth(로그인 · 로그아웃)도 §2.2 에 있어야 한다
    for name, ts in helpers.items():
        for t in sorted(ts - exc - ALWAYS):
            if name == "auth":
                infra.setdefault(t, []).append("auth(로그인)")
    report.add("G-05", "쓰기 경계 — 라우터 정적 스캔", not violations and not unresolved,
               f"라우터 {len(menu_of)} 의 쓰기 테이블 ⊆ 계약(function-list 쓰는 테이블 ∪ 알림 · 로그) · "
               f"위반 {len(violations)} {violations or ''} · 해석 못한 동적 쓰기 {len(unresolved)} {unresolved or ''}")
    # COL — col.py 자신의 쓰기 + erp_upload.py 의 쓰기 ⊆ F-COL 계약 테이블. erp_upload 를 부르는 라우터는 col.py 만 · 다른 라우터가 erp_* 에 쓰지 않는다
    col_contract = contract.get("COL", set())
    col_own = set(_sql_writes((ROUTERS / "col.py").read_text(encoding="utf-8"))) if (ROUTERS / "col.py").exists() else set()
    erp_bad = sorted(erp_tables - col_contract)
    col_bad = sorted(col_own - col_contract - ALWAYS)
    others = sorted(f"{m}→{t}" for m, p in menu_of.items() if m != "COL"
                    for t in set(_sql_writes(p.read_text(encoding="utf-8"))) if t.startswith("erp_"))
    erp_helpers = sorted(h for h in ("notify", "numbering") if re.search(rf"\b{h}\.(emit|handle_by_link|next)\(", erp_text))
    report.add("G-05", "COL 라우터 · erp_upload.py 쓰기 ⊆ F-COL 계약 테이블 ∪ 알림 · 채번",
               not erp_bad and not col_bad and not others and erp_callers == ["COL"],
               f"erp_upload.py 쓰기 {sorted(erp_tables)} · col.py 직접 쓰기 {sorted(col_own)} ⊆ F-COL 「쓰는 테이블」 {sorted(col_contract)} · "
               f"계약 밖 {erp_bad + col_bad or 0} · erp_upload 쓰기 함수를 부르는 라우터 {erp_callers} (기대 ['COL']) · "
               f"다른 라우터의 erp_* 쓰기 {others or 0} · erp_upload 안의 공용 모듈 {erp_helpers} (§2.2 notification · sys_number_seq)")
    helper_in_exc = sorted(t for t in set().union(*helpers.values()) if t in exc)
    report.add("G-05", "쓰기 경계 — 공용 모듈의 계약 밖 쓰기", not infra,
               "db-schema §2.2 「이 표가 전부다」에 없는 쓰기: "
               + (" · ".join(f"{t} ← {', '.join(sorted(set(v)))}" for t, v in sorted(infra.items())) or "0")
               + f" · §2.2 가 허용한 공용 모듈 쓰기 {helper_in_exc} (numbering → sys_number_seq · auth → sys_user)")
    # 예외 줄의 컬럼 — QC 는 lot_stock.status 만, SHP 지시는 allocated_qty 만, SHP 확정은 qty · allocated_qty (db-schema §2.2)
    col_rules = {("qc.py", None): {"status"}, ("shp.py", "create_order"): {"allocated_qty"},
                 ("shp.py", "update_order"): {"allocated_qty"}, ("shp.py", "confirm_order"): {"qty", "allocated_qty"}}
    col_bad, col_seen = [], 0
    for fname in ("qc.py", "shp.py"):
        text = (ROUTERS / fname).read_text(encoding="utf-8")
        for m in re.finditer(r'update\s+lot_stock\s+set\s+([^"]*?)(?:\s|")+where', text, re.I):
            func = re.findall(r"^def (\w+)\(", text[: m.start()], re.M)
            func = func[-1] if func else "?"
            cols = {c.split("=")[0].strip() for c in m.group(1).split(",")} - {"updated_at", "updated_by"}
            allowed = col_rules.get((fname, None)) or col_rules.get((fname, func))
            col_seen += 1
            if allowed is None or not cols <= allowed:
                col_bad.append(f"{fname}:{func} {sorted(cols)}")
    report.add("G-05", "예외 줄의 lot_stock 컬럼 (QC status · SHP allocated_qty/qty)", not col_bad,
               f"update lot_stock {col_seen} 문 · 허용 밖 컬럼 {len(col_bad)} {col_bad or ''}")
    ro_hits = []
    for fname, menu in READ_ONLY.items():
        p = ROUTERS / fname
        if p.exists():
            w = _sql_writes(p.read_text(encoding="utf-8"))
            if w or re.search(r"\b(notify\.emit|numbering\.next|conn\.x|conn\.tx)\(", p.read_text(encoding="utf-8")):
                ro_hits.append(f"{menu}:{w}")
    for fname in READ_ONLY_HELPERS:
        t = (APP / fname).read_text(encoding="utf-8")
        if _sql_writes(t) or re.search(r"\b(conn\.x|conn\.tx)\(", t):
            ro_hits.append(f"{fname}:{_sql_writes(t)}")
    report.add("G-05", "TRK · LOT · PRP · trace · stats 쓰기 0 (정적) + dashboard.py", not ro_hits,
               f"쓰기 문 {len(ro_hits)} {ro_hits or ''} — trk.py · lot.py · prp.py · trace.py · stats.py · dashboard.py")
    # RPT — rpt.py 는 직접 쓰기 0 · 쓰는 것은 dataquery(data_query_log) · ai(sys_ai_call_log) · numbering 만 (F-RPT-132 「쓰는 테이블」 · D-513)
    rpt_text = (ROUTERS / "rpt.py").read_text(encoding="utf-8") if (ROUTERS / "rpt.py").exists() else ""
    rpt_own = set(_sql_writes(rpt_text))
    dq_w = set(_sql_writes((APP / "dataquery.py").read_text(encoding="utf-8"))) if (APP / "dataquery.py").exists() else set()
    ai_w = set(_sql_writes((APP / "ai.py").read_text(encoding="utf-8"))) if (APP / "ai.py").exists() else set()
    rpt_contract = contract.get("RPT", set())
    rpt_bad = sorted((rpt_own | dq_w | ai_w) - rpt_contract - ALWAYS)
    # 질의(쓰기)를 부르는 곳 — dataquery.ask · ai.complete · conn.tx 가 POST 핸들러(run_query) 안에만 있어야 GET 8 이 로그를 쓰지 않는다
    writers = []
    for m in re.finditer(r"\b(dataquery\.ask|ai\.complete|conn\.tx)\(", rpt_text):
        func = re.findall(r"^def (\w+)\(", rpt_text[: m.start()], re.M)
        writers.append((m.group(1), func[-1] if func else "?"))
    get_handlers = re.findall(r"@router\.get\([^)]*\)[^\n]*\ndef (\w+)\(", rpt_text)
    post_handlers = re.findall(r"@router\.post\([^)]*\)[^\n]*\ndef (\w+)\(", rpt_text)
    writers_ok = bool(writers) and all(fn in post_handlers for _, fn in writers)
    report.add("G-05", "RPT 라우터 · dashboard.py 직접 쓰기 0 · 질의 로그는 dataquery · ai 뿐 · dataquery.ask 는 POST run_query 안에서만 (정적)",
               bool(rpt_text) and not rpt_own and not rpt_bad and dq_w <= {"data_query_log"} and ai_w <= {"sys_ai_call_log"} and writers_ok,
               f"rpt.py 직접 쓰기 {sorted(rpt_own) or 0} · dataquery.py 쓰기 {sorted(dq_w)} · ai.py 쓰기 {sorted(ai_w)} ⊆ F-RPT 「쓰는 테이블」 {sorted(rpt_contract)} · "
               f"계약 밖 {rpt_bad or 0} · 쓰기 호출 {sorted(set(writers))} 전부 POST 핸들러 {post_handlers} 안 {writers_ok} · GET 핸들러 {len(get_handlers)} 에는 없음 · "
               f"대시보드 집계 캐시 테이블 없음(dashboard.py 쓰기 0 — 위 행)")
    static_doc(report)


# ── W4 — DOC 문서관리 · sso.py (F-DOC-011 「쓰는 테이블」 sso_token, sys_access_log · F-DOC-021~023 doc_account_map · §2.2 F-DOC-011 → sys_access_log) ──
DOC_TABLES = {"doc_account_map", "sso_token"}
#: 다른 모듈 호출 → 그 호출이 쓰는 테이블 (sso 는 함수 단위로 sso.py 를 읽어 채운다)
_EXT_WRITES = {"audit.write_log": {"sys_access_log"}, "audit.log_change": {"sys_change_log"}}


def contract_fn_writes() -> dict[str, set[str]]:
    """function-list.md §2 → {기능 ID: 쓰는 테이블}."""
    out: dict[str, set[str]] = {}
    for ln in (ROOT / "contracts/function-list.md").read_text(encoding="utf-8").splitlines():
        if ln.startswith("| F-"):
            cells = [c.strip() for c in ln.strip("|").split("|")]
            out[cells[0]] = {t.strip() for t in cells[5].split(",")} if cells[5] and cells[5] != "-" else set()
    return out


def _fn_info(text: str) -> dict[str, dict]:
    """파일의 최상위 함수마다 {sql: 쓰기 테이블, calls: 호출 이름, route: get|post|None, fn_id, dyn: 해석 못한 f-string 쓰기}.

    f-string 은 `{…}` 자리표시로 다시 짜서 스캔한다(동적 테이블 이름이면 dyn 에 남는다)."""
    import ast

    tree = ast.parse(text)
    out = {}
    for node in tree.body:
        if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            continue
        strs, calls, fn_id = [], set(), None
        inner_consts = {id(v) for n in ast.walk(node) if isinstance(n, ast.JoinedStr) for v in n.values}
        for n in ast.walk(node):
            if isinstance(n, ast.JoinedStr):
                strs.append("".join(v.value if isinstance(v, ast.Constant) else "{x}" for v in n.values))
            elif isinstance(n, ast.Constant) and isinstance(n.value, str) and id(n) not in inner_consts:
                strs.append(n.value)
            elif isinstance(n, ast.Call):
                f = n.func
                name = f.id if isinstance(f, ast.Name) else (
                    f"{f.value.id}.{f.attr}" if isinstance(f, ast.Attribute) and isinstance(f.value, ast.Name) else None)
                if name:
                    calls.add(name)
                if name == "rbac.require_fn" and n.args and isinstance(n.args[0], ast.Constant):
                    fn_id = n.args[0].value
        route = None
        for d in node.decorator_list:
            if isinstance(d, ast.Call) and isinstance(d.func, ast.Attribute) and isinstance(d.func.value, ast.Name) \
                    and d.func.value.id == "router":
                route = d.func.attr
        tables = [t for s in strs for t in _sql_writes(s)]
        out[node.name] = {"sql": {t for t in tables if "{" not in t}, "dyn": [t for t in tables if "{" in t],
                          "calls": calls, "route": route, "fn_id": fn_id}
    return out


def _closure(info: dict[str, dict], name: str, ext: dict[str, set[str]], seen=None) -> tuple[set[str], set[str]]:
    """함수 하나가 (같은 파일 안 호출을 따라가며) 쓰는 테이블 · 부르는 외부 쓰기 함수."""
    seen = seen if seen is not None else set()
    if name in seen or name not in info:
        return set(), set()
    seen.add(name)
    tables, ext_calls = set(info[name]["sql"]), set()
    for c in info[name]["calls"]:
        if c in ext:
            tables |= ext[c]
            ext_calls.add(c)
        elif c in info:
            t, e = _closure(info, c, ext, seen)
            tables |= t
            ext_calls |= e
    return tables, ext_calls


def static_doc(report: Report) -> None:
    """G-05 DOC — doc.py · sso.py 의 쓰기를 기능(핸들러) 단위로 F-DOC 「쓰는 테이블」 ∪ 알림 · 로그 · §2.2 와 대조."""
    doc_p, sso_p = ROUTERS / "doc.py", APP / "sso.py"
    if not doc_p.exists() or not sso_p.exists():
        report.add("G-05", "DOC 라우터 · sso.py 쓰기 ⊆ F-DOC 계약 (정적)", FAIL, f"파일 없음 doc.py {doc_p.exists()} · sso.py {sso_p.exists()}")
        return
    fn_contract = contract_fn_writes()
    exc = schema_exceptions()
    sso_info = _fn_info(sso_p.read_text(encoding="utf-8"))
    sso_ext = {}
    for name, i in sso_info.items():
        if not name.startswith("_"):
            t, _ = _closure(sso_info, name, {})
            if t:
                sso_ext[f"sso.{name}"] = t          # 공개 함수 중 쓰기가 있는 것 — issue · consume
    sso_all = set().union(*sso_ext.values()) if sso_ext else set()
    ext = {**_EXT_WRITES, **sso_ext}
    doc_info = _fn_info(doc_p.read_text(encoding="utf-8"))
    dyn = [f"doc.py:{n}:{t}" for n, i in doc_info.items() for t in i["dyn"]] + \
          [f"sso.py:{n}:{t}" for n, i in sso_info.items() for t in i["dyn"]]
    per_fn, bad, get_bad = [], [], []
    for name, i in doc_info.items():
        if not i["route"]:
            continue
        tables, ext_calls = _closure(doc_info, name, ext)
        sso_calls = sorted(c for c in ext_calls if c.startswith("sso."))
        if i["route"] == "get":
            if tables or sso_calls:
                get_bad.append(f"GET {name} {sorted(tables)} {sso_calls}")
            continue
        fid = i["fn_id"]
        allowed = fn_contract.get(fid, set()) | ALWAYS
        over = sorted(tables - allowed)
        if not fid or fid not in fn_contract or over:
            bad.append(f"{name}({fid}) 계약 밖 {over}")
        per_fn.append(f"{fid} {name} {sorted(tables)}{' via ' + ','.join(sso_calls) if sso_calls else ''}")
    # F-DOC-011 진입: sso_token(sso.issue) + sys_access_log(SSO 로그) 둘 다 실제로 쓴다 · §2.2 에 F-DOC-011 → sys_access_log 줄
    go_fn = next((n for n, i in doc_info.items() if i["fn_id"] == "F-DOC-011" and i["route"] == "post"), None)
    go_t = _closure(doc_info, go_fn, ext)[0] if go_fn else set()
    go_ok = bool(go_fn) and {"sso_token", "sys_access_log"} <= go_t and fn_contract.get("F-DOC-011") == {"sso_token", "sys_access_log"}
    sec22 = (ROOT / "contracts/db-schema.md").read_text(encoding="utf-8").split("### 2.2", 1)[1].split("### 2.3", 1)[0]
    exc_line = bool(re.search(r"\|\s*F-DOC-011\s*\|\s*`sys_access_log`", sec22)) and "sys_access_log" in exc
    # 다른 곳 — app 전체(doc.py · sso.py 밖)에서 doc_account_map · sso_token 쓰기 0 · sso 쓰기 함수 호출 0
    others, callers = [], []
    sso_call_re = re.compile(r"\bsso\.(" + "|".join(re.escape(k.split(".", 1)[1]) for k in sso_ext) + r")\(") if sso_ext else None
    for p in sorted(APP.rglob("*.py")):
        if p in (doc_p, sso_p):
            continue
        t = p.read_text(encoding="utf-8")
        hit = sorted(set(_sql_writes(t)) & DOC_TABLES)
        if hit:
            others.append(f"{p.relative_to(APP)}→{hit}")
        if sso_call_re and sso_call_re.search(t):
            callers.append(str(p.relative_to(APP)))
    seed_w = sorted(f"{p.name}→{sorted(set(_sql_writes(p.read_text(encoding='utf-8'))) & DOC_TABLES)}"
                    for p in (ROOT / "src/lcomfine/db").glob("*.py") if set(_sql_writes(p.read_text(encoding="utf-8"))) & DOC_TABLES)
    ok = (not bad and not get_bad and not dyn and go_ok and exc_line and not others and not callers
          and sso_all <= {"sso_token"} and per_fn)
    report.add("G-05", "DOC 라우터 · sso.py 쓰기 ⊆ F-DOC 계약 ∪ 알림 · 로그 · §2.2 — 기능 단위 (정적)", ok,
               f"sso.py 쓰기 함수 { {k: sorted(v) for k, v in sorted(sso_ext.items())} } ⊆ {{sso_token}} · "
               f"POST {len(per_fn)} [{' / '.join(per_fn)}] · 계약 밖 {bad or 0} · "
               f"F-DOC-011 = sso_token + sys_access_log {go_ok} · §2.2 「F-DOC-011 → sys_access_log」 줄 {exc_line} · "
               f"GET 핸들러 쓰기 {get_bad or 0} · 해석 못한 동적 쓰기 {len(dyn)} {dyn or ''} · "
               f"다른 app 모듈의 doc_account_map · sso_token 쓰기 {others or 0} · sso 쓰기 함수를 부르는 다른 모듈 {callers or 0} · "
               f"(참고 시드 {seed_w or 0} — 화면 아님)")


# ════════════════════════════════════════════════════════════════════════════
# G-12 · G-13 정적 — MES 에 수금 금액을 고치는 경로가 없다 (db-schema §2.1 규칙 5 · 정본 「MES 는 수금 데이터를 고치지 않는다」)
# ════════════════════════════════════════════════════════════════════════════
#: erp_upload_row 에서 바뀌어도 되는 컬럼 — 재해석(parsed · error) 과 수정 표시만. raw(원문) · batch_id · row_no 는 불변
ROW_MUTABLE = {"parsed", "error", "updated_at", "updated_by"}
#: erp_upload_batch 에서 바뀌어도 되는 컬럼 — 교체 표시 · 오류 수 갱신만. 기간 · 행 수 · 번호는 불변
BATCH_MUTABLE = {"status", "replaced_by", "error_count", "updated_at", "updated_by"}
#: F-COL-011 · 031 · 032 · 033 — function-list 의 COL 쓰기 기능 API (수금 금액 수정 엔드포인트는 없다)
COL_WRITE_PATHS = {"/col/uploads", "/col/customer-map", "/col/customer-map/{id}", "/col/customer-map/{id}/delete"}


def _set_clauses(text: str, table: str) -> list[tuple[str, set[str]]]:
    """`update <table> [alias] set … where` 의 set 절을 괄호 깊이 0 의 `where` 까지 잘라 대입 컬럼 집합으로. [(set 절, 컬럼들)]."""
    out = []
    for m in re.finditer(rf"\bupdate\s+{table}\s+(?:\w+\s+)?set\s+", text, re.I):
        i, depth, buf = m.end(), 0, []
        while i < len(text):
            ch = text[i]
            if ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
            elif depth == 0 and re.match(r"\s+where\b|\s+returning\b|\"\"\"", text[i:], re.I):
                break
            buf.append(ch)
            i += 1
        clause = "".join(buf)
        cols, depth, cur = set(), 0, []
        for ch in clause + ",":
            if ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
            if ch == "," and depth == 0:
                piece = "".join(cur).strip()
                if "=" in piece:
                    cols.add(piece.split("=", 1)[0].strip().split(".")[-1])
                cur = []
            else:
                cur.append(ch)
        out.append((" ".join(clause.split())[:80], cols))
    return out


def static_col(report: Report) -> None:
    """G-12 · G-13 정적 — col.py 의 POST 는 업로드 · 매핑뿐 · erp_upload_row 쓰기는 parsed · error 재해석만 · raw 불변 · 삭제 0."""
    from lcomfine.app.main import app

    def walk(routes):
        """FastAPI 가 include 한 라우터(`_IncludedRouter.original_router`)까지 펼친다."""
        for r in routes:
            inner = getattr(r, "original_router", None) or (r if hasattr(r, "routes") and not hasattr(r, "endpoint") else None)
            if inner is not None and hasattr(inner, "routes"):
                yield from walk(inner.routes)
            else:
                yield r

    posts = sorted(r.path for r in walk(app.routes) if getattr(r, "path", "").startswith("/col/") and "POST" in (getattr(r, "methods", None) or ()))
    col_text = (ROUTERS / "col.py").read_text(encoding="utf-8") if (ROUTERS / "col.py").exists() else ""
    # col.py 가 수금 행 · 회차 테이블을 직접 쓰지 않는다 (erp_customer_map 만) — 금액을 받는 폼 칸이 없다
    col_direct = sorted(t for t in set(_sql_writes(col_text)) if t in ("erp_upload_row", "erp_upload_batch"))
    form_names = sorted(set(re.findall(r"(\w+):\s*str\s*\|\s*None\s*=\s*Form\(", col_text)))
    report.add("G-12", "col.py 의 POST = 업로드 · 매핑만 (수금 금액 수정 경로 없음 · 정적)",
               set(posts) == COL_WRITE_PATHS and not col_direct,
               f"POST {posts} = function-list F-COL-011 · 031 · 032 · 033 {sorted(COL_WRITE_PATHS)} · col.py 가 직접 쓰는 수금 테이블 {col_direct or 0} "
               f"· 폼 칸 {form_names} (금액 칸 없음) · 수금 데이터는 erp_upload.apply(회차 교체) · remap(재해석)으로만 바뀐다")
    # 앱 전체에서 erp_upload_row · erp_upload_batch 의 update 컬럼 · delete
    bad, seen, deletes = [], [], []
    for p in sorted(APP.rglob("*.py")):
        text = p.read_text(encoding="utf-8")
        for table, allowed in (("erp_upload_row", ROW_MUTABLE), ("erp_upload_batch", BATCH_MUTABLE)):
            for clause, cols in _set_clauses(text, table):
                seen.append(f"{p.name}:{table} {sorted(cols)}")
                if not cols or not cols <= allowed:
                    bad.append(f"{p.name}: update {table} set {sorted(cols)} (허용 {sorted(allowed)})")
        deletes += [f"{p.name}:{t}" for t in _sql_writes(text) if t in ("erp_upload_row", "erp_upload_batch")
                    and re.search(rf"delete\s+from\s+{t}\b", text, re.I)]
    report.add("G-13", "erp_upload_row 쓰기 = parsed · error 재해석만 · raw 불변 · 회차 삭제 0 (정적)", not bad and not deletes,
               f"update 문 {len(seen)} {seen} · 허용 밖 {bad or 0} · delete {deletes or 0} — raw(원문) · 기간 · 행 수는 어디서도 바꾸지 않는다")


# ════════════════════════════════════════════════════════════════════════════
# HTTP 흐름
# ════════════════════════════════════════════════════════════════════════════
ACCOUNTS = ("sl", "pr", "sh", "ds", "ac", "st", "ex", "admin")
READ_SCREENS = ["/trk/orders", "/trk/due", "/trk/customers", "/lot/forward", "/lot/backward", "/prp/utilization",
                "/prp/machines", "/prp/completion", "/wo/status", "/inv/stock", "/qc/history", "/shp/history", "/",
                "/col/uploads", "/col/upload-history", "/col/customer-map", "/col/collections", "/col/receivables",
                "/col/aging", "/col/shipment-vs-collection"]
#: 행 수 diff 에서 빼는 로그 테이블 — 조회 로그(모든 GET 이 쓴다) · 변경 로그 · 질의 · AI 로그는 **다른 세션의 쓰기**(동시에 도는 QA 의 POST · 질의)로
#: 늘어난다. 대시보드 GET 이 질의 로그를 쓰지 않는 것은 정적 행(「RPT 라우터 … dataquery.ask 는 POST run_query 안에서만」)이 보증한다
LOG_TABLES = {"sys_access_log", "sys_change_log", "data_query_log", "sys_ai_call_log"}


def num(text: str) -> Decimal:
    return Decimal(text.replace(",", "").strip() or "0")


class Flow:
    def __init__(self, report: Report) -> None:
        from fastapi.testclient import TestClient

        from lcomfine.app.main import app
        from lcomfine.app.settings import get_settings
        from lcomfine.db import conn

        self.TestClient, self.app, self.conn = TestClient, app, conn
        self.pw = get_settings().seed_password
        if not self.pw:
            raise RuntimeError("LCOMFINE_SEED_PASSWORD 가 .env 에 없다 — 로그인할 수 없다")
        self.r = report
        self.tag = f"QA2-{int(time.time() * 1000) % 10**9}"
        self.today = date.today()
        self.clients: dict[str, object] = {}
        # 만든 것 (정리용)
        self.items: list[int] = []
        self.customer_id: int | None = None
        self.warehouse_id: int | None = None
        self.locations: list[int] = []
        self.req_ids: list[int] = []
        self.req_nos: list[str] = []
        self.wo_ids: list[int] = []
        self.wo_nos: list[str] = []
        self.receipt_ids: list[int] = []
        self.mlot_nos: list[str] = []
        self.plot_nos: list[str] = []
        self.ship_nos: list[str] = []
        self.issue_ids: list[int] = []
        # W2 수금 (정리용)
        self.customer2_id: int | None = None
        self.erp_batches: list[int] = []
        self.erp_codes: list[str] = []
        self.cmap_before: list[dict] | None = None     # 열 매핑을 바꿨으면 원래 행(되돌린다)
        self.v: dict[str, object] = {}

    # ── HTTP ──
    def c(self, login: str):
        if login not in self.clients:
            cl = self.TestClient(self.app)
            resp = cl.post("/login", data={"login_id": login, "password": self.pw}, follow_redirects=False)
            if resp.status_code != 303:
                raise StepFailed(f"{login} 로그인 {resp.status_code}")
            self.clients[login] = cl
        return self.clients[login]

    def post(self, step: str, login: str, path: str, data: dict, expect: int = 200, numbers=None) -> dict:
        resp = self.c(login).post(path, data=data, follow_redirects=False)
        body: dict = {}
        try:
            body = resp.json()
        except Exception:  # noqa: BLE001 — 본문이 JSON 이 아니면 그대로 적는다
            body = {"text": resp.text[:200]}
        msg = body.get("message") if isinstance(body.get("message"), str) else (body.get("detail") or {}).get("message", "") \
            if isinstance(body.get("detail"), dict) else ""
        nums = numbers(body) if (numbers and resp.status_code == expect) else (msg or "")
        self.r.steps.append(Step(step, login, f"POST {path}", resp.status_code, str(nums)))
        if resp.status_code != expect:
            raise StepFailed(f"{step}: POST {path} ({login}) → {resp.status_code} (기대 {expect}) {str(body)[:300]}")
        return body

    def try_post(self, step: str, login: str, path: str, data: dict) -> tuple[int, dict]:
        resp = self.c(login).post(path, data=data, follow_redirects=False)
        try:
            body = resp.json()
        except Exception:  # noqa: BLE001
            body = {"text": resp.text[:200]}
        det = body.get("detail") if isinstance(body.get("detail"), dict) else {}
        self.r.steps.append(Step(step, login, f"POST {path}", resp.status_code, det.get("message", "") or body.get("message", "")))
        return resp.status_code, body

    def get(self, step: str, login: str, path: str, expect: int | None = 200, note: str = ""):
        resp = self.c(login).get(path, follow_redirects=False)
        self.r.steps.append(Step(step, login, f"GET {path}", resp.status_code, note))
        if expect is not None and resp.status_code != expect:
            raise StepFailed(f"{step}: GET {path} ({login}) → {resp.status_code} (기대 {expect})")
        return resp

    def db(self, text: str) -> None:
        """방금 단계의 DB 확인 문장을 기록에 붙인다."""
        if self.r.steps:
            s = self.r.steps[-1]
            s.db = (s.db + " · " if s.db else "") + text

    # ── DB 읽기 (확인용 · 쓰지 않는다) ──
    def q(self, sql: str, params=()):
        return self.conn.q(sql, params)

    def q1(self, sql: str, params=()):
        return self.conn.q1(sql, params)

    def stock(self, kind: str, lot_no: str) -> list[dict]:
        table = "material_lot" if kind == "원자재" else "product_lot"
        return self.q(f"""select s.id, s.qty, s.allocated_qty, s.status, l.location_code
                            from lot_stock s join {table} t on t.id = s.lot_id join location l on l.id = s.location_id
                           where s.lot_kind = %s and t.lot_no = %s order by s.id""", (kind, lot_no))

    def stock1(self, kind: str, lot_no: str) -> dict:
        rows = self.stock(kind, lot_no)
        if len(rows) != 1:
            raise StepFailed(f"{lot_no} 재고 행 {len(rows)}개")
        return rows[0]

    def notes(self, dedupe_key: str) -> dict[str, dict]:
        return {r["to_role"]: r for r in self.q(
            "select id, to_role, handled_at, handled_by, link, event_key from notification where dedupe_key = %s", (dedupe_key,))}

    def expect(self, cond: bool, what: str) -> None:
        self.db(("OK " if cond else "불일치 ") + what)
        if not cond:
            raise StepFailed(what)

    # ════════════════════════════════════════════════════════════════════════
    # G-06 E2E
    # ════════════════════════════════════════════════════════════════════════
    def base(self) -> None:
        t = self.tag
        P = self.post("BAS-010 품목(제품)", "pr", "/bas/items",
                      {"item_code": f"{t}-P", "item_name": f"{t} 제품 (예시)", "item_type": "제품", "unit": "EA"},
                      numbers=lambda b: f"item {b['id']}")
        M = self.post("BAS-010 품목(원자재)", "pr", "/bas/items",
                      {"item_code": f"{t}-M", "item_name": f"{t} 원자재 (예시)", "item_type": "원자재", "unit": "KG"},
                      numbers=lambda b: f"item {b['id']}")
        X = self.post("BAS-010 품목(원자재 · 중지용)", "pr", "/bas/items",
                      {"item_code": f"{t}-X", "item_name": f"{t} 중지 원자재 (예시)", "item_type": "원자재"},
                      numbers=lambda b: f"item {b['id']}")
        self.items += [P["id"], M["id"], X["id"]]
        self.v.update(item_p=P["id"], item_m=M["id"], item_x=X["id"])
        self.post("BAS-010 품목 사용 중지", "pr", f"/bas/items/{X['id']}",
                  {"item_name": f"{t} 중지 원자재 (예시)", "item_type": "원자재", "use_yn": "N"})
        self.expect(self.q1("select use_yn from item where id = %s", (X["id"],))["use_yn"] == "N", "item.use_yn = N")
        C = self.post("BAS-020 거래처", "sl", "/bas/customers",
                      {"customer_code": f"{t}-C", "customer_name": f"{t} 거래처 (예시)"}, numbers=lambda b: f"customer {b['id']}")
        self.customer_id = C["id"]
        W = self.post("BAS-030 창고", "pr", "/bas/warehouses",
                      {"target": "warehouse", "warehouse_code": f"{t}-W", "warehouse_name": f"{t} 창고 (예시)"},
                      numbers=lambda b: f"warehouse {b['id']}")
        self.warehouse_id = W["id"]
        for code in ("L1", "L2"):
            L = self.post(f"BAS-030 위치 {code}", "pr", "/bas/warehouses",
                          {"target": "location", "warehouse_id": str(W["id"]), "location_code": code,
                           "location_name": f"{t} {code} (예시)"}, numbers=lambda b: f"location {b['id']}")
            self.locations.append(L["id"])
        self.v.update(L1=self.locations[0], L2=self.locations[1])

    def e2e(self) -> None:
        v, t = self.v, self.tag
        L1, L2 = str(v["L1"]), str(v["L2"])
        # 1 의뢰 (영업부)
        due1 = self.today + timedelta(days=10)
        b = self.post("WO-010 생산의뢰 등록", "sl", "/wo/requests",
                      {"customer_id": str(self.customer_id), "item_id": str(v["item_p"]), "qty": "100",
                       "due_date": due1.isoformat(), "note": f"{t} (예시)"}, numbers=lambda b: b["request_no"])
        R1, R1id = b["request_no"], b["id"]
        self.req_nos.append(R1)
        self.req_ids.append(R1id)
        v.update(R1=R1, R1id=R1id, due1=due1)
        row = self.q1("select status, qty, due_date from production_request where id = %s", (R1id,))
        self.expect(row["status"] == "확인대기" and row["qty"] == 100, f"production_request {R1} 확인대기 · 100")
        n = self.notes(f"request_created:{R1}")
        self.expect(set(n) == {"PR", "DS"} and all(x["handled_at"] is None for x in n.values()),
                    f"notification request_created:{R1} → PR · DS 1행씩 미처리")
        html = self.get("WO-010 의뢰 목록에서 번호 확인", "sl", f"/wo/requests?q={R1}").text
        self.expect(R1 in html, f"목록 화면에 {R1}")
        # 1b 단가·납기 확인 → 확정 (영업부)
        self.post("WO-010 의뢰 확정(수정)", "sl", f"/wo/requests/{R1id}",
                  {"customer_id": str(self.customer_id), "item_id": str(v["item_p"]), "qty": "100",
                   "due_date": due1.isoformat(), "status": "확정"})
        self.expect(self.q1("select status from production_request where id = %s", (R1id,))["status"] == "확정", "의뢰 확정")
        # 1c 납기 D-3 용 의뢰 (출하 없음)
        due2 = self.today + timedelta(days=2)
        b = self.post("WO-010 생산의뢰 등록(납기 D-2 · 출하 없음)", "sl", "/wo/requests",
                      {"customer_id": str(self.customer_id), "item_id": str(v["item_p"]), "qty": "10",
                       "due_date": due2.isoformat()}, numbers=lambda b: b["request_no"])
        self.req_nos.append(b["request_no"])
        self.req_ids.append(b["id"])
        v.update(R2=b["request_no"], R2id=b["id"])

        # 2 지시 (생산부)
        b = self.post("WO-020 생산지시서 작성", "pr", "/wo/orders",
                      {"request_id": str(R1id), "qty": "100", "equipment_code": "EQ01",
                       "planned_on": self.today.isoformat()}, numbers=lambda b: b["work_order_no"])
        WO, WOid = b["work_order_no"], b["id"]
        self.wo_nos.append(WO)
        self.wo_ids.append(WOid)
        v.update(WO=WO, WOid=WOid)
        w = self.q1("select status, request_id, item_id from work_order where id = %s", (WOid,))
        self.expect(w["status"] == "발행" and w["request_id"] == R1id and w["item_id"] == v["item_p"],
                    f"work_order {WO} 발행 · 의뢰 {R1} · 제품 품목")
        self.expect(self.q1("select status from production_request where id = %s", (R1id,))["status"] == "지시완료",
                    "의뢰 지시완료")
        n = self.notes(f"order_issued:{WO}")
        self.expect(set(n) == {"PR"} and n["PR"]["handled_at"] is None, f"order_issued:{WO} → PR 1행 미처리")
        n = self.notes(f"request_created:{R1}")
        self.expect(n["PR"]["handled_at"] is not None and n["DS"]["handled_at"] is None,
                    "request_created PR 처리됨(지시서 작성이 닫음) · DS 미처리")

        # 3 원자재 입고 (생산부)
        b = self.post("MAT-010 원자재 입고", "pr", "/mat/receipts",
                      {"supplier_name": f"{t} 공급사 (예시)", "received_on": self.today.isoformat(),
                       "item_id": [str(v["item_m"])] * 2, "qty": ["120", "80"], "location_id": [L1, L1],
                       "supplier_lot_no": ["S-1", "S-2"]}, numbers=lambda b: ", ".join(b["lot_nos"]))
        A1, A2 = b["lot_nos"]
        self.receipt_ids.append(b["receipt_id"])
        self.mlot_nos += [A1, A2]
        v.update(A1=A1, A2=A2, rcp1=b["receipt_id"])
        s1, s2 = self.stock1("원자재", A1), self.stock1("원자재", A2)
        self.expect(s1["status"] == s2["status"] == "검사대기" and s1["qty"] == 120 and s2["qty"] == 80
                    and s1["allocated_qty"] == s2["allocated_qty"] == 0, f"lot_stock {A1} 120 · {A2} 80 검사대기")
        n = self.notes(f"material_received:{A1}")
        self.expect(set(n) == {"PR"} and n["PR"]["handled_at"] is None, f"material_received:{A1} → PR 1행")

        # 4 입고검사 (생산부)
        self.post("QC-010 입고검사 합격", "pr", "/qc/incoming",
                  {"lot_no": A1, "result": "합격", "item_name": ["외관"], "item_value": ["양호"]},
                  numbers=lambda b: f"{b['lot_no']} #{b['seq']} → {b['status']}")
        self.expect(self.stock1("원자재", A1)["status"] == "가용"
                    and self.notes(f"material_received:{A1}")["PR"]["handled_at"] is None,
                    f"{A1} 가용 · 입고 알림은 아직 열림(같은 입고의 {A2} 미판정)")
        self.post("QC-010 입고검사 합격", "pr", "/qc/incoming", {"lot_no": A2, "result": "합격"},
                  numbers=lambda b: f"{b['lot_no']} #{b['seq']} → {b['status']}")
        self.expect(self.stock1("원자재", A2)["status"] == "가용"
                    and self.notes(f"material_received:{A1}")["PR"]["handled_at"] is not None,
                    f"{A2} 가용 · material_received 처리됨(입고 LOT 전부 판정)")

        # 4b 분할 (생산부 · INV-020) — 분할 LOT 를 불출해 역추적이 원 LOT 로 잇는지 본다
        b = self.post("INV-020 원자재 LOT 분할", "pr", "/inv/moves",
                      {"lot_kind": "원자재", "lot_no": A1, "from_location_id": L1, "to_location_id": L2, "qty": "20",
                       "split": "1"}, numbers=lambda b: b["lot_no"])
        A1s = b["lot_no"]
        self.mlot_nos.append(A1s)
        v["A1s"] = A1s
        par = self.q1("""select p.lot_no from material_lot c join material_lot p on p.id = c.parent_lot_id
                          where c.lot_no = %s""", (A1s,))
        self.expect(A1s == f"{A1}-01" and par and par["lot_no"] == A1
                    and self.stock1("원자재", A1)["qty"] == 100 and self.stock1("원자재", A1s)["qty"] == 20
                    and self.stock1("원자재", A1s)["status"] == "가용",
                    f"{A1s} parent={A1} · {A1} 100 · {A1s} 20 가용")

        # 5 불출 (생산부)
        b = self.post("MAT-020 불출", "pr", "/mat/issues",
                      {"work_order_no": WO, "issued_on": self.today.isoformat(),
                       "lot_no": [A1, A1s, A2], "location_id": [L1, L2, L1], "qty": ["80", "20", "50"]},
                      numbers=lambda b: f"issue {b['issue_id']} ← {', '.join(b['lot_nos'])}")
        self.issue_ids.append(b["issue_id"])
        q = {x: self.stock1("원자재", x)["qty"] for x in (A1, A1s, A2)}
        self.expect(q == {A1: 20, A1s: 0, A2: 30}, f"lot_stock qty {A1} 20 · {A1s} 0 · {A2} 30")
        mv = self.q1("""select count(*) as n, sum(qty) as s from stock_move
                         where move_kind = '불출' and ref_table = 'material_issue_line'
                           and ref_id in (select id from material_issue_line where issue_id = %s)""", (b["issue_id"],))
        self.expect(mv["n"] == 3 and mv["s"] == 150, "stock_move 불출 3행 · 합 150")
        self.expect(self.notes(f"order_issued:{WO}")["PR"]["handled_at"] is not None, "order_issued 처리됨(불출이 닫음)")

        # 6 생산일보 · 검사일보 (생산부)
        b = self.post("PRD-010 생산일보", "pr", "/prd/daily",
                      {"work_order_no": WO, "work_date": self.today.isoformat(), "equipment_code": "EQ01",
                       "produced_qty": "100", "good_qty": "95", "loss_qty": "5", "run_minutes": "480", "stop_minutes": "30"},
                      numbers=lambda b: f"daily {b['daily_id']}")
        self.expect(self.q1("select good_qty from production_daily where id = %s", (b["daily_id"],))["good_qty"] == 95,
                    "production_daily 양품 95")
        b = self.post("INS-010 검사일보", "pr", "/ins/daily",
                      {"work_order_no": WO, "work_date": self.today.isoformat(), "inspected_qty": "95", "defect_qty": "2",
                       "defect_code": ["DF01"], "defect_code_qty": ["2"]}, numbers=lambda b: f"daily {b['daily_id']}")
        self.expect(self.q1("select inspected_qty, defect_qty from inspection_daily where id = %s",
                            (b["daily_id"],)) == {"inspected_qty": 95, "defect_qty": 2}, "inspection_daily 95 · 불량 2")

        # 7 제품 입고 (생산부) — 수량을 비우면 일보 완료 수량
        b = self.post("MAT-030 제품 입고(수량 비움 = 일보 양품)", "pr", "/mat/product-receipts",
                      {"work_order_no": WO, "qty": "", "location_id": L1, "received_on": self.today.isoformat()},
                      numbers=lambda b: f"{b['lot_no']} · {b['qty']}")
        P1, P1id = b["lot_no"], b["product_lot_id"]
        self.plot_nos.append(P1)
        v.update(P1=P1, P1id=P1id)
        s = self.stock1("제품", P1)
        self.expect(Decimal(b["qty"]) == 95 and s["qty"] == 95 and s["status"] == "검사대기",
                    f"product_lot {P1} 95(= 일보 양품) · lot_stock 검사대기")
        n = self.notes(f"product_received:{P1}")
        self.expect(set(n) == {"PR"} and n["PR"]["handled_at"] is None, f"product_received:{P1} → PR 1행")
        for key, qty in (("P2", "10"), ("P3", "5")):
            b = self.post(f"MAT-030 제품 입고({key} · 차단 시험용)", "pr", "/mat/product-receipts",
                          {"work_order_no": WO, "qty": qty, "location_id": L1}, numbers=lambda b: f"{b['lot_no']} · {b['qty']}")
            self.plot_nos.append(b["lot_no"])
            v[key], v[key + "id"] = b["lot_no"], b["product_lot_id"]

        # 8 제품검사 판정 (생산부)
        self.post("QC-020 제품검사 합격", "pr", "/qc/product", {"lot_no": P1, "result": "합격"},
                  numbers=lambda b: f"{b['lot_no']} #{b['seq']} → {b['status']}")
        n = self.notes(f"judged:{P1}:1")
        self.expect(self.stock1("제품", P1)["status"] == "가용" and set(n) == {"SH", "SL"}
                    and self.notes(f"product_received:{P1}")["PR"]["handled_at"] is not None,
                    f"{P1} 가용 · judged:{P1}:1 → SH · SL · product_received 처리됨")
        self.post("QC-020 제품검사 합격(P2)", "pr", "/qc/product", {"lot_no": v["P2"], "result": "합격"},
                  numbers=lambda b: f"{b['lot_no']} #{b['seq']} → {b['status']}")

        # 9 출하 지시 (영업부 — 정본 영업부 흐름 · D-523) — 판정 알림 링크 화면에서 재고 행을 고른다
        html = self.get("SHP-010 판정 알림 링크로 열기", "sl", f"/shp/orders?lot={P1}").text
        m = re.search(r'<option value="(\d+)" selected>' + re.escape(P1), html)
        self.expect(m is not None, f"SHP-010 ?lot={P1} 가 1번 줄에 그 LOT 재고를 골라 둔다")
        sid = m.group(1)
        v["P1sid"] = sid
        b = self.post("SHP-010 출하 지시 등록(의뢰번호 · LOT 60)", "sl", "/shp/orders",
                      {"request_no": R1, "stock_id": [sid], "qty": ["60"]}, numbers=lambda b: b["shipment_no"])
        S1, S1id = b["shipment_no"], b["id"]
        self.ship_nos.append(S1)
        v.update(S1=S1, S1id=S1id)
        o = self.q1("select status, request_id, customer_id, due_date from shipment_order where id = %s", (S1id,))
        self.expect(o["status"] == "지시" and o["request_id"] == R1id and o["customer_id"] == self.customer_id
                    and o["due_date"] == due1, f"shipment_order {S1} 지시 · 의뢰 {R1} 의 거래처 · 납기")
        s = self.stock1("제품", P1)
        self.expect(s["qty"] == 95 and s["allocated_qty"] == 60 and s["qty"] - s["allocated_qty"] == 35,
                    f"{P1} qty 95 · allocated 60 · 가용 35")
        self.r.add("G-10", "할당 뒤 가용 = qty − allocated", s["qty"] - s["allocated_qty"] == 35,
                   f"{P1} qty {s['qty']} − 할당 {s['allocated_qty']} = {s['qty'] - s['allocated_qty']} (기대 35)")
        n = self.notes(f"shipment_ordered:{S1}")
        j = self.notes(f"judged:{P1}:1")
        self.expect(set(n) == {"SH"} and n["SH"]["handled_at"] is None and j["SL"]["handled_at"] is not None
                    and j["SH"]["handled_at"] is None,
                    f"shipment_ordered:{S1} → SH · judged SL 처리됨(등록 부서) · SH 미처리")
        avail_screen = self.inv_available("sl", P1)
        self.r.add("G-10", "INV-010 화면 가용 = 재계산 (할당 뒤)", avail_screen == [Decimal(35)],
                   f"INV-010 ?lot={P1} 가용 칸 {avail_screen} (기대 [35])")

        # 10 출하 확정 (출하실)
        b = self.post("SHP-020 출하 확정", "sh", f"/shp/confirm/{S1id}/confirm", {}, numbers=lambda b: b["shipment_no"])
        o = self.q1("select status, confirmed_at from shipment_order where id = %s", (S1id,))
        s = self.stock1("제품", P1)
        mv = self.q("""select qty, from_location_id from stock_move where move_kind = '출하' and ref_table = 'shipment_order'
                        and ref_id = %s""", (S1id,))
        self.expect(o["status"] == "확정" and o["confirmed_at"] is not None, f"{S1} 확정 · confirmed_at")
        self.expect(s["qty"] == 35 and s["allocated_qty"] == 0, f"{P1} qty 95→35 · allocated 60→0")
        self.r.add("G-10", "확정 뒤 qty · allocated 함께 감소 · 가용 불변",
                   s["qty"] == 35 and s["allocated_qty"] == 0 and s["qty"] - s["allocated_qty"] == 35,
                   f"확정 전 95/60 가용 35 → 확정 뒤 {s['qty']}/{s['allocated_qty']} 가용 {s['qty'] - s['allocated_qty']}")
        self.expect(len(mv) == 1 and mv[0]["qty"] == 60 and mv[0]["from_location_id"] == v["L1"], "stock_move 출하 1행 60")
        n1, n2 = self.notes(f"shipment_ordered:{S1}"), self.notes(f"shipment_confirmed:{S1}")
        self.expect(n1["SH"]["handled_at"] is not None and set(n2) == {"SL", "AC"},
                    f"shipment_ordered 처리됨(확정이 닫음) · shipment_confirmed:{S1} → SL · AC")
        v["S1_confirmed_at"] = o["confirmed_at"]

        # 의뢰번호로 이어진 행 (G-06)
        chain = self.q1(
            """select r.request_no, w.work_order_no,
                      (select count(*) from material_issue_line l where l.work_order_id = w.id) as issue_lines,
                      (select count(*) from production_daily d where d.work_order_id = w.id) as daily,
                      (select count(*) from inspection_daily d where d.work_order_id = w.id) as insp,
                      (select count(*) from product_receipt p where p.work_order_id = w.id) as receipts,
                      (select count(*) from shipment_order o where o.request_id = r.id and o.status = '확정') as shipped,
                      (select count(*) from shipment_line sl join product_lot p on p.id = sl.product_lot_id
                        where p.work_order_id = w.id) as ship_lines
                 from production_request r join work_order w on w.request_id = r.id where r.request_no = %s""", (R1,))
        ok = chain and chain["issue_lines"] == 3 and chain["daily"] == 1 and chain["insp"] == 1 and chain["receipts"] == 3 \
            and chain["shipped"] == 1 and chain["ship_lines"] == 1
        self.r.add("G-06", "의뢰번호로 이어진 행", bool(ok), f"{R1} → {dict(chain) if chain else '없음'}")
        vw = self.q1("select daily_good_qty, receipt_qty, shipped_qty from v_work_order_status where work_order_id = %s", (WOid,))
        self.r.add("G-06", "WO-040 뷰 수량 = 흐름 (일보 95 · 입고 110 · 출하 60)",
                   vw is not None and (vw["daily_good_qty"], vw["receipt_qty"], vw["shipped_qty"]) == (95, 110, 60),
                   f"v_work_order_status {WO} → {dict(vw) if vw else '없음'}")

    def inv_available(self, login: str, lot_no: str) -> list[Decimal]:
        html = self.c(login).get(f"/inv/stock?lot={lot_no}").text
        out = []
        for m in re.finditer(r'<tr data-lot="([^"]+)">(.*?)</tr>', html, re.S):
            if m.group(1) == lot_no:
                a = re.search(r"data-available>([^<]*)<", m.group(2))
                out.append(num(a.group(1)) if a else None)
        return out

    # ════════════════════════════════════════════════════════════════════════
    # G-07 · G-08 추적
    # ════════════════════════════════════════════════════════════════════════
    def indep_backward(self, shipment_no: str | None = None, product_lot_no: str | None = None) -> tuple[set, set]:
        """QA2 독립 SQL — (원 입고 LOT 번호 집합, 실제 불출 LOT 번호 집합)."""
        if shipment_no:
            start = "select sl.product_lot_id as id from shipment_line sl join shipment_order so on so.id = sl.shipment_order_id " \
                    "where so.shipment_no = %(no)s"
        else:
            start = "select id from product_lot where lot_no = %(no)s"
        rows = self.q(f"""
            with recursive up as (
                select p.id, p.parent_lot_id, p.work_order_id from product_lot p where p.id in ({start})
                union
                select p.id, p.parent_lot_id, p.work_order_id from product_lot p join up on p.id = up.parent_lot_id
            ), issued as (
                select distinct m.id, m.lot_no, m.parent_lot_id
                  from material_issue_line l join material_issue i on i.id = l.issue_id and i.canceled_at is null
                  join material_lot m on m.id = l.material_lot_id
                 where l.work_order_id in (select work_order_id from up)
            ), mup as (
                select id as start_id, id, parent_lot_id, lot_no from issued
                union
                select mup.start_id, m.id, m.parent_lot_id, m.lot_no from material_lot m join mup on m.id = mup.parent_lot_id
            )
            select (select lot_no from issued where issued.id = mup.start_id) as issued_no, mup.lot_no as root_no
              from mup where mup.parent_lot_id is null""", {"no": shipment_no or product_lot_no})
        return {r["root_no"] for r in rows}, {r["issued_no"] for r in rows}

    def indep_forward(self, material_lot_no: str) -> tuple[set, set]:
        rows = self.q("""
            with recursive down as (
                select id from material_lot where lot_no = %(no)s
                union
                select m.id from material_lot m join down on m.parent_lot_id = down.id
            ), wos as (
                select distinct l.work_order_id as id from material_issue_line l
                  join material_issue i on i.id = l.issue_id and i.canceled_at is null
                 where l.material_lot_id in (select id from down)
            ), pls as (
                select p.id, p.lot_no from product_lot p where p.work_order_id in (select id from wos)
                   and not exists (select 1 from product_receipt r where r.product_lot_id = p.id and r.canceled_at is not null)
            )
            select 'P' as k, lot_no as v from pls
            union
            select 'S', so.shipment_no from shipment_line sl join shipment_order so on so.id = sl.shipment_order_id
             where sl.product_lot_id in (select id from pls) and so.status <> '취소'""", {"no": material_lot_no})
        return {r["v"] for r in rows if r["k"] == "P"}, {r["v"] for r in rows if r["k"] == "S"}

    def indep_chain(self, shipment_no: str) -> list[dict]:
        """QA2 독립 SQL — CLAUDE.md 완료 기준 사슬을 한 행씩: 원자재 LOT(원) ← 불출 LOT ← 불출 ← 생산지시 ← 제품 입고 ← 제품 LOT ← 출하 지시.
        trace 모듈 · 뷰를 거치지 않는다. 제품 LOT 은 분할이면 입고 LOT(원)까지 올라가 그 입고 · 지시를 잡는다."""
        return self.q("""
            with recursive pup as (
                select sl.product_lot_id as start_id, p.id, p.parent_lot_id
                  from shipment_line sl join shipment_order so on so.id = sl.shipment_order_id
                  join product_lot p on p.id = sl.product_lot_id
                 where so.shipment_no = %(no)s
                union
                select pup.start_id, p.id, p.parent_lot_id from product_lot p join pup on p.id = pup.parent_lot_id
            ), mup as (
                select l.id as line_id, m.id, m.parent_lot_id from material_issue_line l join material_lot m on m.id = l.material_lot_id
                union
                select mup.line_id, m.id, m.parent_lot_id from material_lot m join mup on m.id = mup.parent_lot_id
            )
            select so.shipment_no, so.status as ship_status, sp.lot_no as shipped_lot, rp.lot_no as received_lot,
                   pr.id as receipt_id, pr.qty as receipt_qty, wo.work_order_no,
                   mi.id as issue_id, l.qty as issue_qty, im.lot_no as issued_lot, om.lot_no as origin_lot
              from shipment_order so
              join shipment_line sl on sl.shipment_order_id = so.id
              join product_lot sp on sp.id = sl.product_lot_id
              join pup on pup.start_id = sp.id and pup.parent_lot_id is null
              join product_lot rp on rp.id = pup.id
              join product_receipt pr on pr.product_lot_id = rp.id and pr.canceled_at is null
              join work_order wo on wo.id = pr.work_order_id
              join material_issue_line l on l.work_order_id = wo.id
              join material_issue mi on mi.id = l.issue_id and mi.canceled_at is null
              join material_lot im on im.id = l.material_lot_id
              join mup on mup.line_id = l.id and mup.parent_lot_id is null
              join material_lot om on om.id = mup.id
             where so.shipment_no = %(no)s
             order by om.lot_no, im.lot_no""", {"no": shipment_no})

    def completion_chain(self) -> None:
        """CLAUDE.md 완료 기준 — 출하 1건 → 원자재 LOT. 독립 사슬 SQL 의 모든 고리가 LOT-020 화면 글자 · trace.backward 와 같다."""
        from lcomfine.app import trace

        v = self.v
        S1 = v["S1"]
        rows = self.indep_chain(S1)
        html = self.get("완료 기준 — LOT-020 출하 → 원자재 LOT", "sh", f"/lot/backward?shipment_no={S1}").text
        bw = trace.backward(shipment_no=S1)
        hops = sorted({(r["origin_lot"], r["issued_lot"], r["work_order_no"], r["received_lot"], r["shipped_lot"]) for r in rows})
        words = {w for h in hops for w in h} | {S1}
        missing = sorted(w for w in words if w not in html)
        chain_txt = " / ".join(f"{o}{'(←' + i + ')' if i != o else ''} → 불출 → {w} → 제품 입고 → {pl}{'(→' + sp + ')' if sp != pl else ''} → {S1}"
                               for o, i, w, pl, sp in hops)
        self.chain = {"S1": S1, "hops": hops, "rows": rows, "missing": missing}   # 보고서 · 테스트용
        ok = (bool(rows) and not missing and {h[0] for h in hops} == set(bw["origin_material_lot_nos"]) == {v["A1"], v["A2"]}
              and {h[1] for h in hops} == set(bw["issued_material_lot_nos"]) and {h[2] for h in hops} == {v["WO"]})
        self.db(f"사슬 {len(rows)} 고리 · 화면에 빠진 글자 {missing or 0}")
        self.r.add("G-07", "완료 기준 사슬 — 출하 1건 → 원자재 LOT (CLAUDE.md · 독립 사슬 SQL = LOT-020 화면 = trace)", ok,
                   f"{chain_txt} · 고리 {len(rows)} (불출 라인 × 원 LOT) · 화면에 빠진 글자 {missing or 0} · "
                   f"origin 독립 {sorted({h[0] for h in hops})} = trace {sorted(bw['origin_material_lot_nos'])} · "
                   f"불출 LOT 독립 {sorted({h[1] for h in hops})} = trace {sorted(bw['issued_material_lot_nos'])}")

    def traces(self) -> None:
        from lcomfine.app import trace

        v = self.v
        S1, P1, A1, A2, A1s = v["S1"], v["P1"], v["A1"], v["A2"], v["A1s"]
        html = self.get("LOT-020 역추적(출하 번호)", "sl", f"/lot/backward?shipment_no={S1}").text
        on_screen = all(x in html for x in (S1, P1, v["WO"], A1, A2, A1s))
        bw = trace.backward(shipment_no=S1)
        origin, issued = self.indep_backward(shipment_no=S1)
        self.db(f"화면에 {S1} · {P1} · {v['WO']} · {A1} · {A2} · {A1s} {'있음' if on_screen else '빠짐'}")
        self.r.add("G-07", "LOT-020 화면이 원자재 LOT 까지", on_screen,
                   f"GET /lot/backward?shipment_no={S1} 200 · 화면 글자 {S1} → {P1} → {v['WO']} → {A1}/{A1s}/{A2} {'전부' if on_screen else '빠짐'}")
        self.r.add("G-07", "trace.backward(출하) = 독립 SQL",
                   set(bw["origin_material_lot_nos"]) == origin == {A1, A2} and set(bw["issued_material_lot_nos"]) == issued,
                   f"origin {sorted(bw['origin_material_lot_nos'])} · 독립 {sorted(origin)} · 기대 {{{A1}, {A2}}} · "
                   f"불출 LOT {sorted(bw['issued_material_lot_nos'])} = 독립 {sorted(issued)} (분할 {A1s} → 원 {A1})")
        self.completion_chain()
        bw2 = trace.backward(product_lot_no=P1)
        o2, _ = self.indep_backward(product_lot_no=P1)
        self.r.add("G-07", "trace.backward(제품 LOT) = 독립 SQL", set(bw2["origin_material_lot_nos"]) == o2 == {A1, A2},
                   f"{P1} → {sorted(bw2['origin_material_lot_nos'])} · 독립 {sorted(o2)}")
        # 전 출하 대조 (DB 에 있는 모든 출하 — 다른 사람이 지우는 중이면 건너뛴다)
        mism, n, gone = [], 0, 0
        for r in self.q("select shipment_no from shipment_order where status <> '취소' order by id"):
            try:
                got = set(trace.backward(shipment_no=r["shipment_no"])["origin_material_lot_nos"])
            except trace.NotFound:
                gone += 1
                continue
            exp, _ = self.indep_backward(shipment_no=r["shipment_no"])
            n += 1
            if got != exp:
                mism.append(f"{r['shipment_no']}: trace {sorted(got)} ≠ 독립 {sorted(exp)}")
        self.r.add("G-07", "DB 의 모든 출하 trace.backward = 독립 SQL", not mism,
                   f"출하 {n} 건 대조 · 불일치 {len(mism)} {mism[:3] or ''} · 대조 중 사라짐 {gone}")
        # G-08
        html = self.get("LOT-010 정방향(원자재 LOT)", "pr", f"/lot/forward?lot_no={A1}").text
        on_screen = all(x in html for x in (A1, A1s, v["WO"], P1, S1))
        self.db(f"화면에 {A1} · {A1s} · {v['WO']} · {P1} · {S1} {'있음' if on_screen else '빠짐'}")
        fw = trace.forward(A1)
        pls, shs = self.indep_forward(A1)
        self.r.add("G-08", "LOT-010 화면이 출하까지", on_screen,
                   f"GET /lot/forward?lot_no={A1} 200 · {A1} → {A1s} → {v['WO']} → {P1} → {S1} {'전부' if on_screen else '빠짐'}")
        self.r.add("G-08", "trace.forward = 독립 SQL", set(fw["product_lot_nos"]) == pls and set(fw["shipment_nos"]) == shs
                   and S1 in shs and P1 in pls,
                   f"{A1}: 제품 LOT {sorted(fw['product_lot_nos'])} = 독립 {sorted(pls)} · 출하 {sorted(fw['shipment_nos'])} = 독립 {sorted(shs)}")
        mism, n = [], 0
        for r in self.q("select lot_no from material_lot where canceled_at is null and parent_lot_id is null order by id"):
            try:
                f = trace.forward(r["lot_no"])
            except trace.NotFound:
                continue
            p, s = self.indep_forward(r["lot_no"])
            n += 1
            if set(f["product_lot_nos"]) != p or set(f["shipment_nos"]) != s:
                mism.append(r["lot_no"])
        self.r.add("G-08", "DB 의 모든 원 입고 LOT trace.forward = 독립 SQL", not mism,
                   f"원자재 LOT {n} 건 대조 · 불일치 {len(mism)} {mism[:3] or ''}")

    # ════════════════════════════════════════════════════════════════════════
    # G-09 차단
    # ════════════════════════════════════════════════════════════════════════
    def blocks(self) -> None:
        v, t = self.v, self.tag
        L1 = str(v["L1"])
        b = self.post("MAT-010 원자재 입고(차단 시험 B · C)", "pr", "/mat/receipts",
                      {"supplier_name": f"{t} 공급사 (예시)", "item_id": [str(v["item_m"])] * 2, "qty": ["10", "10"],
                       "location_id": [L1, L1]}, numbers=lambda b: ", ".join(b["lot_nos"]))
        B, C = b["lot_nos"]
        self.receipt_ids.append(b["receipt_id"])
        self.mlot_nos += [B, C]
        self.post("QC-010 입고검사 불합격(C)", "pr", "/qc/incoming", {"lot_no": C, "result": "불합격"},
                  numbers=lambda b: f"{b['lot_no']} → {b['status']}")
        self.expect(self.stock1("원자재", C)["status"] == "보류", f"{C} 보류")
        before = self.q1("select count(*) as n from material_issue_line where work_order_id = %s", (v["WOid"],))["n"]
        res = {}
        for lot, label in ((B, "검사대기"), (C, "보류")):
            code, body = self.try_post(f"MAT-020 불출({label} LOT)", "pr", "/mat/issues",
                                       {"work_order_no": v["WO"], "lot_no": [lot], "location_id": [L1], "qty": ["1"]})
            res[label] = code
        after = self.q1("select count(*) as n from material_issue_line where work_order_id = %s", (v["WOid"],))["n"]
        qb, qc = self.stock1("원자재", B)["qty"], self.stock1("원자재", C)["qty"]
        self.r.add("G-09", "검사대기 · 보류 LOT 불출 422", res == {"검사대기": 422, "보류": 422} and before == after and qb == qc == 10,
                   f"불출 {B}(검사대기) {res['검사대기']} · {C}(보류) {res['보류']} · 불출 행 {before}→{after} · 재고 {qb}/{qc}")
        code, _ = self.try_post("MAT-010 입고(사용 중지 품목)", "pr", "/mat/receipts",
                                {"supplier_name": f"{t} 공급사 (예시)", "item_id": [str(v["item_x"])], "qty": ["1"],
                                 "location_id": [L1]})
        n = self.q1("select count(*) as n from material_lot where item_id = %s", (v["item_x"],))["n"]
        self.r.add("G-09", "사용 중지 품목 입고 422", code == 422 and n == 0, f"중지 품목 {t}-X 입고 {code} · LOT {n}")
        # 검사대기 제품 LOT 할당 → 422 (화면에는 가용 LOT 만 나온다 — 재고 행 id 를 직접 넣어 본다)
        p3sid = str(self.stock1("제품", v["P3"])["id"])
        html = self.c("sl").get("/shp/orders").text
        listed = f'value="{p3sid}"' in html
        code, _ = self.try_post("SHP-010 출하 지시(검사대기 LOT)", "sl", "/shp/orders",
                                {"customer_id": str(self.customer_id), "due_date": self.today.isoformat(),
                                 "stock_id": [p3sid], "qty": ["1"]})
        self.r.add("G-09", "검사대기 제품 LOT 할당 422 · 목록에 없음", code == 422 and not listed
                   and self.stock1("제품", v["P3"])["allocated_qty"] == 0,
                   f"{v['P3']} 할당 {code} · SHP-010 LOT 목록에 {'있음' if listed else '없음'}")
        # 보류(재판정 불합격) LOT 이 걸린 출하 확정 → 422
        p2sid = str(self.stock1("제품", v["P2"])["id"])
        b = self.post("SHP-010 출하 지시(P2 10)", "sl", "/shp/orders",
                      {"request_no": v["R1"], "stock_id": [p2sid], "qty": ["10"]}, numbers=lambda b: b["shipment_no"])
        S2, S2id = b["shipment_no"], b["id"]
        self.ship_nos.append(S2)
        v.update(S2=S2, S2id=S2id)
        self.post("QC-020 제품검사 재판정 불합격(P2)", "pr", f"/qc/product/{v['P2id']}", {"result": "불합격"},
                  numbers=lambda b: f"{b['lot_no']} #{b['seq']} → {b['status']}")
        st = self.stock1("제품", v["P2"])
        code, body = self.try_post("SHP-020 출하 확정(재판정 불합격 LOT)", "sh", f"/shp/confirm/{S2id}/confirm", {})
        st2 = self.stock1("제품", v["P2"])
        o = self.q1("select status from shipment_order where id = %s", (S2id,))
        self.r.add("G-09", "보류(불합격) LOT 출하 확정 422", code == 422 and o["status"] == "지시" and st == st2
                   and st["status"] == "보류" and st["allocated_qty"] == 10,
                   f"{S2} 확정 {code} · 상태 {o['status']} · {v['P2']} {st2['status']} qty {st2['qty']} 할당 {st2['allocated_qty']} 그대로(D-515 2)")
        # 불합격 → 합격 재판정은 보류 → 가용 (D-515 4) — 원자재 C 로 본다
        cid = self.q1("select id from material_lot where lot_no = %s", (C,))["id"]
        self.post("QC-010 입고검사 재판정 합격(C)", "pr", f"/qc/incoming/{cid}", {"result": "합격"},
                  numbers=lambda b: f"{b['lot_no']} #{b['seq']} → {b['status']}")
        stc = self.stock1("원자재", C)
        seqs = [r["result"] for r in self.q("select result from incoming_inspection where material_lot_id = %s order by seq", (cid,))]
        self.r.add("G-09", "재판정 불합격 → 합격 = 보류 → 가용 · 이전 판정 행 유지", stc["status"] == "가용" and seqs == ["불합격", "합격"],
                   f"{C} 판정 이력 {seqs} · 재고 {stc['status']} (D-515 4 · 새 행)")
        # 동시 불출 — 같은 원자재 LOT(가용 20)에 15 × 2
        A1 = v["A1"]
        s0 = self.stock1("원자재", A1)
        cl = [self.TestClient(self.app) for _ in range(2)]
        for x in cl:
            x.post("/login", data={"login_id": "pr", "password": self.pw}, follow_redirects=False)
        res: list = []
        bar = threading.Barrier(2)

        def go(x):
            bar.wait()
            resp = x.post("/mat/issues", data={"work_order_no": v["WO"], "lot_no": [A1], "location_id": [L1], "qty": ["15"]})
            res.append(resp.status_code)
            try:
                body = resp.json()
                if body.get("issue_id"):
                    self.issue_ids.append(body["issue_id"])
            except Exception:  # noqa: BLE001
                pass

        ths = [threading.Thread(target=go, args=(x,)) for x in cl]
        for th in ths:
            th.start()
        for th in ths:
            th.join(60)
        for code in res:
            self.r.steps.append(Step("MAT-020 동시 불출(스레드 2 · 각 15 / 가용 20)", "pr", "POST /mat/issues", code))
        s1 = self.stock1("원자재", A1)
        self.r.add("G-10", "동시 불출(스레드 2) 초과 없음", sorted(res) == [200, 422] and s1["qty"] == s0["qty"] - 15,
                   f"{A1} 가용 {s0['qty'] - s0['allocated_qty']} 에 15 × 2 동시 → 응답 {sorted(res)} · 남은 qty {s1['qty']}")

    # ════════════════════════════════════════════════════════════════════════
    # G-10 동시 할당 · 전 LOT 재계산 · 사용 중지 품목 출하 지시
    # ════════════════════════════════════════════════════════════════════════
    def allocation(self) -> None:
        v = self.v
        P1, sid = v["P1"], v["P1sid"]
        s0 = self.stock1("제품", P1)          # 확정 뒤 35 / 0
        results: dict[str, tuple[int, dict]] = {}
        clients = {lg: self.c(lg) for lg in ("sl", "sh")}
        barrier = threading.Barrier(2)

        def go(login: str) -> None:
            barrier.wait()
            resp = clients[login].post("/shp/orders", data={"request_no": v["R1"], "stock_id": [sid], "qty": ["30"]})
            try:
                results[login] = (resp.status_code, resp.json())
            except Exception:  # noqa: BLE001
                results[login] = (resp.status_code, {})

        ths = [threading.Thread(target=go, args=(lg,)) for lg in ("sl", "sh")]
        for th in ths:
            th.start()
        for th in ths:
            th.join(60)
        codes = sorted(c for c, _ in results.values())
        for lg, (c, body) in results.items():
            self.r.steps.append(Step("SHP-010 동시 할당(스레드 2 · 각 30 / 가용 35)", lg, "POST /shp/orders", c,
                                     body.get("shipment_no") or (body.get("detail") or {}).get("message", "")))
            if body.get("shipment_no"):
                self.ship_nos.append(body["shipment_no"])
                v.update(S3=body["shipment_no"], S3id=body["id"])
        s1 = self.stock1("제품", P1)
        self.r.add("G-10", "동시 할당(스레드 2) 초과 없음", codes == [200, 422] and s1["allocated_qty"] == 30 and s1["qty"] == 35,
                   f"가용 {s0['qty'] - s0['allocated_qty']} 에 30 × 2 동시 → 응답 {codes} · {P1} qty {s1['qty']} 할당 {s1['allocated_qty']}")
        # F-SHP-012 할당 다시 잡기 → 25
        if "S3id" in v:
            self.post("SHP-010 출하 지시 수정(할당 30 → 25)", "sl", f"/shp/orders/{v['S3id']}", {"stock_id": [sid], "qty": ["25"]})
            s2 = self.stock1("제품", P1)
            self.r.add("G-10", "수정으로 할당 다시 잡기", s2["allocated_qty"] == 25, f"{P1} 할당 30 → {s2['allocated_qty']} (기대 25)")
        # 사용 중지 제품 품목 → 출하 지시 422, 다시 사용
        self.post("BAS-010 제품 품목 사용 중지", "pr", f"/bas/items/{v['item_p']}",
                  {"item_name": f"{self.tag} 제품 (예시)", "item_type": "제품", "unit": "EA", "use_yn": "N"})
        html = self.c("sl").get("/shp/orders").text
        listed = f'value="{sid}"' in html
        code, _ = self.try_post("SHP-010 출하 지시(사용 중지 품목 LOT)", "sl", "/shp/orders",
                                {"request_no": v["R1"], "stock_id": [sid], "qty": ["1"]})
        self.post("BAS-010 제품 품목 다시 사용", "pr", f"/bas/items/{v['item_p']}",
                  {"item_name": f"{self.tag} 제품 (예시)", "item_type": "제품", "unit": "EA", "use_yn": "Y"})
        self.r.add("G-09", "사용 중지 품목 출하 지시 422 · 목록에 없음", code == 422 and not listed
                   and self.stock1("제품", P1)["allocated_qty"] == 25,
                   f"{self.tag}-P 중지 중 {P1} 할당 {code} · SHP-010 LOT 목록에 {'있음' if listed else '없음'}")
        self.global_allocation()

    def global_allocation(self) -> None:
        """전 LOT: allocated_qty = 열린(지시) 출하 행 합 · 가용 = qty − allocated · INV-010 화면 값 = DB 재계산."""
        rows = self.q("""
            select s.id, s.lot_kind, coalesce(m.lot_no, p.lot_no) as lot_no, w.warehouse_code, l.location_code,
                   s.qty, s.allocated_qty,
                   coalesce((select sum(sl.qty) from shipment_line sl join shipment_order so on so.id = sl.shipment_order_id
                              where so.status = '지시' and s.lot_kind = '제품' and sl.product_lot_id = s.lot_id
                                and sl.location_id = s.location_id), 0) as open_lines
              from lot_stock s
              left join material_lot m on s.lot_kind = '원자재' and m.id = s.lot_id
              left join product_lot p on s.lot_kind = '제품' and p.id = s.lot_id
              join location l on l.id = s.location_id join warehouse w on w.id = l.warehouse_id
             where m.canceled_at is null""")
        bad = [f"{r['lot_no']}@{r['location_code']}: 할당 {r['allocated_qty']} ≠ 열린 출하 행 {r['open_lines']}"
               for r in rows if r["allocated_qty"] != r["open_lines"]]
        self.r.add("G-10", "전 LOT 할당 = 열린 출하 행 합 (독립 재계산)", not bad,
                   f"lot_stock {len(rows)} 행 · 불일치 {len(bad)} {bad[:3] or ''}")
        html = self.c("admin").get("/inv/stock").text
        screen = {}
        for m in re.finditer(r'<tr data-lot="([^"]+)">(.*?)</tr>', html, re.S):
            tds = re.findall(r"<td[^>]*>(.*?)</td>", m.group(2), re.S)
            if len(tds) >= 9:
                key = (m.group(1), _html.unescape(re.sub("<[^>]+>", "", tds[4])).strip(),
                       _html.unescape(re.sub("<[^>]+>", "", tds[5])).strip())
                screen[key] = num(re.sub("<[^>]+>", "", tds[8]))
        dbv = {(r["lot_no"], r["warehouse_code"], r["location_code"]): r["qty"] - r["allocated_qty"] for r in rows}
        mism = [f"{k}: 화면 {screen.get(k)} ≠ {dbv[k]}" for k in dbv if k in screen and screen[k] != dbv[k]]
        missing = [k for k in dbv if k not in screen]
        self.r.add("G-10", "INV-010 화면 가용 = qty − allocated (전 행)", not mism and not missing,
                   f"화면 {len(screen)} 행 · DB {len(dbv)} 행 · 값 불일치 {len(mism)} {mism[:3] or ''} · 화면에 없음 {len(missing)} {missing[:3] or ''}")

    # ════════════════════════════════════════════════════════════════════════
    # 재판정 후 출하 (Review Focus 3 · D-515 3 · D-702)
    # ════════════════════════════════════════════════════════════════════════
    def rejudge_after_shipment(self) -> None:
        from lcomfine.app import trace

        v = self.v
        S1, P1 = v["S1"], v["P1"]
        before = self.q1("""select o.status, o.confirmed_at, (select sum(qty) from shipment_line where shipment_order_id = o.id) as q,
                                   (select count(*) from shipment_adjust where shipment_order_id = o.id) as adj
                              from shipment_order o where o.id = %s""", (v["S1id"],))
        b = self.post("QC-020 출하 뒤 재판정 불합격(P1)", "pr", f"/qc/product/{v['P1id']}", {"result": "불합격"},
                      numbers=lambda b: f"{b['lot_no']} #{b['seq']} → {b['status']} · shipped_after {b['shipped_after']}")
        after = self.q1("""select o.status, o.confirmed_at, (select sum(qty) from shipment_line where shipment_order_id = o.id) as q,
                                  (select count(*) from shipment_adjust where shipment_order_id = o.id) as adj
                             from shipment_order o where o.id = %s""", (v["S1id"],))
        st = self.stock1("제품", P1)
        flag = [x for x in trace.backward(shipment_no=S1)["rejected_after_shipment"] if x["product_lot_no"] == P1]
        html = self.get("LOT-020 역추적(재판정 뒤)", "sl", f"/lot/backward?shipment_no={S1}").text
        self.db(f"출하 {before['status']}→{after['status']} · 수량 {before['q']}→{after['q']} · 반품정정 {after['adj']} · "
                f"재고 {st['status']} · 「출하 후 불합격」 {'있음' if '출하 후 불합격' in html else '없음'}")
        self.r.add("G-07", "출하 후 불합격 — 출하 유지 · 역추적 표시",
                   before == after and b.get("shipped_after") is True and st["status"] == "보류" and bool(flag)
                   and "출하 후 불합격" in html,
                   f"{S1} {after['status']} · 확정 시각 · 수량 {after['q']} 그대로 · shipment_adjust {after['adj']} · "
                   f"남은 재고 {st['qty']} {st['status']} · trace.rejected_after_shipment {len(flag)} · LOT-020 「출하 후 불합격」 "
                   f"{'있음' if '출하 후 불합격' in html else '없음'}")
        if "S3id" in v:
            code, _ = self.try_post("SHP-020 출하 확정(재판정 불합격 P1 할당분)", "sh", f"/shp/confirm/{v['S3id']}/confirm", {})
            self.r.add("G-09", "출하 뒤 재판정 불합격 LOT 의 남은 지시 확정 422", code == 422, f"{v['S3']} 확정 {code}")

    # ════════════════════════════════════════════════════════════════════════
    # G-11 알림
    # ════════════════════════════════════════════════════════════════════════
    def notifications(self) -> None:
        v = self.v
        expected = {
            "request_created": (f"request_created:{v['R1']}", {"PR", "DS"}),
            "order_issued": (f"order_issued:{v['WO']}", {"PR"}),
            "material_received": (f"material_received:{v['A1']}", {"PR"}),
            "product_received": (f"product_received:{v['P1']}", {"PR"}),
            "judged": (f"judged:{v['P1']}:1", {"SH", "SL"}),
            "shipment_ordered": (f"shipment_ordered:{v['S1']}", {"SH"}),
            "shipment_confirmed": (f"shipment_confirmed:{v['S1']}", {"SL", "AC"}),
        }
        closed_by_screen = {"request_created": {"PR"}, "order_issued": {"PR"}, "material_received": {"PR"},
                            "product_received": {"PR"}, "judged": {"SL"}, "shipment_ordered": {"SH"}}
        for ev, (key, roles) in expected.items():
            n = self.notes(key)
            ok = set(n) == roles and all(x["event_key"] == ev for x in n.values())
            closed = closed_by_screen.get(ev, set())
            ok_closed = all(n.get(r, {}).get("handled_at") is not None for r in closed)
            self.r.add("G-11", f"{ev}", ok and ok_closed,
                       f"{key} → {sorted(n)} (기대 {sorted(roles)}) · 처리 화면이 닫음 {sorted(closed) or '-'} "
                       f"{'OK' if ok_closed else '안 닫힘'}")
        # dedupe — 같은 키 재발생은 무시 (notify.emit 을 같은 키로 다시 불러 본다: on conflict do nothing → 0행)
        from lcomfine.app import notify

        key = f"request_created:{v['R1']}"
        with self.conn.tx() as cur:
            again = notify.emit(cur, "request_created", link="/wo/requests", dedupe_key=key, text="QA2 dedupe 재발생")
        rows = self.q("select count(*) as n from notification where dedupe_key = %s", (key,))[0]["n"]
        self.r.add("G-11", "같은 키 재발생 무시 (D-522)", again == 0 and rows == 2, f"emit 재호출 새 행 {again} · {key} 행 {rows}")
        # HOME-020 처리 — 디자인·개발실의 request_created(처리 화면이 없는 부서)
        ds_id = self.notes(key)["DS"]["id"]
        html = self.get("HOME-020 알림함(ds)", "ds", "/home/notifications").text
        listed = f"/home/notifications/{ds_id}/handle" in html
        code, _ = self.try_post("HOME-020 다른 부서 알림 처리", "sh", f"/home/notifications/{ds_id}/handle", {})
        self.post("HOME-020 알림 처리", "ds", f"/home/notifications/{ds_id}/handle", {})
        n = self.notes(key)["DS"]
        html2 = self.c("ds").get("/home/notifications").text
        still = f"/home/notifications/{ds_id}/handle" in html2
        self.r.add("G-11", "HOME-020 처리 → 부서 처리됨 · 남의 부서 404",
                   listed and code == 404 and n["handled_by"] == "ds" and not still,
                   f"ds 알림함에 처리 버튼 {'있음' if listed else '없음'} → 처리 후 {'남음' if still else '사라짐'} · "
                   f"handled_by {n['handled_by']} · sh 가 처리 시도 {code}")
        code2, _ = self.try_post("HOME-020 이미 처리한 알림 다시 처리", "ds", f"/home/notifications/{ds_id}/handle", {})
        self.r.add("G-11", "이미 처리된 알림 다시 처리 200", code2 == 200, f"재처리 {code2}")
        # due_d3_unshipped — 배치
        rc1, out1 = self.due_check()
        k = f"due:{v['R2']}:{self.today.isoformat()}"
        n1 = self.notes(k)
        rc2, out2 = self.due_check()
        n2 = self.notes(k)
        r1_due = self.q("select count(*) as n from notification where dedupe_key like %s", (f"due:{v['R1']}:%",))[0]["n"]
        self.r.add("G-11", "due_d3_unshipped (jobs due-check)",
                   rc1 == rc2 == 0 and set(n1) == {"SL", "PR"} and set(n2) == {"SL", "PR"}
                   and {r: x["id"] for r, x in n1.items()} == {r: x["id"] for r, x in n2.items()} and r1_due == 0,
                   f"`python -m lcomfine.jobs due-check` ×2 rc {rc1}/{rc2} · {k} → {sorted(n1)} · 2회째 같은 행 · "
                   f"납기 D-10 의뢰 {v['R1']} 알림 {r1_due} · 출력 「{out2.strip().splitlines()[-1] if out2.strip() else ''}」")
        # erp_upload_error 는 수금 단계(collections)가 잰다 — 행이 없으면 run() 끝에서 FAIL 로 채운다

    def due_check(self) -> tuple[int, str]:
        p = subprocess.run(["uv", "run", "python", "-m", "lcomfine.jobs", "due-check"], cwd=ROOT, capture_output=True,
                           text=True, timeout=300)
        self.r.steps.append(Step("jobs due-check", "(배치)", "uv run python -m lcomfine.jobs due-check", p.returncode,
                                 (p.stdout or p.stderr).strip()[:160]))
        return p.returncode, p.stdout + p.stderr

    # ════════════════════════════════════════════════════════════════════════
    # G-05 동적 — 조회 화면 GET 전후 행 수
    # ════════════════════════════════════════════════════════════════════════
    def table_counts(self) -> dict[str, int]:
        names = [r["table_name"] for r in self.q("""select table_name from information_schema.tables
                                                    where table_schema = 'public' and table_type = 'BASE TABLE'""")]
        return {n: self.q1(f'select count(*) as n from "{n}"')["n"] for n in names if n not in LOG_TABLES}

    def readonly_diff(self) -> None:
        v = self.v
        paths = READ_SCREENS + [f"/lot/backward?shipment_no={v.get('S1', '')}", f"/lot/forward?lot_no={v.get('A1', '')}",
                                f"/trk/orders?q={v.get('R1', '')}", "/doc/go", "/doc/account-map"]   # W4 DOC-010 · 020 GET
        diff = None
        for _ in range(3):          # 같은 DB 를 다른 QA 가 함께 쓴다 — 바깥 변화면 다시 잰다
            before = self.table_counts()
            codes = [self.c("admin").get(p).status_code for p in paths]
            after = self.table_counts()
            diff = {k: (before[k], after.get(k)) for k in before if before[k] != after.get(k)}
            if not diff:
                break
        self.r.add("G-05", "조회 화면 GET 전후 업무 테이블 diff 0",
                   not diff and all(c == 200 for c in codes),
                   f"admin GET {len(paths)} 화면(TRK 3 · LOT 2 · PRP 3 · WO-040 · INV-010 · QC-030 · SHP-030 · HOME · COL 7 · DOC 2) "
                   f"응답 {sorted(set(codes))} · 테이블 diff {diff or 0} (로그 4 — 조회 · 변경 · 질의 · AI 로그 — 는 다른 세션이 동시에 쓰므로 제외 · 질의 로그는 정적 행이 보증)")

    # ════════════════════════════════════════════════════════════════════════
    # G-10 — F-SHP-012 라인 전부 제거 = 지시 취소 (function-list 계약 문장 · Ruling 21) → 할당 해제 · 가용 복원
    # ════════════════════════════════════════════════════════════════════════
    def shp012_cancel(self) -> None:
        """취소 뒤 할당 0 · 가용 = qty 복원 · 독립 SQL(열린 출하 행 합) 대조 · INV-010 화면 값 · 알림 없음."""
        v = self.v
        if "S2id" not in v:
            self.r.add("G-10", "F-SHP-012 라인 전부 제거 = 지시 취소 · 할당 해제 · 가용 복원 (Ruling 21)", FAIL,
                       "앞 단계에서 S2 출하 지시가 만들어지지 않았다")
            return
        P2 = v["P2"]
        before = self.stock1("제품", P2)
        lb = self.q1("select count(*) as n, coalesce(sum(qty), 0) as s from shipment_line where shipment_order_id = %s", (v["S2id"],))
        ev_before = self.q1("select count(*) as n from notification where dedupe_key like %s", (f"%:{v['S2']}",))["n"]
        code, body = self.try_post("SHP-010 출하 지시 수정(라인 전부 제거 = 취소)", "sl", f"/shp/orders/{v['S2id']}",
                                   {"stock_id": [""], "qty": [""]})
        o = self.q1("select status from shipment_order where id = %s", (v["S2id"],))
        after = self.stock1("제품", P2)
        la = self.q1("select count(*) as n from shipment_line where shipment_order_id = %s", (v["S2id"],))["n"]
        # 독립 재계산 — 그 재고 행의 할당 = 열린(지시) 출하 행 합. 취소 뒤 0 이어야 하고 가용 = qty − 0 = qty
        indep = self.q1("""select coalesce(sum(sl.qty), 0) as open_alloc
                             from shipment_line sl join shipment_order so on so.id = sl.shipment_order_id
                            where so.status = '지시' and sl.product_lot_id = %s and sl.location_id = %s""",
                        (v["P2id"], v["L1"]))["open_alloc"]
        ev_after = self.q1("select count(*) as n from notification where dedupe_key like %s", (f"%:{v['S2']}",))["n"]
        avail_screen = self.inv_available("sl", P2)
        self.db(f"응답 {code} · 상태 {o['status']} · {P2} 할당 {before['allocated_qty']}→{after['allocated_qty']} · 독립 열린 출하 행 합 {indep} · "
                f"shipment_line {lb['n']}→{la} · INV-010 가용 {avail_screen} · 알림 행 {ev_before}→{ev_after}")
        ok = (code == 200 and o["status"] == "취소" and lb["n"] == 1 and la == 0
              and before["allocated_qty"] == lb["s"] and after["allocated_qty"] == 0 == indep
              and after["qty"] == before["qty"] and avail_screen == [after["qty"]] and ev_after == ev_before)
        self.r.add("G-10", "F-SHP-012 라인 전부 제거 = 지시 취소 · 할당 해제 · 가용 복원 (Ruling 21)", ok,
                   f"{v['S2']} 라인 {lb['n']}→{la} → 응답 {code} · 상태 {o['status']} · {P2} qty {after['qty']} 할당 "
                   f"{before['allocated_qty']}→{after['allocated_qty']} = 독립 SQL(열린 출하 행 합) {indep} · 가용 {after['qty'] - after['allocated_qty']} "
                   f"= INV-010 {avail_screen} · 알림 없음({ev_before}→{ev_after})")

    # ════════════════════════════════════════════════════════════════════════
    # W2 수금 — G-11 erp_upload_error · G-12 회차 교체 · G-13 오류 행 · 미매핑 · G-14 출하 대비 수금
    # ════════════════════════════════════════════════════════════════════════
    @staticmethod
    def xlsx(headers: list[str], rows: list[list]) -> bytes:
        """(예시) ERP 수금현황 엑셀 — 첫 시트 · 1행 열 이름. 값은 글자 그대로(원문 보존을 본다)."""
        import io

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        for r in rows:
            ws.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def post_file(self, step: str, login: str, period: tuple[date, date], headers: list[str], rows: list[list],
                  action: str = "apply", expect: int = 200) -> dict:
        """COL-010 `POST /col/uploads` — 화면 폼과 같은 multipart(action · period_from · period_to · file)."""
        f, t = period
        resp = self.c(login).post("/col/uploads", data={"action": action, "period_from": f.isoformat(), "period_to": t.isoformat()},
                                  files={"file": (f"{self.tag}-erp (예시).xlsx", self.xlsx(headers, rows), XLSX_MIME)},
                                  follow_redirects=False)
        try:
            body = resp.json()
        except Exception:  # noqa: BLE001
            body = {"text": resp.text[:200]}
        det = body.get("detail") if isinstance(body.get("detail"), dict) else {}
        nums = (f"{body.get('batch_no')} {body.get('round_no')}회차 · 행 {body.get('rows')} · 오류 {body.get('errors')} · 교체 {body.get('replaced')}"
                if resp.status_code == 200 and action == "apply" else (det.get("message") or body.get("message", "")))
        self.r.steps.append(Step(step, login, f"POST /col/uploads action={action} {f}~{t} ({len(rows)}행)", resp.status_code, nums))
        if resp.status_code != expect:
            raise StepFailed(f"{step}: POST /col/uploads ({login}) → {resp.status_code} (기대 {expect}) {str(body)[:300]}")
        if action == "apply" and resp.status_code == 200:
            self.erp_batches.append(body["id"])
        return body

    def free_period(self, f: date, t: date) -> tuple[date, date]:
        """다른 사람의 회차와 (기준 기간 시작, 끝)이 겹치지 않는 기간 — 같은 기간 재업로드는 그 회차를 교체하기 때문이다."""
        while t > f and self.q1("select 1 from erp_upload_batch where period_from = %s and period_to = %s", (f, t)):
            t -= timedelta(days=1)
        return f, t

    @staticmethod
    def row_cells(html: str, attrs: str) -> list[str] | None:
        """`<tr {attrs}>` 한 행의 td 글자(태그 제거). 없으면 None."""
        m = re.search(r"<tr " + attrs + r">(.*?)</tr>", html, re.S)
        if not m:
            return None
        return [_html.unescape(re.sub("<[^>]+>", "", td)).strip() for td in re.findall(r"<td[^>]*>(.*?)</td>", m.group(1), re.S)]

    @staticmethod
    def cell(text: str | None) -> Decimal | None:
        """화면 숫자 칸 → Decimal. 빈 칸(`-`) 은 None."""
        from lcomfine.app.util.screen import EMPTY

        if text is None or text.strip() in ("", EMPTY):
            return None
        return num(text)

    def col070(self, ym: str, customer_id: int | None = None) -> dict[int, tuple]:
        """COL-070 화면의 (월 = ym) 거래처 행 → {customer_id: (출하 건수, 출하 수량, 매출, 수금)} (빈 칸 None)."""
        q = f"/col/shipment-vs-collection?ym_from={ym}&ym_to={ym}" + (f"&customer={customer_id}" if customer_id else "")
        html = self.get("COL-070 출하 대비 수금", "ac", q).text
        out = {}
        for m in re.finditer(r'<tr data-month="' + re.escape(ym) + r'" data-customer="(\d+)">(.*?)</tr>', html, re.S):
            tds = [_html.unescape(re.sub("<[^>]+>", "", td)).strip() for td in re.findall(r"<td[^>]*>(.*?)</td>", m.group(2), re.S)]
            out[int(m.group(1))] = tuple(self.cell(x) for x in tds[2:6])
        return out

    def indep_ship_vs_coll(self, ym: str, customer_id: int) -> tuple:
        """QA2 독립 SQL — (확정 출하 건수, 출하 수량, ERP 매출 합, ERP 수금 합). 뷰를 거치지 않고 테이블에서 직접:
        출하 = shipment_order status '확정' · confirmed_at 의 월 / ERP = erp_upload_row ← 반영 회차 · error 없음 · 사용 중 매핑 · txn_date 의 월."""
        s = self.q1("""select count(distinct o.id) as n, sum(l.qty) as q
                         from shipment_order o join shipment_line l on l.shipment_order_id = o.id
                        where o.status = '확정' and o.customer_id = %s and to_char(o.confirmed_at, 'YYYY-MM') = %s""", (customer_id, ym))
        e = self.q1("""select sum(nullif(u.parsed ->> 'sales_amount', '')::numeric) as s,
                              sum(nullif(u.parsed ->> 'collected_amount', '')::numeric) as c, count(*) as n
                         from erp_upload_row u
                         join erp_upload_batch b on b.id = u.batch_id and b.status = '반영'
                         join erp_customer_map m on m.erp_code = u.parsed ->> 'customer_code' and m.use_yn = 'Y'
                        where u.error is null and m.customer_id = %s
                          and to_char((u.parsed ->> 'txn_date')::date, 'YYYY-MM') = %s""", (customer_id, ym))
        ship_n = s["n"] or None
        return (ship_n, s["q"] if ship_n else None, e["s"] if e["n"] else None, e["c"] if e["n"] else None)

    def collections(self) -> None:
        from lcomfine.app import erp_upload, notify, stats

        v, t = self.v, self.tag
        C1 = self.customer_id
        # 월 · 날짜는 DB 가 센 값(세션 시간대) — 확정 출하 S1 의 confirmed_at 월 = ERP 일자 월
        anchor = self.q1("select to_char(confirmed_at, 'YYYY-MM') as ym, confirmed_at::date as d from shipment_order where id = %s", (v["S1id"],))
        ym, d = anchor["ym"], anchor["d"]
        m1 = date(int(ym[:4]), int(ym[5:7]), 1)
        m_last = (m1 + timedelta(days=32)).replace(day=1) - timedelta(days=1)
        P = self.free_period(m1, m_last)                    # 기준 기간 P — 회차 1 · 2 (교체)
        Q = self.free_period(m1, max(m1, min(m1 + timedelta(days=14), P[1] - timedelta(days=1))))   # 기준 기간 Q — ERP 만 있는 거래처
        v.update(ym=ym, P=P, Q=Q)
        # 열 매핑 — 지금 저장된 것을 쓴다. 거래처코드 · 일자 · 매출액 · 수금액이 없으면 시드 매핑(정본 낱말 7)을 HTTP 로 저장하고 끝에 되돌린다
        cmap = erp_upload.column_map()
        if not all(cmap.get(f) for f in ("customer_code", "txn_date", "sales_amount", "collected_amount")):
            from lcomfine.db import seed_dev3

            self.cmap_before = self.q("select field, column_name, created_by from erp_column_map order by id")
            self.post("COL-010 열 매핑 저장(시드 낱말 7)", "ac", "/col/uploads",
                      {"action": "map", **{f"map_{f}": c for f, c in seed_dev3.COLUMN_MAP.items()}})
            cmap = erp_upload.column_map()
        fields = [f for f in erp_upload.FIELDS if cmap.get(f)]
        heads = [cmap[f] for f in fields]

        def row(code, day, sales=None, collected=None, method=None, bill_due=None, balance=None):
            vals = {"customer_code": code, "txn_date": day, "sales_amount": sales, "collected_amount": collected,
                    "balance_amount": balance, "method": method, "bill_due_date": bill_due}
            return [vals[f] for f in fields]

        E, U, E2 = f"{t}-E", f"{t}-U", f"{t}-E2"           # E → 거래처 C1(출하 있음) · U 미매핑 · E2 → 거래처 C2(출하 없음)
        self.erp_codes += [E, U, E2]
        d1, d2 = m1.isoformat(), d.isoformat()
        link = erp_upload.error_link

        def notes_for(bid: int) -> list[dict]:
            return self.q("select to_role, handled_at, event_key, dedupe_key, message from notification where link = %s order by to_role",
                          (link(bid),))

        # ── 회차 1 (기간 P) — E 는 아직 미매핑이라 전부 오류 행 · 알림 1행 (AC)
        A = self.post_file("COL-010 ERP 업로드 1회차(기간 P · 미매핑 E · U · 날짜 오류)", "ac", P, heads,
                           [row(E, d1, "1000", "0"), row(E, d2, "0", "300"), row(U, d1, "77", "77"), row(E, f"{m1.year}/13/01", "1", "1")])
        bA = self.q1("select * from erp_upload_batch where id = %s", (A["id"],))
        nA = notes_for(A["id"])
        all_err = self.q1("select count(*) as n from notification where event_key = 'erp_upload_error' and link = %s", (link(A["id"]),))["n"]
        self.expect(bA["status"] == "반영" and bA["row_count"] == 4 and bA["error_count"] == 4 and A["round_no"] == 1 and A["replaced"] is None,
                    f"{bA['batch_no']} 반영 · 1회차 · 행 4 · 오류 4(E 미매핑 3 + 날짜 1 · U 1 — 겹침)")
        ok_note = ([n["to_role"] for n in nA] == ["AC"] and nA[0]["handled_at"] is None and nA[0]["event_key"] == "erp_upload_error"
                   and nA[0]["dedupe_key"] == f"erp_upload_error:{bA['batch_no']}" and all_err == 1
                   and "4행" in nA[0]["message"] and U in nA[0]["message"])
        with self.conn.tx() as cur:
            again = notify.emit(cur, "erp_upload_error", link=link(A["id"]), dedupe_key=f"erp_upload_error:{bA['batch_no']}", text="QA2 dedupe")
        n_after = len(notes_for(A["id"]))
        self.db(f"notification erp_upload_error → {[n['to_role'] for n in nA]} · link {link(A['id'])} · dedupe 재호출 새 행 {again}")
        # 매핑 전 COL-070 — C1 은 출하만 (MES 쪽만)
        before_map = self.col070(ym, C1)
        s1q = self.q1("select sum(l.qty) as q from shipment_line l where l.shipment_order_id = %s", (v["S1id"],))["q"]
        self.r.add("G-14", "출하만 있는 거래처 = MES 쪽만 (ERP 코드 매핑 전)",
                   before_map.get(C1) == (1, s1q, None, None),
                   f"COL-070 {ym} {t}-C → 출하 {before_map.get(C1)} (기대 (1, {s1q}, None, None)) — E 미매핑이라 ERP 쪽 비어 있음")

        # ── 회차 2 (같은 기간 P) — 교체 · 이전 회차 행 보존 · 교체된 회차의 알림 닫힘
        rows2 = [row(E, d1, "1000", "0"), row(E, d2, "0", "300", "PM (예시)", d2),      # 2 · 3 정상(매핑 뒤)
                 row(U, d1, "77", "77"),                                                  # 4 미매핑
                 row(E, f"{m1.year}/13/01", "1", "1"),                                    # 5 날짜 형식
                 row(E, d2, "abc", "0"),                                                  # 6 금액 형식
                 row("", d2, "5", "0")]                                                   # 7 필수(거래처코드) 누락
        A2 = self.post_file("COL-010 ERP 재업로드 2회차(같은 기간 P · 6행)", "ac", P, heads, rows2)
        bA_after = self.q1("select * from erp_upload_batch where id = %s", (A["id"],))
        bA2 = self.q1("select * from erp_upload_batch where id = %s", (A2["id"],))
        rowsA = self.q("select row_no, error from erp_upload_row where batch_id = %s order by row_no", (A["id"],))
        live_n = self.q1("select count(*) as n from erp_upload_batch where period_from = %s and period_to = %s and status = '반영'", P)["n"]
        viewA = self.q1("select count(*) as n from v_dq_collections where batch_no = %s", (bA["batch_no"],))["n"]
        nA_after = notes_for(A["id"])
        nA2 = notes_for(A2["id"])
        self.db(f"{bA['batch_no']} {bA_after['status']} replaced_by {bA_after['replaced_by']} · 행 {len(rowsA)} · {bA2['batch_no']} {bA2['status']} {bA2['round_no']}회차")
        self.r.add("G-12", "같은 기간 재업로드 = 새 회차 반영 · 이전 회차 교체 + replaced_by · 이전 회차 행 보존",
                   A2["replaced"] == bA["batch_no"] and A2["round_no"] == 2 and bA_after["status"] == "교체"
                   and bA_after["replaced_by"] == A2["id"] and bA2["status"] == "반영" and bA2["replaced_by"] is None
                   and len(rowsA) == 4 and bA_after["row_count"] == 4 and all(r["error"] for r in rowsA),
                   f"기간 {P[0]}~{P[1]}: {bA['batch_no']} 1회차 → {bA2['batch_no']} 2회차(응답 replaced {A2['replaced']}) · 이전 status {bA_after['status']} "
                   f"replaced_by {bA_after['replaced_by']} (= 새 회차 id {A2['id']}) · 이전 회차 행 4→{len(rowsA)} 그대로 · 새 회차 {bA2['status']} · 행 {A2['rows']} 오류 {A2['errors']}")
        self.r.add("G-12", "기간마다 반영 하나 · v_dq_collections 는 반영 회차만",
                   live_n == 1 and viewA == 0,
                   f"기간 {P[0]}~{P[1]} 의 반영 회차 {live_n} · 교체 회차 {bA['batch_no']} 의 뷰 행 {viewA} (이전엔 오류 4 라 0 · 교체 뒤에도 0)")
        self.r.add("G-11", "erp_upload_error",
                   ok_note and again == 0 and n_after == 1 and nA_after and nA_after[0]["handled_at"] is not None
                   and [n["to_role"] for n in nA2] == ["AC"] and nA2[0]["handled_at"] is None,
                   f"회차 {bA['batch_no']}(오류 4) → AC 1행 {'OK' if ok_note else '불일치'} · dedupe erp_upload_error:{bA['batch_no']} 재발생 새 행 {again} · "
                   f"같은 기간 교체 → 그 알림 닫힘 {'OK' if nA_after and nA_after[0]['handled_at'] else '안 닫힘'} · 새 회차 {bA2['batch_no']}(오류 {A2['errors']}) → AC 1행 미처리 "
                   f"{[n['to_role'] for n in nA2]} · 재해석으로 오류 0 → 닫힘은 다음 행(기간 Q 회차)")
        # COL-020 — 반영 / 교체됨 구분 · 회차별 오류 행
        html = self.get("COL-020 업로드 이력", "ac", "/col/upload-history").text
        segA = html[html.find(f'data-batch="{A["id"]}"'):][:240] if f'data-batch="{A["id"]}"' in html else ""
        segA2 = html[html.find(f'data-batch="{A2["id"]}"'):][:240] if f'data-batch="{A2["id"]}"' in html else ""
        detail = self.get("COL-020 회차별 오류 행 (2회차)", "ac", f"/col/upload-history?batch={A2['id']}").text
        err_rows_shown = sorted(int(x) for x in re.findall(r'<tr data-row="(\d+)"', detail))
        self.r.add("G-12", "COL-020 이력 — 반영 / 교체됨 구분 · 교체한 회차 번호",
                   'data-status="교체"' in segA and "교체됨" in segA and 'data-status="반영"' in segA2 and bA2["batch_no"] in detail,
                   f"{bA['batch_no']} 행 data-status {'교체 · 「교체됨」' if '교체됨' in segA else segA[:60]} · {bA2['batch_no']} 행 {'반영' if '반영' in segA2 else segA2[:60]}")
        self.r.add("G-13", "COL-020 회차별 오류 행 보기 (매핑 전 — 6행 전부 오류)", err_rows_shown == [2, 3, 4, 5, 6, 7],
                   f"?batch={A2['id']} 오류 행 data-row {err_rows_shown} (엑셀 행 2~7 · E 미매핑 포함)")

        # ── 회차 B (기간 Q) — ERP 만 있는 거래처 C2 · 미매핑 E2 → 매핑 등록으로 오류 0 → 알림 닫힘 (G-11 · G-13)
        C2 = self.post("BAS-020 거래처(ERP 만 · 출하 없음)", "sl", "/bas/customers",
                       {"customer_code": f"{t}-C2", "customer_name": f"{t} 거래처 2 (예시)"}, numbers=lambda b: f"customer {b['id']}")["id"]
        self.customer2_id = C2
        B = self.post_file("COL-010 ERP 업로드(기간 Q · 미매핑 E2 2행)", "ac", Q, heads, [row(E2, d1, "500", "0"), row(E2, d2, "0", "200")])
        bB = self.q1("select * from erp_upload_batch where id = %s", (B["id"],))
        nB = notes_for(B["id"])
        html = self.get("COL-030 거래처 코드 매핑(미매핑 목록)", "ac", "/col/customer-map").text
        listed_before = {c: f'data-unmapped="{c}"' in html for c in (E, U, E2)}
        unm = {u["erp_code"]: u for u in erp_upload.unmapped_codes() if u["erp_code"] in (E, U, E2)}
        self.db(f"COL-030 미매핑 목록 {listed_before} · unmapped_codes {sorted(unm)}")
        m2 = self.post("COL-030 거래처 코드 매핑 등록(E2 → C2)", "ac", "/col/customer-map",
                       {"erp_code": E2, "erp_name": f"{t} ERP 거래처 2 (예시)", "customer_id": str(C2)},
                       numbers=lambda b: f"map {b['id']} · 재해석 {b['reinterpreted']} · 반영 {b['applied']} · 여전히 오류 {b['still_error']}")
        bB_after = self.q1("select * from erp_upload_batch where id = %s", (B["id"],))
        rowsB = self.q("select row_no, error, updated_at from erp_upload_row where batch_id = %s order by row_no", (B["id"],))
        nB_after = notes_for(B["id"])
        viewB = self.q1("select count(*) as n from v_dq_collections where batch_no = %s and customer_code = %s", (bB["batch_no"], f"{t}-C2"))["n"]
        self.r.add("G-11", "erp_upload_error — 매핑 뒤 재해석으로 오류 0 이면 닫힘",
                   bB["error_count"] == 2 and [n["to_role"] for n in nB] == ["AC"] and nB[0]["handled_at"] is None
                   and (m2["reinterpreted"], m2["applied"], m2["still_error"], m2["remapped"]) == (2, 2, 0, 2)
                   and bB_after["error_count"] == 0 and nB_after and nB_after[0]["handled_at"] is not None,
                   f"회차 {bB['batch_no']} 오류 {bB['error_count']}(E2 미매핑) → AC 1행 미처리 → E2 매핑 등록 → 재해석 {m2['reinterpreted']} · 반영 {m2['applied']} · "
                   f"여전히 오류 {m2['still_error']} (remapped = applied {m2['remapped']}) · "
                   f"error_count {bB_after['error_count']} · 알림 handled_at {'있음(닫힘)' if nB_after and nB_after[0]['handled_at'] else '없음'}")
        # ── E → C1 매핑 — 반영 회차 A2 의 E 행만 재해석 · 교체 회차 A 는 그대로 · 오류 남은 알림은 열린 채
        m1_ = self.post("COL-030 거래처 코드 매핑 등록(E → C1)", "ac", "/col/customer-map",
                        {"erp_code": E, "erp_name": f"{t} ERP 거래처 (예시)", "customer_id": str(C1)},
                        numbers=lambda b: f"map {b['id']} · 재해석 {b['reinterpreted']} · 반영 {b['applied']} · 여전히 오류 {b['still_error']} · 「{b.get('message', '')[-60:]}」")
        rowsA2 = self.q("select row_no, raw, parsed, error, updated_at from erp_upload_row where batch_id = %s order by row_no", (A2["id"],))
        rowsA_after = self.q("select row_no, error from erp_upload_row where batch_id = %s order by row_no", (A["id"],))
        bA2_after = self.q1("select * from erp_upload_batch where id = %s", (A2["id"],))
        bA_final = self.q1("select * from erp_upload_batch where id = %s", (A["id"],))
        nA2_after = notes_for(A2["id"])
        err = {r["row_no"]: (r["error"] or "") for r in rowsA2}
        UP = erp_upload.UNMAPPED_PREFIX
        kept = (rowsA2[2]["raw"].get(cmap["customer_code"]) == U and rowsA2[3]["raw"].get(cmap["txn_date"]) == f"{m1.year}/13/01"
                and rowsA2[4]["raw"].get(cmap["sales_amount"]) == "abc")
        ok_err = (err[2] == "" and err[3] == "" and err[4].startswith(UP) and U in err[4]
                  and "일자" in err[5] and UP not in err[5] and "매출액" in err[6] and UP not in err[6]
                  and "거래처코드" in err[7] and "필수" in err[7])
        self.db(f"A2 error 행 {dict((k, x[:40]) for k, x in err.items())} · 재해석 {m1_['reinterpreted']} · 반영 {m1_['applied']} · 여전히 오류 {m1_['still_error']}")
        self.r.add("G-13", "오류 행 보존 — 날짜 · 금액 · 필수 · 미매핑 행이 error 사유 · 원문(raw)과 함께 저장",
                   kept and ok_err and bA2_after["error_count"] == 4 and bA2_after["row_count"] == 6,
                   f"{bA2['batch_no']} 행 6: 2 · 3 정상(error null) · 4 「{err[4][:30]}」 · 5 「{err[5][:24]}…」 · 6 「{err[6][:26]}…」 · 7 「{err[7][:30]}」 "
                   f"· raw 원문 그대로 {'OK' if kept else '불일치'} · error_count {bA2_after['error_count']}")
        self.r.add("G-13", "COL-030 미매핑 목록 · 매핑 뒤 반영 회차만 재해석(error is null) · 교체 회차 그대로",
                   listed_before == {E: True, U: True, E2: True} and set(unm) == {E, U, E2} and unm[E]["rows"] == 4 and unm[U]["rows"] == 1
                   and (m1_["reinterpreted"], m1_["applied"], m1_["still_error"], m1_["remapped"]) == (4, 2, 2, 2)
                   and "다시 해석한 행 4" in m1_.get("message", "") and "그중 반영 2" in m1_.get("message", "") and "여전히 제외 2" in m1_.get("message", "")
                   and all(r["error"] and UP in r["error"] for r in rowsA_after) and bA_final["error_count"] == 4
                   and nA2_after and nA2_after[0]["handled_at"] is None
                   and f'data-unmapped="{E}"' not in self.c("ac").get("/col/customer-map").text
                   and f'data-unmapped="{U}"' in self.c("ac").get("/col/customer-map").text,
                   f"매핑 전 COL-030 data-unmapped {listed_before} · unmapped_codes E {unm.get(E, {}).get('rows')}행(반영 회차만 · 교체 회차의 3행은 세지 않음) U {unm.get(U, {}).get('rows')}행 · "
                   f"E 매핑 → 반영 회차 {bA2['batch_no']} 응답 재해석 {m1_['reinterpreted']} · 반영 {m1_['applied']} · 여전히 오류 {m1_['still_error']} "
                   f"(remapped = applied {m1_['remapped']} · 문장 「{m1_.get('message', '')[-48:]}」 — DEF-QA2-004 해소) · 교체 회차 {bA['batch_no']} 오류 4 그대로 · "
                   f"남은 오류 4 라 알림 열린 채 · 매핑 뒤 목록에서 E 사라짐 · U 남음")
        # 반영에서 빠진다 — 뷰 · COL-040 · COL-050 · COL-020 오류 행
        viewA2 = self.q("select erp_customer_code, sales_amount, collected_amount from v_dq_collections where batch_no = %s order by txn_date",
                        (bA2["batch_no"],))
        viewU = self.q1("select count(*) as n from v_dq_collections where erp_customer_code = %s", (U,))["n"]
        c040 = self.get("COL-040 수금 현황(거래처 C1 · 이 달)", "ac",
                        f"/col/collections?date_from={m1.isoformat()}&date_to={m_last.isoformat()}&customer={C1}").text
        c040_rows = re.findall(r'<tr data-erp-code="([^"]+)"', c040)
        c050 = self.get("COL-050 거래처별 미수", "ac", f"/col/receivables?asof={m_last.isoformat()}").text
        c050_cells = self.row_cells(c050, f'data-customer="{C1}"')
        c050_nums = tuple(self.cell(x) for x in c050_cells[3:6]) if c050_cells and len(c050_cells) >= 6 else None
        detail2 = self.get("COL-020 회차별 오류 행 (E 매핑 뒤)", "ac", f"/col/upload-history?batch={A2['id']}").text
        err_rows_after = sorted(int(x) for x in re.findall(r'<tr data-row="(\d+)"', detail2))
        self.db(f"뷰 {len(viewA2)}행 · U {viewU} · COL-040 행 {c040_rows} · COL-050 {c050_nums} · COL-020 오류 행 {err_rows_after}")
        self.r.add("G-13", "오류 행 · 미매핑 행은 반영에서 제외 (v_dq_collections · COL-040 · COL-050) · COL-020 에는 보임",
                   len(viewA2) == 2 and {r["erp_customer_code"] for r in viewA2} == {E} and viewU == 0
                   and c040_rows == [E, E] and U not in c040 and c050_nums == (Decimal(1000), Decimal(300), Decimal(700)) and U not in c050
                   and err_rows_after == [4, 5, 6, 7],
                   f"뷰 {bA2['batch_no']} 행 {len(viewA2)}(E 만) · 미매핑 {U} 뷰 {viewU} · COL-040 행 {len(c040_rows)} {U} {'없음' if U not in c040 else '있음'} · "
                   f"COL-050 {t}-C 매출 · 수금 · 미수 {c050_nums} (기대 1000 · 300 · 700) · COL-020 ?batch= 오류 행 {err_rows_after} (기대 [4, 5, 6, 7])")

        # ── G-14 — COL-070 화면 = 독립 SQL (C1 양쪽 · C2 ERP 만 · U 제외)
        screen_all = self.col070(ym)
        screen_c1 = self.col070(ym, C1)
        ind1, ind2 = self.indep_ship_vs_coll(ym, C1), self.indep_ship_vs_coll(ym, C2)
        mod = {r["customer_id"]: (r["shipments"], r["shipped_qty"], r["sales_amount"], r["collected_amount"])
               for r in stats.shipment_vs_collection(ym, ym)["rows"] if r["ym"] == ym}
        view_c1 = self.q1("""select sum(sales_amount) as s, sum(collected_amount) as c from v_dq_collections
                              where customer_code = %s and to_char(txn_date, 'YYYY-MM') = %s""", (f"{t}-C", ym))
        v["col070"] = {"ym": ym, "C1": C1, "C2": C2, "screen": {C1: screen_all.get(C1), C2: screen_all.get(C2)}, "indep": {C1: ind1, C2: ind2},
                       "E": E, "U": U, "E2": E2}
        self.db(f"화면 C1 {screen_all.get(C1)} C2 {screen_all.get(C2)} · 독립 C1 {ind1} C2 {ind2} · 모듈 C1 {mod.get(C1)} C2 {mod.get(C2)}")
        self.r.add("G-14", "COL-070 화면 (월 · 거래처) = QA2 독립 SQL — 출하 · ERP 양쪽 있는 거래처",
                   screen_all.get(C1) == ind1 == (1, s1q, Decimal(1000), Decimal(300)) and screen_c1.get(C1) == ind1,
                   f"{ym} {t}-C: 화면 (출하 건수 · 수량 · 매출 · 수금) {screen_all.get(C1)} = 독립 SQL {ind1} (기대 (1, {s1q}, 1000, 300)) · ?customer= 필터 {screen_c1.get(C1)}")
        self.r.add("G-14", "ERP 만 있는 거래처 = ERP 쪽만 · 미매핑 코드는 어느 쪽에도 없음",
                   screen_all.get(C2) == ind2 == (None, None, Decimal(500), Decimal(200)) and viewU == 0 and U not in self.c("ac").get(
                       f"/col/shipment-vs-collection?ym_from={ym}&ym_to={ym}").text,
                   f"{ym} {t}-C2: 화면 {screen_all.get(C2)} = 독립 SQL {ind2} (기대 (None, None, 500, 200)) · {U}(미매핑 77/77) 뷰 {viewU} · 화면에 없음")
        self.r.add("G-14", "stats.shipment_vs_collection = 독립 SQL · 뷰(v_dq_collections) = 테이블 직접 집계",
                   mod.get(C1) == ind1 and mod.get(C2) == ind2 and (view_c1["s"], view_c1["c"]) == ind1[2:],
                   f"모듈 C1 {mod.get(C1)} C2 {mod.get(C2)} · 뷰 C1 매출 · 수금 ({view_c1['s']}, {view_c1['c']}) = 테이블 직접 {ind1[2:]} · 건 단위 짝짓기 없음(월 · 거래처 합만)")

    # ════════════════════════════════════════════════════════════════════════
    # W3 대시보드 — G-15 (dashboard 8 함수 = QA2 독립 SQL = 화면 KPI · 표 · SVG <title>) · G-05 RPT 8 화면 diff 0
    # ════════════════════════════════════════════════════════════════════════
    # ── 화면 조각 읽기 (템플릿 표식 — progress-dev3 §1.5: data-kpi · data-series · data-series-table · data-table · data-section) ──
    @staticmethod
    def _strip(s: str) -> str:
        return " ".join(_html.unescape(re.sub("<[^>]+>", " ", s)).split())

    def kpi_text(self, html: str, key: str) -> str | None:
        """`data-kpi="<key>"` 카드의 값 글자 — 「2 건」 · 「80.000」 · 「미수집」 · 「미확정 (D-519)」. 카드가 없으면 None."""
        m = re.search(r'<div class="kpi" data-kpi="' + re.escape(key) + r'">(.*?)</div>', html, re.S)
        if not m:
            return None
        s = re.search(r'<strong class="kpi__value num">(.*?)</strong>', m.group(1), re.S)
        return self._strip(s.group(1)) if s else None

    def series_html(self, html: str, name: str) -> tuple[list[tuple[str, str]], list[tuple[str, str]]] | None:
        """`data-series="<name>"` 구역의 SVG `<title>` 「라벨: 값」 과 `table[data-series-table]` 행(라벨, 값). 구역이 없으면 None."""
        i = html.find(f'data-series="{name}"')
        if i < 0:
            return None
        seg = html[i:html.find("</section>", i)]
        titles = [(self._strip(a), self._strip(b)) for a, b in re.findall(r"<title>(.*): ([^<]*)</title>", seg)]
        rows = [(self._strip(a), self._strip(b)) for a, b in re.findall(r'<tr><td>(.*?)</td><td class="num">(.*?)</td></tr>', seg, re.S)]
        return titles, rows

    def table_html(self, html: str, name: str) -> list[list[str]] | None:
        """`data-table="<name>"` 구역의 tbody 행(td 글자). 구역이 없으면 None · 빈 표(미수집 문구)는 []."""
        i = html.find(f'data-table="{name}"')
        if i < 0:
            return None
        seg = html[i:html.find("</section>", i)]
        body = re.search(r"<tbody>(.*?)</tbody>", seg, re.S)
        if not body:
            return []
        return [[self._strip(td) for td in re.findall(r"<td[^>]*>(.*?)</td>", tr, re.S)]
                for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", body.group(1), re.S)]

    @staticmethod
    def kpi_expected(k: dict) -> str:
        """KPI 카드가 보여야 하는 글자 — 수 대신 text · None 은 미수집 · 그 밖 `num(digits)` + 단위 (템플릿 `rpt/_d.html` kpis 매크로의 규칙)."""
        from lcomfine.app.util.screen import num as fmt

        if k.get("text"):
            return k["text"]
        if k["value"] is None:
            return "미수집"
        return fmt(k["value"], k["digits"]) + (f" {k['unit']}" if k["unit"] else "")

    def kpi_mismatch(self, html: str, data: dict, keys=None) -> list[str]:
        """화면 KPI 글자 ≠ 함수 반환으로 기대되는 글자 인 키 목록(빈 리스트 = 전부 일치)."""
        bad = []
        for k in data["kpis"]:
            if keys is not None and k["key"] not in keys:
                continue
            got = self.kpi_text(html, k["key"])
            exp = self.kpi_expected(k)
            if got != exp:
                bad.append(f"{k['key']} 화면 {got!r} ≠ {exp!r}")
        return bad

    def series_mismatch(self, html: str, data: dict) -> list[str]:
        """series 마다 SVG <title> 「라벨: 값」 · 표 둘째 열 = 함수 반환(`num` 소수 0 · None = 미수집) — 어긋난 것 목록."""
        from lcomfine.app.util.screen import num as fmt

        bad = []
        for name, pts in data["series"].items():
            got = self.series_html(html, name)
            if got is None:
                bad.append(f"{name}: 화면에 data-series 없음")
                continue
            titles, rows = got
            exp = [(p["label"], "미수집" if p["value"] is None else fmt(p["value"])) for p in pts]
            all_none = not pts or all(p["value"] is None for p in pts)
            if rows != exp:
                bad.append(f"{name}: 표 {rows[:4]}… ≠ {exp[:4]}…")
            if all_none:
                if titles:
                    bad.append(f"{name}: 전부 미수집인데 SVG 가 있다 {titles[:3]}")
            elif titles != exp:
                bad.append(f"{name}: SVG title {titles[:4]}… ≠ {exp[:4]}…")
        return bad

    @staticmethod
    def _eq(a, b) -> bool:
        """수 · None 비교 — Decimal 80.000 == 80 · None == None."""
        if a is None or b is None:
            return a is None and b is None
        return Decimal(str(a)) == Decimal(str(b))

    def _kv(self, data: dict) -> dict:
        return {k["key"]: k["value"] for k in data["kpis"]}

    def _kt(self, data: dict) -> dict:
        return {k["key"]: k["text"] for k in data["kpis"]}

    @staticmethod
    def _nz(x):
        """QA2 쪽 규칙(D-805 「KPI 도 0건이면 None」) — 0 · None → None."""
        return None if x is None or x == 0 else x

    @staticmethod
    def _rate100(a, denom) -> Decimal | None:
        return (Decimal(a) * 100 / Decimal(denom)).quantize(Decimal("0.1")) if denom else None

    @staticmethod
    def _days(f: date, t: date) -> list[str]:
        out, d = [], f
        while d <= t:
            out.append(d.isoformat())
            d += timedelta(days=1)
        return out

    def _series_from(self, labels: list[str], rows: list[dict], key: str, val: str) -> list:
        got = {str(r[key]): r[val] for r in rows}
        return [got.get(lb) for lb in labels]

    @staticmethod
    def _snapshot_cm():
        """한 스냅샷(REPEATABLE READ 트랜잭션 하나) 안에서 `conn.q` · `conn.q1` 을 읽는다 — 그 동안 dashboard · stats · 라우터(TestClient GET) ·
        QA2 독립 SQL 이 **같은 순간의 DB** 를 본다(다른 QA 가 같은 DB 에 동시에 쓴다 · 쓰기 `conn.x` · `conn.tx` 는 건드리지 않는다 · 끝에 rollback)."""
        import contextlib

        import psycopg
        from psycopg.rows import dict_row

        from lcomfine.db import conn as _conn

        @contextlib.contextmanager
        def cm():
            c = psycopg.connect(_conn.dsn(), row_factory=dict_row)
            c.isolation_level = psycopg.IsolationLevel.REPEATABLE_READ
            lock = threading.Lock()
            orig_q, orig_q1 = _conn.q, _conn.q1

            def q(sql, params=None):
                with lock, c.cursor() as cur:
                    cur.execute(sql, params)
                    return [dict(r) for r in cur.fetchall()] if cur.description is not None else []

            def q1(sql, params=None):
                rows = q(sql, params)
                return rows[0] if rows else None

            _conn.q, _conn.q1 = q, q1
            try:
                yield c
            finally:
                _conn.q, _conn.q1 = orig_q, orig_q1
                c.rollback()
                c.close()
        return cm()

    def snapshot(self):
        return self._snapshot_cm()

    def stable(self, indep, dash, html_fn=None, tries: int = 3):
        """같은 DB 를 다른 QA 가 함께 쓴다 — 독립 SQL → 함수 → 화면 → 독립 SQL 을 한 번에 읽고, 앞뒤 독립값이 같을 때만 그 묶음을 쓴다(최대 tries 회)."""
        last = None
        for n in range(1, tries + 1):
            a = indep()
            d = dash()
            h = html_fn() if html_fn else None
            b = indep()
            last = (a, d, h, n)
            if a == b:
                return a, d, h, n
        return last

    # ── 독립 SQL (D-805 의 정의를 QA2 가 따로 쓴 문장 — 집계는 SQL group by · filter · 뷰를 거치지 않는다) ──
    def indep_orders(self, f: date, t: date, cid: int | None) -> dict:
        p = {"f": f, "t": t, "c": cid}
        where = "r.request_date between %(f)s and %(t)s and r.status <> '취소' and (%(c)s::bigint is null or r.customer_id = %(c)s::bigint)"
        k = self.q1(f"""select count(*) as n, sum(r.qty) as qty, count(distinct r.customer_id) as customers,
                               count(*) filter (where not exists (select 1 from work_order w where w.request_id = r.id and w.status <> '취소')) as not_started
                          from production_request r where {where}""", p)
        days = self.q(f"select r.request_date::text as d, count(*) as n, sum(r.qty) as qty from production_request r where {where} group by 1", p)
        due = self.q(f"select to_char(r.due_date, 'YYYY-MM') as m, count(*) as n from production_request r where {where} group by 1 order by 1", p)
        top = self.q(f"""select c.customer_code, count(*) as n, sum(r.qty) as qty from production_request r join customer c on c.id = r.customer_id
                          where {where} group by 1 order by 2 desc, 3 desc, 1 limit 10""", p)
        labels = self._days(f, t)
        return {"requests": self._nz(k["n"]), "qty": k["qty"] if k["n"] else None, "customers": self._nz(k["customers"]),
                "not_started": self._nz(k["not_started"]),
                "s_requests": self._series_from(labels, days, "d", "n"), "s_qty": self._series_from(labels, days, "d", "qty"),
                "s_due": [(d["m"], d["n"]) for d in due], "top": [(r["customer_code"], r["n"], r["qty"]) for r in top]}

    def indep_production(self, f: date, t: date) -> dict:
        p = {"f": f, "t": t}
        k = self.q1("""select count(distinct work_order_id) as wos, sum(produced_qty) as produced, sum(good_qty) as good, sum(loss_qty) as loss,
                              sum(run_minutes) as run, sum(stop_minutes) as stop, count(*) as n
                         from production_daily where work_date between %(f)s and %(t)s""", p)
        oq = self.q1("""select sum(w.qty) as q from work_order w
                         where w.id in (select distinct work_order_id from production_daily where work_date between %(f)s and %(t)s)""", p)
        days = self.q("select work_date::text as d, sum(good_qty) as good, sum(produced_qty) as produced from production_daily where work_date between %(f)s and %(t)s group by 1", p)
        eq = self.q("select equipment_code, sum(good_qty) as good, sum(produced_qty) as produced from production_daily where work_date between %(f)s and %(t)s group by 1", p)
        run, stop = k["run"], k["stop"]
        total = (run or 0) + (stop or 0)
        util = None if (run is None and stop is None) or total == 0 else (Decimal(run or 0) * 100 / Decimal(total)).quantize(Decimal("0.1"))
        labels = self._days(f, t)
        return {"work_orders": self._nz(k["wos"]), "order_qty": oq["q"] if k["n"] else None, "produced_qty": self._nz(k["produced"]),
                "good_qty": self._nz(k["good"]), "loss_qty": k["loss"] if k["n"] else None, "utilization": util,
                "s_good": self._series_from(labels, days, "d", "good"), "s_produced": self._series_from(labels, days, "d", "produced"),
                "by_eq": {r["equipment_code"]: (r["good"], r["produced"]) for r in eq}}

    def indep_quality(self, f: date, t: date) -> dict:
        p = {"f": f, "t": t}
        k = self.q1("select sum(inspected_qty) as insp, sum(defect_qty) as dft, count(*) as n from inspection_daily where work_date between %(f)s and %(t)s", p)
        days = self.q("select work_date::text as d, sum(inspected_qty) as insp, sum(defect_qty) as dft from inspection_daily where work_date between %(f)s and %(t)s group by 1", p)
        dtypes = self.q("""select e.key as code, sum(nullif(e.value, '')::numeric) as qty from inspection_daily d, jsonb_each_text(d.defects) e
                            where d.work_date between %(f)s and %(t)s group by 1""", p)
        j = self.q1("""select count(*) as n, count(*) filter (where seq > 1) as rejudged,
                              count(distinct (kind, lot_id)) filter (where result = '불합격') as rejected
                         from (select '입고' as kind, material_lot_id as lot_id, seq, result, inspected_at from incoming_inspection
                               union all
                               select '제품', product_lot_id, seq, result, inspected_at from product_inspection) x
                        where x.inspected_at::date between %(f)s and %(t)s""", p)
        labels = self._days(f, t)
        return {"inspected_qty": self._nz(k["insp"]), "defect_qty": k["dft"] if k["n"] else None, "judgements": self._nz(j["n"]),
                "rejected_lots": self._nz(j["rejected"]), "rejudged": self._nz(j["rejudged"]),
                "s_defect": self._series_from(labels, days, "d", "dft"), "s_inspected": self._series_from(labels, days, "d", "insp"),
                "defects": {r["code"]: r["qty"] for r in dtypes}}

    def indep_inventory(self, f: date, t: date) -> dict:
        p = {"f": f, "t": t}
        rec = self.q1("""select count(*) as n, sum(s) as qty from (
                            select r.id, sum(l.qty) as s from material_receipt r join material_lot l on l.receipt_id = r.id
                             where r.canceled_at is null and l.canceled_at is null and l.parent_lot_id is null
                               and r.received_on between %(f)s and %(t)s group by r.id) x""", p)
        rec_days = self.q("""select r.received_on::text as d, sum(l.qty) as qty from material_receipt r join material_lot l on l.receipt_id = r.id
                              where r.canceled_at is null and l.canceled_at is null and l.parent_lot_id is null
                                and r.received_on between %(f)s and %(t)s group by 1""", p)
        iss = self.q1("""select count(*) as n, sum(s) as qty from (
                            select m.id, sum(l.qty) as s from material_issue m join material_issue_line l on l.issue_id = m.id
                             where m.canceled_at is null and m.issued_on between %(f)s and %(t)s group by m.id) x""", p)
        iss_days = self.q("""select m.issued_on::text as d, sum(l.qty) as qty from material_issue m join material_issue_line l on l.issue_id = m.id
                              where m.canceled_at is null and m.issued_on between %(f)s and %(t)s group by 1""", p)
        prod = self.q1("select count(*) as n, sum(qty) as qty from product_receipt where canceled_at is null and received_on between %(f)s and %(t)s", p)
        adj = self.q1("select count(*) as n, sum(diff_qty) as diff from stock_adjust where adjusted_at::date between %(f)s and %(t)s", p)
        top = self.q("""select sum(qty) as q from stock_move where moved_at::date between %(f)s and %(t)s group by lot_kind, lot_id
                         order by 1 desc limit 10""", p)
        stay = self.q("""select days from (
                            select current_date - m.created_at::date as days from lot_stock s join material_lot m on s.lot_kind = '원자재' and m.id = s.lot_id where s.qty > 0
                            union all
                            select current_date - pl.created_at::date from lot_stock s join product_lot pl on s.lot_kind = '제품' and pl.id = s.lot_id where s.qty > 0) x
                          order by days desc limit 10""")
        labels = self._days(f, t)
        return {"material_receipts": self._nz(rec["n"]), "material_receipt_qty": self._nz(rec["qty"]), "issues": self._nz(iss["n"]),
                "issue_qty": self._nz(iss["qty"]), "product_receipts": self._nz(prod["n"]), "product_receipt_qty": self._nz(prod["qty"]),
                "adjusts": self._nz(adj["n"]), "adjust_diff": adj["diff"] if adj["n"] else None,
                "s_receipt": self._series_from(labels, rec_days, "d", "qty"), "s_issue": self._series_from(labels, iss_days, "d", "qty"),
                "top_moves": [r["q"] for r in top], "long_stay": [r["days"] for r in stay]}

    def indep_shipment(self, f: date, t: date, cid: int | None) -> dict:
        p = {"f": f, "t": t, "c": cid}
        base = """select o.id, o.customer_id, o.confirmed_at::date as d, sum(l.qty) as q
                    from shipment_order o join shipment_line l on l.shipment_order_id = o.id
                   where o.status = '확정' and o.confirmed_at::date between %(f)s and %(t)s
                     and (%(c)s::bigint is null or o.customer_id = %(c)s::bigint) group by o.id, o.customer_id, o.confirmed_at"""
        k = self.q1(f"select count(*) as n, sum(q) as qty, count(distinct customer_id) as customers from ({base}) x", p)
        days = self.q(f"select d::text as d, count(*) as n, sum(q) as qty from ({base}) x group by 1", p)
        cust = self.q(f"select c.customer_code, sum(x.q) as qty, count(*) as n from ({base}) x join customer c on c.id = x.customer_id group by 1 order by 2 desc, 1", p)
        due = self.q1("""select count(*) filter (where done and last_on <= due_date) as on_time,
                                count(*) filter (where done and last_on > due_date) as late_done,
                                count(*) filter (where not done and due_date < current_date) as late_open,
                                count(*) filter (where not done and due_date >= current_date) as in_progress
                           from (select q.due_date, (q.qty > 0 and coalesce(s.sq, 0) >= q.qty) as done, s.last_on
                                   from production_request q
                                   left join (select o.request_id, sum(l.qty) as sq, max(o.confirmed_at)::date as last_on
                                                from shipment_order o join shipment_line l on l.shipment_order_id = o.id
                                               where o.status = '확정' group by o.request_id) s on s.request_id = q.id
                                  where q.status <> '취소' and q.due_date between %(f)s and %(t)s
                                    and (%(c)s::bigint is null or q.customer_id = %(c)s::bigint)) x""", p)
        denom = due["on_time"] + due["late_done"] + due["late_open"]
        labels = self._days(f, t)
        return {"shipments": self._nz(k["n"]), "qty": k["qty"] if k["n"] else None, "customers": self._nz(k["customers"]),
                "due_rate": self._rate100(due["on_time"], denom), "on_time": self._nz(due["on_time"]), "late_done": self._nz(due["late_done"]),
                "late_open": self._nz(due["late_open"]),
                "s_qty": self._series_from(labels, days, "d", "qty"), "s_shipments": self._series_from(labels, days, "d", "n"),
                "by_customer": [(r["customer_code"], r["qty"], r["n"]) for r in cust]}

    def indep_collection(self, f: date, t: date, cid: int | None) -> dict:
        """RPT-060 — ERP 반영 회차 행을 테이블에서 직접(뷰 · stats 를 거치지 않음) · 월 라벨 = 기간의 달들 · 미수 연령은 기준일 − 일자."""
        months, y, m = [], f.year, f.month
        while date(y, m, 1) <= t:
            months.append(f"{y:04d}-{m:02d}")
            y, m = (y + 1, 1) if m == 12 else (y, m + 1)
        asof = self.q1("select max(period_to) as d from erp_upload_batch where status = '반영'")["d"]
        sales_mapped = bool(self.q1("select 1 from erp_column_map where field = 'sales_amount' and coalesce(column_name, '') <> ''"))
        per = {ym: self.indep_ship_vs_coll_any(ym, cid) for ym in months}
        coll = [per[ym][3] for ym in months]
        sales = [per[ym][2] for ym in months]
        shipped = [per[ym][1] for ym in months]
        some = lambda xs: sum((x for x in xs if x is not None), Decimal(0)) if any(x is not None for x in xs) else None  # noqa: E731
        buckets = ["0~30일", "31~60일", "61~90일", "90일 초과"]
        aging_series = [None] * 4
        over90 = None
        if asof is not None:
            rows = self.q("""select m.customer_id,
                                    case when %(asof)s::date - (u.parsed ->> 'txn_date')::date <= 30 then '0~30일'
                                         when %(asof)s::date - (u.parsed ->> 'txn_date')::date <= 60 then '31~60일'
                                         when %(asof)s::date - (u.parsed ->> 'txn_date')::date <= 90 then '61~90일' else '90일 초과' end as b,
                                    sum(coalesce(nullif(u.parsed ->> 'sales_amount', '')::numeric, 0)) - sum(coalesce(nullif(u.parsed ->> 'collected_amount', '')::numeric, 0)) as net
                               from erp_upload_row u join erp_upload_batch b on b.id = u.batch_id and b.status = '반영'
                               join erp_customer_map m on m.erp_code = u.parsed ->> 'customer_code' and m.use_yn = 'Y'
                              where u.error is null and (u.parsed ->> 'txn_date')::date <= %(asof)s
                                and (%(c)s::bigint is null or m.customer_id = %(c)s::bigint)
                              group by 1, 2""", {"asof": asof, "c": cid})
            if sales_mapped:
                if rows:
                    aging_series = [sum((r["net"] for r in rows if r["b"] == b), Decimal(0)) for b in buckets]
                over90 = self._nz(len({r["customer_id"] for r in rows if r["b"] == "90일 초과" and r["net"] > 0}))
        return {"asof": asof, "asof_text": asof.isoformat() if asof else "미수집", "collected": some(coll),
                "sales": some(sales) if sales_mapped else None, "sales_text": None if sales_mapped else "미확정 (D-516)",
                "shipped_qty": some(shipped), "over90": over90,
                "over90_text": "미수집" if asof is None else (None if sales_mapped else "미확정 (D-516)"),
                "months": months, "s_collected": coll, "s_shipped": shipped, "s_aging": aging_series, "buckets": buckets}

    def indep_ship_vs_coll_any(self, ym: str, cid: int | None) -> tuple:
        """indep_ship_vs_coll 의 거래처 선택형(None = 전체) — (확정 출하 건수, 출하 수량, ERP 매출 합, ERP 수금 합)."""
        p = {"c": cid, "ym": ym}
        s = self.q1("""select count(distinct o.id) as n, sum(l.qty) as q from shipment_order o join shipment_line l on l.shipment_order_id = o.id
                        where o.status = '확정' and (%(c)s::bigint is null or o.customer_id = %(c)s::bigint) and to_char(o.confirmed_at, 'YYYY-MM') = %(ym)s""", p)
        e = self.q1("""select sum(nullif(u.parsed ->> 'sales_amount', '')::numeric) as s, sum(nullif(u.parsed ->> 'collected_amount', '')::numeric) as c, count(*) as n
                         from erp_upload_row u join erp_upload_batch b on b.id = u.batch_id and b.status = '반영'
                         join erp_customer_map m on m.erp_code = u.parsed ->> 'customer_code' and m.use_yn = 'Y'
                        where u.error is null and (%(c)s::bigint is null or m.customer_id = %(c)s::bigint)
                          and to_char((u.parsed ->> 'txn_date')::date, 'YYYY-MM') = %(ym)s""", p)
        n = s["n"] or None
        return (n, s["q"] if n else None, e["s"] if e["n"] else None, e["c"] if e["n"] else None)

    STAGES = (("의뢰 → 지시", "request_date", "ordered_on"), ("지시 → 생산 완료", "ordered_on", "produced_on"),
              ("생산 완료 → 제품 입고", "produced_on", "received_on"), ("제품 입고 → 출하", "received_on", "shipped_on"),
              ("의뢰 → 출하", "request_date", "shipped_on"))     # D-805 RPT-070 의 단계 5 (QA2 가 결정문에서 옮겨 적음)

    def indep_order_flow(self, f: date, t: date, cid: int | None) -> dict:
        rows = self.q("""select r.request_no, r.request_date,
                                min(w.created_at)::date as ordered_on, max(d.work_date) as produced_on,
                                max(pr.received_on) as received_on, max(so.confirmed_at)::date as shipped_on
                           from production_request r
                           left join work_order w on w.request_id = r.id and w.status <> '취소'
                           left join production_daily d on d.work_order_id = w.id
                           left join product_receipt pr on pr.work_order_id = w.id and pr.canceled_at is null
                           left join shipment_order so on so.request_id = r.id and so.status = '확정'
                          where r.request_date between %(f)s and %(t)s and r.status <> '취소'
                            and (%(c)s::bigint is null or r.customer_id = %(c)s::bigint)
                          group by r.id, r.request_no, r.request_date order by r.request_date, r.id""", {"f": f, "t": t, "c": cid})
        per = {}
        vals = {name: [] for name, _, _ in self.STAGES}
        for r in rows:
            per[r["request_no"]] = {}
            for name, a, b in self.STAGES:
                d = (r[b] - r[a]).days if r[a] is not None and r[b] is not None else None
                per[r["request_no"]][name] = d
                if d is not None:
                    vals[name].append(d)
        avg = {name: ((Decimal(sum(v)) / Decimal(len(v))).quantize(Decimal("0.1")) if v else None) for name, v in vals.items()}
        return {"requests": self._nz(len(rows)), "shipped": self._nz(sum(1 for r in rows if r["shipped_on"] is not None)),
                "avg": avg, "count": {n: len(v) for n, v in vals.items()}, "max": {n: (max(v) if v else None) for n, v in vals.items()},
                "rows": per}

    # ── 단계 ──
    def dashboards(self) -> None:
        from lcomfine.app import dashboard

        v, t = self.v, self.tag
        C1, C2 = self.customer_id, self.customer2_id
        ym = v["ym"]
        m1 = date(int(ym[:4]), int(ym[5:7]), 1)
        m_last = (m1 + timedelta(days=32)).replace(day=1) - timedelta(days=1)
        P = (m1, m_last)
        qs = f"date_from={m1.isoformat()}&date_to={m_last.isoformat()}"
        who = "admin"       # 전 화면 R 이상 — 경영진(ex)도 RPT 8 전부 R 이다(permission-matrix)
        g15 = v.setdefault("g15", {"period": P, "C1": C1, "C2": C2})

        def page(step: str, path: str, login: str = who) -> str:
            return self.get(step, login, path).text

        # 8 화면 · 함수 · 독립 SQL · 화면(GET)을 **한 스냅샷** 에서 읽는다 — 다른 QA 의 동시 쓰기가 끼어들지 않는다(020 · 030 · 040 은 전 DB 집계)
        with self.snapshot():
            # ── RPT-010 수주 (거래처 C1 · 10월) ──
            a, d, h, n = self.stable(lambda: self.indep_orders(m1, m_last, C1), lambda: dashboard.orders(P, C1),
                                     lambda: page("RPT-010 수주 리포트(거래처 C1 · 이 달)", f"/rpt/orders?{qs}&customer={C1}"))
            kv = self._kv(d)
            kpi_ok = all(self._eq(kv[k], a[k]) for k in ("requests", "qty", "customers", "not_started"))
            ser_ok = ([p["value"] for p in d["series"]["requests"]] == a["s_requests"] and [p["value"] for p in d["series"]["qty"]] == a["s_qty"]
                      and [(p["label"], p["value"]) for p in d["series"]["due"]] == a["s_due"])
            top_ok = [(r["customer_code"], r["requests"], r["qty"]) for r in d["tables"]["top_customers"]] == a["top"]
            ns_ok = {r["request_no"] for r in d["tables"]["not_started"]} >= {v["R2"]} and v["R1"] not in {r["request_no"] for r in d["tables"]["not_started"]}
            hk, hs = self.kpi_mismatch(h, d), self.series_mismatch(h, d)
            ht = self.table_html(h, "not_started")
            ht_ok = ht is not None and len(ht) == len(d["tables"]["not_started"])
            g15["orders"] = {"kpis": {k: kv[k] for k in ("requests", "qty", "customers", "not_started")}, "indep": {k: a[k] for k in ("requests", "qty", "customers", "not_started")}}
            self.db(f"KPI 함수 {g15['orders']['kpis']} · 독립 {g15['orders']['indep']} · 미착수 {sorted(r['request_no'] for r in d['tables']['not_started'])} · 읽기 {n}회")
            self.r.add("G-15", "RPT-010 수주 — dashboard.orders = 독립 SQL = 화면 KPI · 표 · SVG title (거래처 C1 · 이 달)",
                       kpi_ok and ser_ok and top_ok and ns_ok and not hk and not hs and ht_ok and d["bucket"] == "day",
                       f"{ym} {t}-C: 의뢰 {kv['requests']} · 수량 {kv['qty']} · 거래처 {kv['customers']} · 미착수 {kv['not_started']} = 독립 SQL "
                       f"{a['requests']} · {a['qty']} · {a['customers']} · {a['not_started']} · 구간(일 {len(a['s_requests'])}) 건수 · 수량 · 납기 월 분포 {a['s_due']} 같음 {ser_ok} · "
                       f"거래처 상위 {top_ok} · 미착수에 {v['R2']} 있고 {v['R1']}(지시 있음) 없음 {ns_ok} · 화면 KPI 어긋남 {hk or 0} · SVG title/표 어긋남 {hs or 0} · 미착수 표 {len(ht or [])}행")

            # ── RPT-020 생산 (기간만 — 전 DB · 앞뒤 독립값이 같은 묶음만) ──
            a, d, h, n = self.stable(lambda: self.indep_production(m1, m_last), lambda: dashboard.production(P),
                                     lambda: page("RPT-020 생산 리포트(이 달)", f"/rpt/production?{qs}"))
            kv, kt = self._kv(d), self._kt(d)
            keys = ("work_orders", "order_qty", "produced_qty", "good_qty", "loss_qty", "utilization")
            kpi_ok = all(self._eq(kv[k], a[k]) for k in keys) and kv["loss_rate"] is None and kt["loss_rate"] == "미확정 (D-519)"
            ser_ok = ([p["value"] for p in d["series"]["good_qty"]] == a["s_good"] and [p["value"] for p in d["series"]["produced_qty"]] == a["s_produced"])
            eq_tab = {r["equipment_code"]: (r["good_qty"], r["produced_qty"]) for r in d["tables"]["by_equipment"] if r["good_qty"] is not None or r["produced_qty"] is not None}
            eq_ok = eq_tab == a["by_eq"] and {p["value"] for p in d["series"]["by_equipment"] if p["value"] is not None} == {g for g, _ in a["by_eq"].values()}
            mine = [r for r in d["tables"]["by_work_order"] if r["work_order_no"] == v["WO"]]
            mine_ok = len(mine) == 1 and mine[0]["produced_qty"] == 100 and mine[0]["good_qty"] == 95 and mine[0]["loss_qty"] == 5 and mine[0]["order_qty"] == 100
            hk, hs = self.kpi_mismatch(h, d), self.series_mismatch(h, d)
            g15["production"] = {"kpis": {k: kv[k] for k in keys}, "indep": {k: a[k] for k in keys}}
            self.db(f"KPI 함수 {g15['production']['kpis']} · 독립 {g15['production']['indep']} · 설비 {a['by_eq']} · 읽기 {n}회")
            self.r.add("G-15", "RPT-020 생산 — dashboard.production = 독립 SQL = 화면 · 로스율 미확정 (D-519) · 가동률 D-801 · 기계별 (이 달 · 거래처 필터 없음)",
                       kpi_ok and ser_ok and eq_ok and mine_ok and not hk and not hs and self.kpi_text(h, "loss_rate") == "미확정 (D-519)",
                       f"{ym}: 지시 {kv['work_orders']} · 지시 수량 {kv['order_qty']} · 생산 {kv['produced_qty']} · 양품 {kv['good_qty']} · 로스 {kv['loss_qty']} · 가동률 {kv['utilization']} "
                       f"= 독립 {[a[k] for k in keys]} · 로스율 text 「{kt['loss_rate']}」 · 화면 「{self.kpi_text(h, 'loss_rate')}」 · 구간 양품 · 생산 같음 {ser_ok} · "
                       f"기계별(설비 코드 → 양품 · 생산) {eq_ok} · {v['WO']} 행 100/95/5 {mine_ok} · 화면 KPI 어긋남 {hk or 0} · SVG/표 어긋남 {hs or 0} · 읽기 {n}회")

            # ── RPT-030 품질 ──
            a, d, h, n = self.stable(lambda: self.indep_quality(m1, m_last), lambda: dashboard.quality(P),
                                     lambda: page("RPT-030 품질 리포트(이 달)", f"/rpt/quality?{qs}"))
            kv, kt = self._kv(d), self._kt(d)
            keys = ("inspected_qty", "defect_qty", "judgements", "rejected_lots", "rejudged")
            kpi_ok = all(self._eq(kv[k], a[k]) for k in keys) and kt["loss_rate"] == "미확정 (D-519)"
            ser_ok = ([p["value"] for p in d["series"]["defect_qty"]] == a["s_defect"] and [p["value"] for p in d["series"]["inspected_qty"]] == a["s_inspected"])
            dt_ok = {r["defect_code"]: r["qty"] for r in d["tables"]["defects_by_type"]} == a["defects"]
            rj_ok = len(d["tables"]["rejudged"]) == (a["rejudged"] or 0) and all(j["seq"] > 1 for j in d["tables"]["rejudged"])
            rc_ok = len({(j["inspection_kind"], j["lot_no"]) for j in d["tables"]["rejected"]}) == (a["rejected_lots"] or 0)
            hk, hs = self.kpi_mismatch(h, d), self.series_mismatch(h, d)
            g15["quality"] = {"kpis": {k: kv[k] for k in keys}, "indep": {k: a[k] for k in keys}}
            self.db(f"KPI 함수 {g15['quality']['kpis']} · 독립 {g15['quality']['indep']} · 불량유형 {a['defects']} · 읽기 {n}회")
            self.r.add("G-15", "RPT-030 품질 — dashboard.quality = 독립 SQL = 화면 · 검사 로스율 미확정 (D-519) · 불합격 LOT · 재판정 (이 달)",
                       kpi_ok and ser_ok and dt_ok and rj_ok and rc_ok and not hk and not hs and self.kpi_text(h, "loss_rate") == "미확정 (D-519)",
                       f"{ym}: 검사 {kv['inspected_qty']} · 불량 {kv['defect_qty']} · 판정 {kv['judgements']} · 불합격 LOT {kv['rejected_lots']} · 재판정 {kv['rejudged']} "
                       f"= 독립 {[a[k] for k in keys]} · 로스율 「{kt['loss_rate']}」 화면 「{self.kpi_text(h, 'loss_rate')}」 · 구간 불량 · 검사 같음 {ser_ok} · "
                       f"불량유형 합 {a['defects']} {dt_ok} · 재판정 표 seq>1 {rj_ok} · 불합격 표 (종류, LOT) 수 {rc_ok} · 화면 KPI 어긋남 {hk or 0} · SVG/표 어긋남 {hs or 0}")

            # ── RPT-040 재고·창고 ──
            a, d, h, n = self.stable(lambda: self.indep_inventory(m1, m_last), lambda: dashboard.inventory(P),
                                     lambda: page("RPT-040 재고·창고 리포트(이 달)", f"/rpt/inventory?{qs}"))
            kv = self._kv(d)
            keys = ("material_receipts", "material_receipt_qty", "issues", "issue_qty", "product_receipts", "product_receipt_qty", "adjusts", "adjust_diff")
            kpi_ok = all(self._eq(kv[k], a[k]) for k in keys)
            ser_ok = ([p["value"] for p in d["series"]["receipt_qty"]] == a["s_receipt"] and [p["value"] for p in d["series"]["issue_qty"]] == a["s_issue"])
            tm_ok = [r["qty"] for r in d["tables"]["top_moves"]] == a["top_moves"]
            ls_ok = [r["days"] for r in d["tables"]["long_stay"]] == a["long_stay"]
            hk, hs = self.kpi_mismatch(h, d), self.series_mismatch(h, d)
            tm_html = self.table_html(h, "top_moves")
            tm_html_ok = tm_html is not None and len(tm_html) == len(d["tables"]["top_moves"])
            g15["inventory"] = {"kpis": {k: kv[k] for k in keys}, "indep": {k: a[k] for k in keys}}
            self.db(f"KPI 함수 {g15['inventory']['kpis']} · 독립 {g15['inventory']['indep']} · 변동 상위 {a['top_moves']} · 체류일 {a['long_stay']} · 읽기 {n}회")
            self.r.add("G-15", "RPT-040 재고·창고 — dashboard.inventory = 독립 SQL = 화면 · 변동 상위 · 장기 체류 · 실사 차이 (이 달)",
                       kpi_ok and ser_ok and tm_ok and ls_ok and not hk and not hs and tm_html_ok,
                       f"{ym}: 원자재 입고 {kv['material_receipts']}건 {kv['material_receipt_qty']} · 불출 {kv['issues']}건 {kv['issue_qty']} · 제품 입고 {kv['product_receipts']}건 {kv['product_receipt_qty']} · "
                       f"조정 {kv['adjusts']} 차이 {kv['adjust_diff']} = 독립 {[a[k] for k in keys]} · 구간 입고 · 불출 같음 {ser_ok} · 변동 상위 10 수량 {tm_ok} · "
                       f"장기 체류 10 체류일 {ls_ok} · 화면 KPI 어긋남 {hk or 0} · SVG/표 어긋남 {hs or 0} · 변동 상위 표 {len(tm_html or [])}행 · 읽기 {n}회")

            # ── RPT-050 출하 (거래처 C1) ──
            a, d, h, n = self.stable(lambda: self.indep_shipment(m1, m_last, C1), lambda: dashboard.shipment(P, C1),
                                     lambda: page("RPT-050 출하 리포트(거래처 C1 · 이 달)", f"/rpt/shipment?{qs}&customer={C1}"))
            kv = self._kv(d)
            keys = ("shipments", "qty", "customers", "due_rate", "on_time", "late_done", "late_open")
            kpi_ok = all(self._eq(kv[k], a[k]) for k in keys)
            ser_ok = ([p["value"] for p in d["series"]["qty"]] == a["s_qty"] and [p["value"] for p in d["series"]["shipments"]] == a["s_shipments"]
                      and [(p["label"], p["value"]) for p in d["series"]["by_customer"]] == [(c, q) for c, q, _ in a["by_customer"]][:10])
            bc_ok = [(r["customer_code"], r["qty"], r["shipments"]) for r in d["tables"]["by_customer"]] == a["by_customer"]
            s1 = [r for r in d["tables"]["recent"] if r["shipment_no"] == v["S1"]]
            s1_ok = len(s1) == 1 and s1[0]["qty"] == 60 and s1[0]["request_no"] == v["R1"]
            due_rows_ok = all(r["customer_code"] == f"{t}-C" for r in d["tables"]["due_rows"])
            hk, hs = self.kpi_mismatch(h, d), self.series_mismatch(h, d)
            g15["shipment"] = {"kpis": {k: kv[k] for k in keys}, "indep": {k: a[k] for k in keys}}
            self.db(f"KPI 함수 {g15['shipment']['kpis']} · 독립 {g15['shipment']['indep']} · 거래처별 {a['by_customer']} · 읽기 {n}회")
            self.r.add("G-15", "RPT-050 출하 — dashboard.shipment = 독립 SQL = 화면 · 납기 준수율 D-804(의뢰번호 있는 출하만) (거래처 C1 · 이 달)",
                       kpi_ok and ser_ok and bc_ok and s1_ok and due_rows_ok and not hk and not hs,
                       f"{ym} {t}-C: 출하 {kv['shipments']}건 {kv['qty']} · 거래처 {kv['customers']} · 준수율 {kv['due_rate']} (납기 안 {kv['on_time']} · 지연 완료 {kv['late_done']} · 지연 {kv['late_open']}) "
                       f"= 독립 {[a[k] for k in keys]} · 구간 수량 · 건수 · 거래처별 같음 {ser_ok} {bc_ok} · {v['S1']} 60 · 의뢰 {v['R1']} {s1_ok} · 준수율 모집단 전부 {t}-C {due_rows_ok} · "
                       f"화면 KPI 어긋남 {hk or 0} · SVG/표 어긋남 {hs or 0}")

            # ── RPT-060 수금 (거래처 C1) — asof 명시 ──
            a, d, h, n = self.stable(lambda: self.indep_collection(m1, m_last, C1), lambda: dashboard.collection(P, C1),
                                     lambda: page("RPT-060 수금 리포트(거래처 C1 · 이 달)", f"/rpt/collection?{qs}&customer={C1}"))
            kv, kt = self._kv(d), self._kt(d)
            kpi_ok = (self._eq(kv["collected"], a["collected"]) and self._eq(kv["sales"], a["sales"]) and kt["sales"] == a["sales_text"]
                      and self._eq(kv["shipped_qty"], a["shipped_qty"]) and self._eq(kv["over90"], a["over90"]) and kt["over90"] == a["over90_text"]
                      and kt["asof"] == a["asof_text"] and d["asof"] == a["asof"])
            ser_ok = ([(p["label"], p["value"]) for p in d["series"]["collected"]] == list(zip(a["months"], a["s_collected"]))
                      and [(p["label"], p["value"]) for p in d["series"]["shipped_qty"]] == list(zip(a["months"], a["s_shipped"]))
                      and [(p["label"], p["value"]) for p in d["series"]["aging"]] == list(zip(a["buckets"], a["s_aging"])))
            asof_line = re.search(r"ERP 업로드 기준일:\s*<strong>([^<]+)</strong>", h)
            asof_ok = asof_line is not None and asof_line.group(1).strip() == a["asof_text"] and self.kpi_text(h, "asof") == a["asof_text"]
            hk, hs = self.kpi_mismatch(h, d), self.series_mismatch(h, d)
            mon = self.table_html(h, "monthly")
            mon_ok = mon is not None and len(mon) == len(a["months"]) and all(
                row[0] == m and row[4] == ("미수집" if c is None else f"{c:,.0f}") for row, m, c in zip(mon, a["months"], a["s_collected"]))
            a060 = a
            g15["collection"] = {"kpis": {k: kv[k] for k in ("collected", "sales", "shipped_qty", "over90")}, "asof": kt["asof"],
                                 "indep": {k: a[k] for k in ("collected", "sales", "shipped_qty", "over90")}, "indep_asof": a["asof_text"]}
            self.db(f"KPI 함수 {g15['collection']['kpis']} asof {kt['asof']} · 독립 {g15['collection']['indep']} asof {a['asof_text']} · 연령 {a['s_aging']} · 읽기 {n}회")
            self.r.add("G-15", "RPT-060 수금 — dashboard.collection = 독립 SQL(ERP 반영 회차 테이블 직접) = 화면 · ERP 업로드 기준일(asof) 머리 · KPI 명시 (거래처 C1)",
                       kpi_ok and ser_ok and asof_ok and not hk and not hs and mon_ok,
                       f"{ym} {t}-C: 수금 {kv['collected']} · 매출 {kv['sales']} · 출하 수량 {kv['shipped_qty']} · 90일 초과 {kv['over90']}/{kt['over90']} · asof {kt['asof']} "
                       f"= 독립 {a['collected']} · {a['sales']} · {a['shipped_qty']} · {a['over90']}/{a['over90_text']} · asof {a['asof_text']}(= max(period_to) of 반영) · "
                       f"월별 수금 · 출하 · 연령 4구간 {a['s_aging']} 같음 {ser_ok} · 화면 머리 「ERP 업로드 기준일: {asof_line.group(1) if asof_line else '없음'}」 {asof_ok} · "
                       f"월 표 {len(mon or [])}행 {mon_ok} · 화면 KPI 어긋남 {hk or 0} · SVG/표 어긋남 {hs or 0}")

            # ── RPT-070 주문 흐름 (거래처 C1) ──
            a, d, h, n = self.stable(lambda: self.indep_order_flow(m1, m_last, C1), lambda: dashboard.order_flow(P, C1),
                                     lambda: page("RPT-070 주문 흐름 리포트(거래처 C1 · 이 달)", f"/rpt/order-flow?{qs}&customer={C1}"))
            kv, kt = self._kv(d), self._kt(d)
            stages = [s["stage"] for s in d["tables"]["stages"]]
            names = [nm for nm, _, _ in self.STAGES]
            kpi_ok = (self._eq(kv["requests"], a["requests"]) and self._eq(kv["shipped"], a["shipped"]) and stages == names
                      and all(self._eq(kv[f"avg_{i}"], a["avg"][nm]) for i, nm in enumerate(names, start=1))
                      and kv["bottleneck"] is None and kt["bottleneck"] == "미확정 (D-520)")
            st_ok = all(self._eq(s["avg_days"], a["avg"][s["stage"]]) and s["count"] == a["count"][s["stage"]] and self._eq(s["max_days"], a["max"][s["stage"]])
                        for s in d["tables"]["stages"])
            rows_ok = {r["request_no"]: r["days"] for r in d["tables"]["rows"]} == a["rows"] and all(r["bottleneck"] == "미확정 (D-520)" for r in d["tables"]["rows"])
            r1 = a["rows"].get(v["R1"], {})
            r1_ok = r1 and all(r1[nm] is not None for nm in names) and all(a["rows"].get(v["R2"], {nm: 0})[nm] is None for nm in names)
            ser_ok = [(p["label"], p["value"]) for p in d["series"]["avg_days"]] == [(nm, a["avg"][nm]) for nm in names]
            coll_ok = [(p["label"], p["value"]) for p in d["series"]["collected"]] == [(m, self.indep_ship_vs_coll_any(m, C1)[3]) for m in
                                                                                       [p["label"] for p in d["series"]["collected"]]]
            hk, hs = self.kpi_mismatch(h, d), self.series_mismatch(h, d)
            html_rows = self.table_html(h, "rows")
            row_html_ok = html_rows is not None and len(html_rows) == len(a["rows"]) and all(row[-1] == "미확정 (D-520)" for row in html_rows)
            asof_in_070 = "ERP 업로드 기준일" in h          # 날짜 글자만으로는 안 본다(조회 조건의 종료일과 겹친다)
            g15["order_flow"] = {"kpis": {"requests": kv["requests"], "shipped": kv["shipped"], **{f"avg_{i}": kv[f"avg_{i}"] for i in range(1, 6)}},
                                 "indep": {"requests": a["requests"], "shipped": a["shipped"], **{f"avg_{i}": a["avg"][nm] for i, nm in enumerate(names, start=1)}},
                                 "R1_days": r1, "asof": d["asof"], "asof_shown": asof_in_070}
            self.db(f"KPI 함수 {g15['order_flow']['kpis']} · 독립 {g15['order_flow']['indep']} · {v['R1']} 소요일 {r1} · asof 반환 {d['asof']} · 화면에 기준일 {asof_in_070} · 읽기 {n}회")
            self.r.add("G-15", "RPT-070 주문 흐름 — dashboard.order_flow = 독립 SQL = 화면 · 단계 5 소요일 · 병목 미확정 (D-520) · 수금 월 합 (거래처 C1)",
                       kpi_ok and st_ok and rows_ok and r1_ok and ser_ok and coll_ok and not hk and not hs and row_html_ok,
                       f"{ym} {t}-C: 의뢰 {kv['requests']} · 출하까지 {kv['shipped']} · 평균 소요일 {[kv[f'avg_{i}'] for i in range(1, 6)]} = 독립 {a['requests']} · {a['shipped']} · "
                       f"{[a['avg'][nm] for nm in names]} · 의뢰별 소요일 표 {rows_ok} ({v['R1']} 5단계 전부 수 {dict(r1)} · {v['R2']} 미착수 전부 None) · 단계 요약 {st_ok} · "
                       f"병목 KPI 「{kt['bottleneck']}」 · 표 병목 열 「미확정 (D-520)」 {row_html_ok} · 수금 월 합 = 독립 {coll_ok} · 화면 KPI 어긋남 {hk or 0} · SVG/표 어긋남 {hs or 0}")
            self.r.add("G-15", "ERP 업로드 기준일(asof) — 수금 대시보드 RPT-060 머리 · KPI 명시(G-15 기대값) · RPT-070 · 080 반환 asof 와 화면",
                       asof_ok,
                       f"RPT-060 화면 머리 · KPI asof {a060['asof_text']} {asof_ok} / RPT-070 반환 asof {d['asof']} · 화면에 「ERP 업로드 기준일」 글자 {asof_in_070}"
                       + (" — 거래처·월별 수금을 보이면서 기준일 글자가 없다 · 계약(F-RPT-071 · D-512 · D-805)이 070 에 명시를 요구하지 않아 게이트 수치는 아니다 · "
                          "DEF-QA2-005 경미(개발3 · 아키텍트 판단)" if (not asof_in_070 and d["asof"] is not None) else "")
                       + " / RPT-080 은 다음 행(주간 · 월간 머리에 표시 · 일일은 수금 절이 없어 없음)")

            # ── RPT-080 경영 요약 — 일일 · 주간 · 월간 (D-512) · 절 KPI = 개별 함수 = 독립 SQL = 화면 ──
            today = self.today
            spans = {"daily": (today, today), "weekly": (today - timedelta(days=6), today), "monthly": (today.replace(day=1), today)}
            sum_ok, sum_notes = True, []
            for tab, (f, tt) in spans.items():
                d = dashboard.summary((f, tt), C1)
                h = page(f"RPT-080 경영 요약 {tab}(거래처 C1 · 기준일 오늘)", f"/rpt/summary?tab={tab}&asof={today.isoformat()}&customer={C1}")
                secs = [s["screen_id"] for s in d["tables"]["sections"]]
                exp_secs = ["RPT-040", "RPT-050"] + ([] if tab == "daily" else ["RPT-010", "RPT-020", "RPT-060"])
                kind = {"daily": "일일", "weekly": "주간", "monthly": "월간"}[tab]
                # 절 KPI = 개별 함수(같은 기간 · 거래처 규칙) = 독립 SQL
                parts = {"RPT-040": (dashboard.inventory((f, tt)), self.indep_inventory(f, tt)),
                         "RPT-050": (dashboard.shipment((f, tt), C1), self.indep_shipment(f, tt, C1))}
                if tab != "daily":
                    parts.update({"RPT-010": (dashboard.orders((f, tt), C1), self.indep_orders(f, tt, C1)),
                                  "RPT-020": (dashboard.production((f, tt)), self.indep_production(f, tt)),
                                  "RPT-060": (dashboard.collection((f, tt), C1), self.indep_collection(f, tt, C1))})
                bad = []
                for s in d["tables"]["sections"]:
                    fn_d, ind = parts[s["screen_id"]]
                    fn_kv, fn_kt = self._kv(fn_d), self._kt(fn_d)
                    for k in s["kpis"]:
                        if not self._eq(k["value"], fn_kv.get(k["key"])) or k["text"] != fn_kt.get(k["key"]):
                            bad.append(f"{s['screen_id']}.{k['key']} 요약 {k['value']}/{k['text']} ≠ 함수 {fn_kv.get(k['key'])}/{fn_kt.get(k['key'])}")
                        if k["key"] in ind and k["key"] != "asof" and not self._eq(k["value"], ind[k["key"]]):
                            bad.append(f"{s['screen_id']}.{k['key']} 요약 {k['value']} ≠ 독립 {ind[k['key']]}")
                        if k["key"] == "asof" and k["text"] != ind["asof_text"]:
                            bad.append(f"RPT-060.asof 「{k['text']}」 ≠ 독립 「{ind['asof_text']}」")
                    sec_html = re.search(r'<section class="panel" data-section="' + s["screen_id"] + r'">(.*?)</section>', h, re.S)
                    if not sec_html:
                        bad.append(f"{s['screen_id']} 화면 절 없음")
                    else:
                        bad += [f"{s['screen_id']} {x}" for x in self.kpi_mismatch(sec_html.group(1), {"kpis": s["kpis"]})]
                hs = self.series_mismatch(h, d)
                kind_ok = d["kind"] == kind and f"<strong>{kind}</strong>" in h
                asof_ok80 = (("ERP 업로드 기준일" in h) == (d["asof"] is not None)) and (tab == "daily") == (d["asof"] is None)
                ok = secs == exp_secs and kind_ok and not bad and not hs and asof_ok80
                sum_ok &= ok
                sum_notes.append(f"{kind}({f}~{tt}) 절 {secs} {'OK' if secs == exp_secs else '≠ ' + str(exp_secs)} · KPI {len(d['kpis'])}개 어긋남 {bad or 0} · "
                                 f"구간 SVG/표 어긋남 {hs or 0} · 기준일 표시 {'있음' if 'ERP 업로드 기준일' in h else '없음'}(asof {d['asof']})")
                self.db(f"{kind} 절 {secs} · KPI {len(d['kpis'])} · 어긋남 {len(bad)}")
            g15["summary"] = sum_notes
            self.r.add("G-15", "RPT-080 경영 요약 — 일일(040 · 050) · 주간 · 월간(+ 010 · 020 · 060) 구성(D-512) · 절 KPI = 개별 함수 = 독립 SQL = 화면 · 해설 없음",
                       sum_ok, " / ".join(sum_notes))

            # ── 미수집 · 미확정 글자 — 0건은 None → 화면 「미수집」 (거래처 C2 = 의뢰 · 출하 없음 · ERP 만) ──
            d0 = dashboard.orders(P, C2)
            h0 = page("RPT-010 수주 리포트(거래처 C2 · 의뢰 없음)", f"/rpt/orders?{qs}&customer={C2}")
            none_kpi = all(k["value"] is None for k in d0["kpis"]) and all(self.kpi_text(h0, k["key"]) == "미수집" for k in d0["kpis"])
            s_req = self.series_html(h0, "requests")
            none_ser = s_req is not None and s_req[0] == [] and len(s_req[1]) == (m_last - m1).days + 1 and all(val == "미수집" for _, val in s_req[1])
            empty_msg = re.search(r'data-series="requests">.*?<p class="empty" role="status">미수집</p>', h0, re.S) is not None
            # 거래처 C1 의 일 구간 — 행이 없는 날은 미수집(0 으로 적지 않는다)
            h1 = page("RPT-010 수주 리포트(거래처 C1 — 일 구간 미수집 확인)", f"/rpt/orders?{qs}&customer={C1}")
            s1h = self.series_html(h1, "requests")
            a1 = self.indep_orders(m1, m_last, C1)
            day_none_ok = s1h is not None and [val for _, val in s1h[1]] == ["미수집" if x is None else f"{x:,.0f}" for x in a1["s_requests"]] \
                and sum(1 for _, val in s1h[0] if val == "미수집") == sum(1 for x in a1["s_requests"] if x is None)
            zero_text = "0</td>" in "".join(f"<td class=\"num\">{val}</td>" for _, val in s1h[1]) if s1h else True
            self.r.add("G-15", "미수집 · 미확정 글자 — 0건 None → 「미수집」(KPI · 표 · SVG title) · 로스율 「미확정 (D-519)」 · 병목 「미확정 (D-520)」 · 0 으로 지어내지 않음",
                       none_kpi and none_ser and empty_msg and day_none_ok and not zero_text,
                       f"{t}-C2(의뢰 0): KPI 4 전부 None → 화면 「미수집」 {none_kpi} · 일 구간 {len(s_req[1]) if s_req else 0}칸 표 전부 미수집 · SVG 없음(empty) {none_ser and empty_msg} / "
                       f"{t}-C(의뢰 있음): 행 없는 날 미수집 {sum(1 for x in a1['s_requests'] if x is None)}칸 = 화면 표 · SVG title 미수집 수 {day_none_ok} · 「0」 칸 {'있음' if zero_text else '없음'} / "
                       f"로스율 · 병목 글자는 RPT-020 · 030 · 070 행에서 확인 / 매출 열 미매핑 → 「미확정 (D-516)」 분기는 이 DB 의 열 매핑(시드 7 · 매핑됨)을 바꾸지 않아 미실측")

            # ── 거래처 필터 — 010 · 050 · 060 · 070 · 080 에만 (D-805) ──
            f_ok, f_notes = True, []
            for sid, path, fn in (("RPT-020", "production", dashboard.production), ("RPT-030", "quality", dashboard.quality), ("RPT-040", "inventory", dashboard.inventory)):
                dn = fn(P)
                dc = fn(P, C2)
                hc = page(f"{sid} ?customer= 무시 확인", f"/rpt/{path}?{qs}&customer={C2}")
                same = self._kv(dn) == self._kv(dc) and dn["series"] == dc["series"]
                no_sel = 'name="customer"' not in hc
                shown_same = not self.kpi_mismatch(hc, dn)
                f_ok &= same and no_sel and shown_same
                f_notes.append(f"{sid} 거래처 인자를 줘도 KPI · series 같음 {same}(반환 customer_id 는 인자를 되읊음 {dc['customer_id'] == C2} — 집계엔 안 씀) · "
                               f"화면에 거래처 칸 없음 {no_sel} · 화면 값 = 필터 없음 {shown_same}")
            d_c2 = {"RPT-010": dashboard.orders(P, C2), "RPT-050": dashboard.shipment(P, C2), "RPT-060": dashboard.collection(P, C2),
                    "RPT-070": dashboard.order_flow(P, C2), "RPT-080": dashboard.summary((m1, today), C2)}
            kv060 = self._kv(d_c2["RPT-060"])
            ind060 = self.indep_collection(m1, m_last, C2)
            applied = (all(k["value"] is None for k in d_c2["RPT-010"]["kpis"]) and all(k["value"] is None for k in d_c2["RPT-050"]["kpis"])
                       and all(k["value"] is None for k in d_c2["RPT-070"]["kpis"])
                       and self._eq(kv060["collected"], ind060["collected"]) and self._eq(kv060["sales"], ind060["sales"]) and kv060["shipped_qty"] is None
                       and all(k["value"] is None for k in d_c2["RPT-080"]["kpis"] if k["group"] in ("수주", "출하"))
                       and all(d_c2[s]["customer_id"] == C2 for s in d_c2))
            hsel = page("RPT-050 출하 리포트(거래처 C2 — 필터 적용 확인)", f"/rpt/shipment?{qs}&customer={C2}")
            sel_ok = f'<option value="{C2}" selected>' in hsel and self.kpi_text(hsel, "shipments") == "미수집"
            f_notes.append(f"RPT-010 · 050 · 070 (C2 = 의뢰 · 출하 없음) KPI 전부 None · RPT-060 C2 수금 {kv060['collected']} 매출 {kv060['sales']} = 독립 {ind060['collected']} · {ind060['sales']} · 출하 None · "
                           f"RPT-080 월간 C2 수주 · 출하 절 None · customer_id 반환 {applied} · 화면 select 에 C2 selected · 출하 미수집 {sel_ok}")
            self.r.add("G-15", "거래처 필터 — RPT-010 · 050 · 060 · 070 · 080 에만 적용 · 020 · 030 · 040 은 기간만 (D-805)",
                       f_ok and applied and sel_ok, " / ".join(f_notes))

        # ── G-05 — 대시보드 8 화면 GET 전후 업무 테이블 diff 0 (캐시 테이블 없음) ──
        paths = [f"/rpt/orders?{qs}&customer={C1}", f"/rpt/production?{qs}", f"/rpt/quality?{qs}", f"/rpt/inventory?{qs}",
                 f"/rpt/shipment?{qs}&customer={C1}", f"/rpt/collection?{qs}&customer={C1}", f"/rpt/order-flow?{qs}&customer={C1}",
                 f"/rpt/summary?tab=monthly&asof={today.isoformat()}&customer={C1}", "/rpt/summary?tab=daily", "/rpt/summary?tab=weekly",
                 "/rpt/orders", "/rpt/collection"]
        diff, codes = None, []
        for _ in range(3):
            before = self.table_counts()
            codes = [self.c("ex").get(p).status_code for p in paths]
            after = self.table_counts()
            diff = {k: (before[k], after.get(k)) for k in before if before[k] != after.get(k)}
            if not diff:
                break
        self.r.steps.append(Step("RPT 8 화면 GET 전후 행 수 (경영진 ex)", "ex", f"GET {len(paths)} 경로", sorted(set(codes)), f"diff {diff or 0}"))
        self.r.add("G-05", "RPT 대시보드 8 화면 GET 전후 업무 테이블 diff 0 (캐시 없음 · D-512)",
                   not diff and all(c == 200 for c in codes),
                   f"ex GET {len(paths)} (RPT-010~080 · 기본 기간 2 · 요약 탭 3) 응답 {sorted(set(codes))} · 업무 테이블 {len(before)} diff {diff or 0} "
                   f"(로그 4 제외 — 다른 세션이 동시에 쓴다 · GET 이 질의 로그를 쓰지 않는 것은 정적 행 「dataquery.ask 는 POST run_query 안에서만」)")

    # ════════════════════════════════════════════════════════════════════════
    # 정리
    # ════════════════════════════════════════════════════════════════════════
    def cleanup(self) -> str:
        c = self.conn
        self.cleanup_erp()
        mlots = [r["id"] for r in c.q("select id from material_lot where lot_no = any(%s) or receipt_id = any(%s)",
                                      (self.mlot_nos, self.receipt_ids or [0]))]
        wos = [r["id"] for r in c.q("select id from work_order where request_id = any(%s)", (self.req_ids or [0],))] + self.wo_ids
        plots = [r["id"] for r in c.q("select id from product_lot where work_order_id = any(%s) or lot_no = any(%s)",
                                      (wos or [0], self.plot_nos))]
        orders = [r["id"] for r in c.q("select id from shipment_order where customer_id = %s", (self.customer_id or 0,))]
        ship_nos = [r["shipment_no"] for r in c.q("select shipment_no from shipment_order where id = any(%s)", (orders,))]
        plot_nos = [r["lot_no"] for r in c.q("select lot_no from product_lot where id = any(%s)", (plots,))]
        mlot_nos = [r["lot_no"] for r in c.q("select lot_no from material_lot where id = any(%s)", (mlots,))]
        wo_nos = [r["work_order_no"] for r in c.q("select work_order_no from work_order where id = any(%s)", (wos or [0],))]
        keys = ([f"request_created:{n}" for n in self.req_nos] + [f"order_issued:{n}" for n in wo_nos]
                + [f"material_received:{n}" for n in mlot_nos] + [f"product_received:{n}" for n in plot_nos]
                + [f"shipment_ordered:{n}" for n in ship_nos] + [f"shipment_confirmed:{n}" for n in ship_nos])
        likes = [f"judged:{n}:%" for n in plot_nos] + [f"due:{n}:%" for n in self.req_nos]
        adj = c.q1("select count(*) as n from shipment_adjust where shipment_order_id = any(%s)", (orders,))["n"]
        sadj = c.q1("select count(*) as n from stock_adjust where (lot_kind = '원자재' and lot_id = any(%s)) "
                    "or (lot_kind = '제품' and lot_id = any(%s))", (mlots, plots))["n"]
        if adj or sadj:
            return f"정리 보류 — 불변 이력 shipment_adjust {adj} · stock_adjust {sadj} 행이 있어 LOT · 출하를 지우지 않았다"
        with c.tx() as cur:
            cur.execute("delete from notification where dedupe_key = any(%s) or dedupe_key like any(%s)", (keys, likes or [""]))
            cur.execute("delete from stock_move where (lot_kind = '원자재' and lot_id = any(%s)) or (lot_kind = '제품' and lot_id = any(%s))"
                        " or (ref_table = 'shipment_order' and ref_id = any(%s))", (mlots, plots, orders))
            cur.execute("delete from shipment_line where shipment_order_id = any(%s)", (orders,))
            cur.execute("delete from shipment_order where id = any(%s)", (orders,))
            cur.execute("delete from product_inspection where product_lot_id = any(%s)", (plots,))
            cur.execute("delete from incoming_inspection where material_lot_id = any(%s)", (mlots,))
            cur.execute("delete from lot_stock where (lot_kind = '원자재' and lot_id = any(%s)) or (lot_kind = '제품' and lot_id = any(%s))",
                        (mlots, plots))
            cur.execute("delete from product_receipt where product_lot_id = any(%s) or work_order_id = any(%s)", (plots, wos or [0]))
            cur.execute("delete from product_lot where id = any(%s) and parent_lot_id is not null", (plots,))
            cur.execute("delete from product_lot where id = any(%s)", (plots,))
            cur.execute("delete from production_daily where work_order_id = any(%s)", (wos or [0],))
            cur.execute("delete from inspection_daily where work_order_id = any(%s)", (wos or [0],))
            cur.execute("delete from material_issue_line where work_order_id = any(%s) or material_lot_id = any(%s)", (wos or [0], mlots))
            cur.execute("delete from material_issue where id = any(%s)", (self.issue_ids or [0],))
            cur.execute("delete from material_lot where id = any(%s) and parent_lot_id is not null", (mlots,))
            cur.execute("delete from material_lot where id = any(%s)", (mlots,))
            cur.execute("delete from material_receipt where id = any(%s)", (self.receipt_ids or [0],))
            cur.execute("delete from work_sequence where work_order_id = any(%s)", (wos or [0],))
            cur.execute("delete from work_order where id = any(%s)", (wos or [0],))
            cur.execute("delete from production_request where id = any(%s)", (self.req_ids or [0],))
            if self.locations:
                cur.execute("delete from location where id = any(%s)", (self.locations,))
            if self.warehouse_id:
                cur.execute("delete from warehouse where id = %s", (self.warehouse_id,))
            if self.customer_id:
                cur.execute("delete from customer where id = %s", (self.customer_id,))
            if self.customer2_id:
                cur.execute("delete from customer where id = %s", (self.customer2_id,))
            if self.items:
                cur.execute("delete from item where id = any(%s)", (self.items,))
        left = c.q1("""select (select count(*) from item where item_code like %(t)s) + (select count(*) from customer where customer_code like %(t)s)
                            + (select count(*) from warehouse where warehouse_code like %(t)s)
                            + (select count(*) from erp_customer_map where erp_code like %(t)s)
                            + (select count(*) from erp_upload_batch where source_file like %(t)s) as n""", {"t": f"{self.tag}%"})["n"]
        return f"정리 완료 — {self.tag} 기준정보 · ERP 회차 · 매핑 남은 행 {left} (sys_change_log · sys_access_log 감사 행과 채번 카운터는 남긴다)"

    def cleanup_erp(self) -> None:
        """W2 가 만든 ERP 회차(행 · 알림) · 거래처 매핑을 지우고, 바꿨으면 열 매핑을 원래대로 되돌린다."""
        from lcomfine.app import erp_upload

        c = self.conn
        bids = [r["id"] for r in c.q("select id from erp_upload_batch where id = any(%s) or source_file like %s",
                                     (self.erp_batches or [0], f"{self.tag}%"))]
        with c.tx() as cur:
            if bids:
                cur.execute("delete from notification where event_key = 'erp_upload_error' and link = any(%s)",
                            ([erp_upload.error_link(b) for b in bids],))
                cur.execute("delete from erp_upload_row where batch_id = any(%s)", (bids,))
                cur.execute("delete from erp_upload_batch where id = any(%s)", (bids,))      # replaced_by FK 는 문장 끝에 검사된다
            cur.execute("delete from erp_customer_map where erp_code = any(%s) or customer_id = any(%s)",
                        (self.erp_codes or [""], [x for x in (self.customer_id, self.customer2_id) if x] or [0]))
            if self.cmap_before is not None:
                cur.execute("delete from erp_column_map")
                for b in self.cmap_before:
                    cur.execute("insert into erp_column_map (field, column_name, created_by) values (%s, %s, %s)",
                                (b["field"], b["column_name"], b["created_by"]))

    # ════════════════════════════════════════════════════════════════════════
    def run(self, keep: bool = False, only: set[str] | None = None) -> None:
        """`only` = 돌릴 단계의 메서드 이름 집합(테스트가 일부만 돌릴 때 · 예 {"base", "e2e", "collections"}). 기본 전부."""
        stages = [("G-06", "기준정보 BAS (HTTP)", self.base), ("G-06", "E2E 의뢰 → 출하 확정", self.e2e),
                  ("G-07", "추적", self.traces), ("G-09", "차단", self.blocks), ("G-10", "할당", self.allocation),
                  ("G-07", "재판정 후 출하", self.rejudge_after_shipment), ("G-11", "알림", self.notifications),
                  ("G-05", "조회 diff", self.readonly_diff), ("G-10", "F-SHP-012 취소 = 할당 해제", self.shp012_cancel),
                  ("G-12", "수금 — 회차 교체 · 오류 행 · erp_upload_error · 출하 대비 수금", self.collections),
                  ("G-15", "대시보드 8 — 독립 재계산 · 화면 대조", self.dashboards)]
        if only:
            stages = [s for s in stages if s[2].__name__ in only]
        failed = None
        try:
            for gate, name, fn in stages:
                if failed:
                    self.r.add(gate, f"{name}", FAIL, f"앞 단계 실패로 돌리지 못함 — {failed}")
                    continue
                try:
                    fn()
                    if name.startswith("E2E"):
                        self.r.add("G-06", "E2E 의뢰 → 출하 확정 (HTTP 만)", True,
                                   f"{self.v['R1']} → {self.v['WO']} → {self.v['A1']}/{self.v['A1s']}/{self.v['A2']} → "
                                   f"{self.v['P1']} → {self.v['S1']} 확정 · 부서 sl · pr · sh · 요청 {len(self.r.steps)} 전부 기대 응답")
                except StepFailed as exc:
                    failed = f"{name}: {exc}"
                    self.r.add(gate, f"{name}", FAIL, str(exc))
                except Exception as exc:  # noqa: BLE001 — 예상 못한 실패도 그대로 FAIL 로 보인다(조용히 넘기지 않는다)
                    failed = f"{name}: {type(exc).__name__} {exc}"
                    self.r.add(gate, f"{name}", FAIL, f"{type(exc).__name__}: {str(exc)[:300]}")
            # 재지 못한 게이트는 행이 없어 PASS 로 보이면 안 된다
            need = set()
            if not only or "collections" in only:
                need |= MUST_MEASURE
            if not only or "dashboards" in only:
                need |= MUST_MEASURE_W3
            if not only or "traces" in only:
                need |= MUST_MEASURE_W4
            for gate, prefix in sorted(need):
                if not any(c.gate == gate and c.item.startswith(prefix) for c in self.r.checks):
                    self.r.add(gate, f"{prefix} ({GATE_NAMES[gate]})", FAIL, f"수금 · 대시보드 단계가 돌지 못해 재지 못함 — {failed or '행 없음'}")
        finally:
            if keep:
                self.r.notes.append(f"--keep — {self.tag} 행을 지우지 않았다")
            else:
                try:
                    self.r.notes.append(self.cleanup())
                except Exception as exc:  # noqa: BLE001
                    self.r.notes.append(f"정리 실패 — {type(exc).__name__}: {exc}")


# ════════════════════════════════════════════════════════════════════════════
def seed_idempotence(report: Report) -> None:
    """개발 시드 2회 — 행 수 diff 0 (G-21 은 QA3 판정 · 여기서는 참고 행)."""
    import importlib

    from lcomfine.db import conn

    def counts():
        names = [r["table_name"] for r in conn.q("""select table_name from information_schema.tables
                                                     where table_schema = 'public' and table_type = 'BASE TABLE'""")]
        return {n: conn.q1(f'select count(*) as n from "{n}"')["n"] for n in names if n not in ("sys_access_log", "sys_change_log")}

    mods = [m for m in ("seed_dev1", "seed_dev2", "seed_dev3") if (ROOT / f"src/lcomfine/db/{m}.py").exists()]
    out = []
    for m in mods:
        mod = importlib.import_module(f"lcomfine.db.{m}")
        mod.run()
        a = counts()
        mod.run()
        b = counts()
        d = {k: (a[k], b.get(k)) for k in a if a[k] != b.get(k)}
        out.append((m, d))
    report.notes.append("시드 멱등(참고 · G-21 은 QA3) — " + " · ".join(f"{m}.run() 2회 diff {d or 0}" for m, d in out)
                        + f" · 없는 개발 시드 {sorted(set(['seed_dev1', 'seed_dev2', 'seed_dev3']) - set(mods))}")
    report.seed = out  # type: ignore[attr-defined]


def run_all(keep: bool = False) -> Report:
    report = Report()
    try:
        static_scan(report)
    except Exception as exc:  # noqa: BLE001
        report.add("G-05", "쓰기 경계 — 정적 스캔", FAIL, f"스캔 실패 {type(exc).__name__}: {exc}")
    try:
        static_col(report)
    except Exception as exc:  # noqa: BLE001
        report.add("G-13", "수금 수정 경로 없음 — 정적 스캔", FAIL, f"스캔 실패 {type(exc).__name__}: {exc}")
    try:
        Flow(report).run(keep=keep)
    except Exception as exc:  # noqa: BLE001
        report.add("G-06", "E2E 흐름", FAIL, f"흐름을 시작하지 못함 {type(exc).__name__}: {exc}")
    for gid, w in LATER.items():
        report.add(gid, GATE_NAMES[gid], FAIL, f"{w} — 화면 · 기능이 아직 없다 (다음 웨이브에서 잰다)")
    return report


def main(argv: list[str] | None = None) -> int:
    argv = sys.argv[1:] if argv is None else argv
    from lcomfine.app.settings import get_settings

    get_settings()
    report = run_all(keep="--keep" in argv)
    if "--seed" in argv:
        seed_idempotence(report)
    for gid in GATE_NAMES:
        for c in [c for c in report.checks if c.gate == gid]:
            print(f"{c.gate}  {c.item}  {c.status}  {c.measured}")
    print()
    for n in report.notes:
        print(f"참고: {n}")
    if "--transcript" in argv:
        print()
        print("| # | 단계 | 계정 | 요청 | 응답 | 번호 · 메시지 | DB 확인 |")
        print("|---|---|---|---|---|---|---|")
        for i, s in enumerate(report.steps, 1):
            cell = lambda x: str(x).replace("|", "\\|")  # noqa: E731
            print(f"| {i} | {cell(s.step)} | {s.account} | `{cell(s.request)}` | {s.status} | {cell(s.numbers)} | {cell(s.db)} |")
    fails = sum(1 for c in report.checks if c.status == FAIL)
    by_gate = {g: ("FAIL" if any(c.status == FAIL for c in report.checks if c.gate == g) else
                   "미검증" if any(c.status != PASS for c in report.checks if c.gate == g) else "PASS") for g in GATE_NAMES}
    print()
    print("게이트 요약: " + " · ".join(f"{g} {s}" for g, s in by_gate.items()) + f" — FAIL 행 {fails} → 종료코드 {min(fails, 125)}")
    return min(fails, 125)


if __name__ == "__main__":
    raise SystemExit(main())
