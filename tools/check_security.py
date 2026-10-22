#!/usr/bin/env python
"""G-16~G-22 보안 · 운영 판정 (QA3) — `make check-security`. 서버를 띄우지 않고 TestClient 로 앱을 직접 부른다.

출력은 `G-nn  항목  PASS|FAIL  실측` 행이다(gate.py `_ROW_RE`). 종료코드 = FAIL 행 수.
  · G-16 데이터 질의 · G-17 키 없음 · G-18 SSO 는 W3/W4 범위 — 아직 기능이 없으므로 FAIL 「W3/W4」 로 찍는다(게이트를 낮추지 않는다).
  · G-19 권한 403 전수 · G-20 세션 상태 · G-21 비밀 · G-22 변경 감사 로그 를 실제 요청으로 잰다.
  · G-24 는 compose 문법 · Dockerfile 정적 검사만(컨테이너는 띄우지 않는다 — gate.py 가 G-24 를 따로 판정한다).
  · `X-nn` 행은 조용한 실패 사냥(DB 끊김 · 이상한 입력)이다 — 게이트 번호가 아니라 gate.py 는 읽지 않지만 FAIL 이면 종료코드에 센다.
  · W2 수금(COL): 화면 7 × 부서 8 · 쓰기 4 × 부서 8 · 비활성/로그아웃 세션의 COL 401 · 쓰기 4 의 변경 로그 · 업로드 보안(X-05~X-08) 을 더 잰다.
    업로드 회차는 1985년 기준 기간(다른 테스트는 1990~1993 · 자기 접두) · 파일 이름 `q3…` · 끝에 지운다. 열 매핑은 **바꾸지 않는다**(지금 매핑의 열 이름으로 파일을 만든다).

계정: 부서 8 의 **임시 계정**(접두 `q3…` · 난수 비밀번호 · 이름 `(예시)`)을 만들어 쓰고 끝나면 지운다. 시드 계정(sl pr sh ds ac st ex admin)은
로그인 확인에만 쓴다 — 시드 계정의 상태를 바꾸면 `session_epoch` 트리거가 같은 DB 를 쓰는 다른 검사·테스트의 세션까지 끊기 때문이다.
권한 판정은 부서 단위(`sys_permission`)라 임시 계정과 시드 계정의 판정은 같다. 비밀번호 · 비밀 값은 어디에도 찍지 않는다.

    uv run python tools/check_security.py               # 판정표
    uv run python tools/check_security.py --run-seeds   # + 시드 1회 더 돌려 행 수 diff (G-21 시드 멱등 — gate-full 과 같은 측정)
    uv run --offline --with playwright python tools/check_security.py --e2e   # 화면 캡처(1280 · 390) → outputs/e2e/
"""

from __future__ import annotations

import warnings

warnings.filterwarnings("ignore", message="Using `httpx` with `starlette.testclient`")

import importlib
import inspect
import io
import json
import os
import re
import secrets
import subprocess
import sys
import time
import uuid
import zipfile
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

from lcomfine.app.settings import get_settings  # noqa: E402 — .env 를 읽는다

PASS, FAIL, WARN = "PASS", "FAIL", "WARN"
SEED_ACCOUNTS = ("sl", "pr", "sh", "ds", "ac", "st", "ex", "admin")
ABSENT_DB_DSN = "postgresql:///lcomfine_qa3_absent_db"   # 없는 DB — DB 끊김 재현용 (같은 소켓 · 이름만 없는 것)
E2E_DIR = ROOT / "outputs" / "e2e"
E2E_PATHS = ("/", "/wo/requests", "/shp/orders", "/lot/backward", "/ia")
E2E_WIDTHS = (1280, 390)


@dataclass
class Result:
    gid: str
    item: str
    status: str
    measured: str
    bad: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.status == PASS

    def line(self) -> str:
        return f"{self.gid}  {self.item}  {self.status}  {self.measured}"


def res(gid: str, item: str, ok: bool, measured: str, bad: list[str] | None = None, *, soft: bool = False) -> Result:
    st = PASS if ok else (WARN if soft else FAIL)
    bad = bad or []
    if bad:
        measured += " — " + " / ".join(bad[:6]) + (f" … 외 {len(bad) - 6}" if len(bad) > 6 else "")
    return Result(gid, item, st, measured, bad)


# ── 앱 · 계정 ───────────────────────────────────────────────────────────
def _app():
    from lcomfine.app.main import app

    return app


def new_client():
    from fastapi.testclient import TestClient

    return TestClient(_app(), raise_server_exceptions=False)


def login(login_id: str, password: str):
    c = new_client()
    r = c.post("/login", data={"login_id": login_id, "password": password}, follow_redirects=False)
    if r.status_code != 303:
        raise AssertionError(f"{login_id} 로그인 {r.status_code}")
    return c


class Lab:
    """임시 계정 · 임시 행의 수명. `with Lab() as lab:` — 끝나면 만든 것을 지운다(로그 행은 남는다 — 감사 로그는 지우지 않는다)."""

    def __init__(self) -> None:
        self.prefix = "q3" + uuid.uuid4().hex[:8]
        self.users: dict[str, tuple[str, str, int]] = {}    # 부서 코드 → (login_id, 비밀번호, user_id)
        self._clients: dict[str, object] = {}
        self.extra_users: list[str] = []

    def __enter__(self) -> "Lab":
        from lcomfine.app import rbac

        for r in rbac.roles():
            self.users[r.code] = self.make_user(r.code, r.code)
        return self

    def __exit__(self, *exc) -> None:
        self.cleanup()

    def make_user(self, role_code: str, tag: str) -> tuple[str, str, int]:
        from lcomfine.app import auth
        from lcomfine.db import conn

        login_id, pw = f"{self.prefix}{tag}".lower(), secrets.token_urlsafe(12)
        uid = conn.q1("insert into sys_user (login_id, user_name, password_hash, role_code, created_by) "
                      "values (%s, %s, %s, %s, 'qa3') returning user_id",
                      (login_id, f"QA3 임시 {tag} (예시)", auth.hash_password(pw), role_code))["user_id"]
        return login_id, pw, uid

    def client(self, role_code: str):
        if role_code not in self._clients:
            lid, pw, _ = self.users[role_code]
            self._clients[role_code] = login(lid, pw)
        return self._clients[role_code]

    def fresh(self, role_code: str):
        lid, pw, _ = self.users[role_code]
        return login(lid, pw)

    def cleanup(self) -> None:
        from lcomfine.db import conn

        like = self.prefix + "%"
        up = self.prefix.upper() + "%"
        with conn.tx() as cur:
            cur.execute("select request_no from production_request where item_id in (select id from item where item_code ilike %s)",
                        (like,))
            reqs = [r["request_no"] for r in cur.fetchall()]
            # 이상한 입력이 잘못 저장된 경우까지 — 원자재 LOT · 입고 · 출하 지시 (만든 품목 · 거래처 기준)
            cur.execute("select id, lot_no, receipt_id from material_lot where item_id in (select id from item where item_code ilike %s)",
                        (like,))
            mls = cur.fetchall()
            m_ids, r_ids = [r["id"] for r in mls], [r["receipt_id"] for r in mls]
            cur.execute("delete from notification where split_part(dedupe_key, ':', 2) = any(%s)", ([r["lot_no"] for r in mls],))
            cur.execute("delete from stock_move where lot_kind = '원자재' and lot_id = any(%s)", (m_ids,))
            cur.execute("delete from lot_stock where lot_kind = '원자재' and lot_id = any(%s)", (m_ids,))
            cur.execute("delete from incoming_inspection where material_lot_id = any(%s)", (m_ids,))
            cur.execute("delete from material_lot where id = any(%s)", (m_ids,))
            cur.execute("delete from material_receipt where id = any(%s) or supplier_name ilike %s", (r_ids, like))
            cur.execute("select id from shipment_order where customer_id in (select id from customer where customer_code ilike %s)", (like,))
            s_ids = [r["id"] for r in cur.fetchall()]
            cur.execute("delete from shipment_line where shipment_order_id = any(%s)", (s_ids,))
            cur.execute("delete from shipment_order where id = any(%s)", (s_ids,))
            cur.execute("delete from notification where dedupe_key ilike %s or message ilike %s or dedupe_key = any(%s)",
                        (f"%{self.prefix}%", f"%{self.prefix}%", [f"request_created:{n}" for n in reqs]))
            cur.execute("delete from production_request where item_id in (select id from item where item_code ilike %s)", (like,))
            cleanup_col(cur, like)
            cur.execute("delete from standard where standard_code ilike %s", (like,))
            cur.execute("delete from location where warehouse_id in (select id from warehouse where warehouse_code ilike %s)", (like,))
            cur.execute("delete from location where location_code ilike %s", (like,))
            cur.execute("delete from warehouse where warehouse_code ilike %s", (like,))
            cur.execute("delete from item where item_code ilike %s or item_code ilike %s", (like, up))
            cur.execute("delete from customer where customer_code ilike %s", (like,))
            cur.execute("delete from sys_common_code where code ilike %s or group_code ilike %s", (like, like))
            cur.execute("delete from notification_setting where user_id in (select user_id from sys_user where login_id ilike %s)",
                        (like,))
            cur.execute("delete from sys_user where login_id ilike %s", (like,))


def router_routes() -> list[tuple[str, str, str, object, list[str]]]:
    """라우터 모듈의 (모듈, 메서드, 경로, endpoint, 의존성이 잡은 화면/기능 ID) — `main.py` 가 include 하는 그대로."""
    from lcomfine.app import nav

    def deps(d, out):
        for s in d.dependencies:
            for cell in (getattr(s.call, "__closure__", None) or ()):
                try:
                    v = cell.cell_contents
                except ValueError:
                    continue
                if isinstance(v, str) and re.fullmatch(r"(F-[A-Z]+-\d{3}|[A-Z]+-\d{3})", v):
                    out.append(v)
            deps(s, out)
        return out

    out = []
    for m in nav.MODULES:
        try:
            mod = importlib.import_module(f"lcomfine.app.routers.{m}")
        except ModuleNotFoundError:
            continue
        for r in mod.router.routes:
            for meth in sorted(getattr(r, "methods", None) or ()):
                out.append((m, meth, r.path, r.endpoint, deps(r.dependant, [])))
    return out


def write_functions_routed():
    """W1 에 라우트가 있는 쓰기 기능 — [(Function, 라우트 경로)]."""
    from lcomfine.app import contracts

    routed = {(meth, path) for _, meth, path, _, _ in router_routes()}
    return [f for f in contracts.write_functions() if (f.method, f.path) in routed]


def concrete(path: str, value: str = "0") -> str:
    return re.sub(r"\{[^}]+\}", value, path)


# ── G-19 권한 ───────────────────────────────────────────────────────────
def check_static_require_fn() -> Result:
    """라우터의 POST 마다 자기 기능 ID 의 `require_fn` 이 걸렸는가 (기능 목록의 API 열과 글자 그대로)."""
    from lcomfine.app import contracts

    by_api = {f.api: f for f in contracts.functions()}
    posts = [r for r in router_routes() if r[1] == "POST"]
    bad = []
    for m, meth, path, _, ids in posts:
        fn = by_api.get(f"POST {path}")
        if fn is None:
            bad.append(f"POST {path}: 기능 목록에 없는 쓰기 라우트")
        elif fn.id not in ids:
            bad.append(f"POST {path}: require_fn({fn.id}) 없음 (잡힌 것 {ids})")
    gets = [r for r in router_routes() if r[1] == "GET" and not r[4]]
    bad += [f"GET {p}: 권한 의존성 없음" for _, _, p, _, _ in gets]
    return res("G-19", "쓰기 라우트 require_fn", not bad, f"POST {len(posts)} · GET {len([r for r in router_routes() if r[1] == 'GET'])} 중 빠짐 {len(bad)}", bad)


def check_no_session_401() -> Result:
    """세션 없이 모든 W1 쓰기 API POST → 401 (JSON). 화면 GET 은 JSON 이면 401 · 브라우저면 /login 303."""
    from lcomfine.app import nav

    c = new_client()
    fns = write_functions_routed()
    bad = []
    for f in fns:
        r = c.post(concrete(f.path), data={})
        if r.status_code != 401:
            bad.append(f"{f.id} {r.status_code}")
    n_get = 0
    for sc in nav.SCREENS:
        if sc.common:
            continue
        n_get += 1
        r = c.get(sc.path)
        if r.status_code != 401:
            bad.append(f"GET {sc.path} {r.status_code}")
        r = c.get(sc.path, headers={"accept": "text/html"}, follow_redirects=False)
        if r.status_code != 303 or not r.headers.get("location", "").startswith("/login"):
            bad.append(f"GET {sc.path} (html) {r.status_code} → {r.headers.get('location')}")
    return res("G-19", "세션 없음 401", not bad, f"쓰기 POST {len(fns)} · 화면 GET {n_get}×2 중 어긋남 {len(bad)}", bad)


def check_write_403(lab: Lab) -> Result:
    """쓰기 기능 × 부서 8 전수 — 계약(function-list 권한 열)에 W 가 없는 부서는 403, W 부서는 403·401 이 아니다(본문이 비어 422/404)."""
    from lcomfine.app import rbac

    fns = write_functions_routed()
    bad, n403, nother = [], 0, 0
    for f in fns:
        for r in rbac.roles():
            expect_403 = (r.code not in f.roles_w) and not f.permission_exception
            got = lab.client(r.code).post(concrete(f.path), data={}).status_code
            if expect_403:
                n403 += 1
                if got != 403:
                    bad.append(f"{f.id}×{r.code} {got} (403 기대)")
            else:
                nother += 1
                if got in (401, 403) or got >= 500:
                    bad.append(f"{f.id}×{r.code} {got} (W 칸 — 403/401/5xx 아님 기대)")
    return res("G-19", "R·없음 부서 쓰기 403", not bad,
               f"쓰기 기능 {len(fns)} × 부서 {len(rbac.roles())} — 403 기대 {n403} · 통과 기대 {nother} · 어긋남 {len(bad)}", bad)


def check_screen_403(lab: Lab) -> Result:
    """화면 67 × 부서 8 — 계약 칸이 없으면 403 · 있으면 200, 그리고 메뉴(메인 화면의 링크)에 칸 없는 화면 경로가 없다."""
    from lcomfine.app import contracts, nav, rbac

    cells = {(sid, rc): lv for sid, rc, lv in contracts.permission_cells()}
    routed = {(meth, path) for _, meth, path, _, _ in router_routes()}
    screen_paths = {sc.path for sc in nav.SCREENS}
    read_fns = [f for f in contracts.functions() if not f.is_write and (f.method, f.path) in routed and f.path not in screen_paths]
    bad, n = [], 0
    for r in rbac.roles():
        c = lab.client(r.code)
        for sc in nav.SCREENS:
            if sc.common:
                continue
            n += 1
            want = 200 if cells.get((sc.screen_id, r.code)) else 403
            got = c.get(sc.path).status_code
            if got != want:
                bad.append(f"{sc.screen_id}×{r.code} {got} ({want} 기대)")
        for f in read_fns:   # 화면 경로가 아닌 조회·출력 기능(출력 · 하위 경로) — 칸 없으면 403, 칸 있으면 403/401/5xx 가 아니다
            n += 1
            has = bool(cells.get((f.screen_id, r.code)))
            got = c.get(concrete(f.path)).status_code
            if (not has and got != 403) or (has and (got in (401, 403) or got >= 500)):
                bad.append(f"{f.id}×{r.code} GET {got} ({'통과' if has else '403'} 기대)")
        html = c.get("/", headers={"accept": "text/html"}).text
        hrefs = set(re.findall(r'href="([^"?#]+)', html))
        for sc in nav.SCREENS:
            if not sc.common and sc.path != "/" and not cells.get((sc.screen_id, r.code)) and sc.path in hrefs:
                bad.append(f"메뉴 {r.code}: 칸 없는 {sc.screen_id} 링크")
    return res("G-19", "칸 밖 화면 GET 403 · 메뉴 숨김", not bad,
               f"화면 {len([s for s in nav.SCREENS if not s.common])} + 하위 조회·출력 {len(read_fns)} × 부서 8 = GET {n} · 어긋남 {len(bad)}", bad)


def check_permission_change(lab: Lab) -> Result:
    """권한 칸은 DB 데이터 — SYS-021 로 칸 하나(EX × WO-040 R)를 지우면 다음 요청 403, 되돌리면 200. 칸은 finally 에서 원래 값으로."""
    from lcomfine.app import rbac
    from lcomfine.db import conn

    rc, sid, path = "EX", "WO-040", "/wo/status"
    orig = conn.q1("select level from sys_permission where role_code = %s and screen_id = %s", (rc, sid))
    orig_lv = orig["level"] if orig else ""
    c, adm = lab.client(rc), lab.client("ADM")
    try:
        a = c.get(path).status_code
        s1 = adm.post("/sys/roles", data={f"cell:{rc}:{sid}": ""}).status_code
        b = c.get(path).status_code
        s2 = adm.post("/sys/roles", data={f"cell:{rc}:{sid}": orig_lv or "R"}).status_code
        d = c.get(path).status_code
    finally:
        if orig_lv:
            conn.x("insert into sys_permission (role_code, screen_id, level) values (%s, %s, %s) "
                   "on conflict (role_code, screen_id) do update set level = excluded.level", (rc, sid, orig_lv))
        else:
            conn.x("delete from sys_permission where role_code = %s and screen_id = %s", (rc, sid))
        rbac.invalidate()
    ok = (a, s1, b, s2, d) == (200, 200, 403, 200, 200) and orig_lv == "R"
    return res("G-19", "권한 칸 변경 다음 요청 반영", ok,
               f"{rc}×{sid} 칸 {orig_lv or '없음'} — 전 {a} · 칸 지움 {s1} → 다음 요청 {b} · 되돌림 {s2} → {d}")


def check_idor(lab: Lab) -> Result:
    """ID 바꿔치기 — 다른 부서 알림 처리 404(F-HOME-022 계약) · W 없는 부서의 남의 의뢰 수정 403(F-WO-012) · 행은 그대로."""
    from lcomfine.db import conn

    bad = []
    p = lab.prefix
    n = conn.q1("insert into notification (event_key, to_role, link, dedupe_key, message) values "
                "('request_created', 'PR', '/wo/requests', %s, %s) returning id",
                (f"qa3:{p}:idor", f"QA3 IDOR 시험 {p} (예시)"))["id"]
    for role in ("SL", "SH", "DS", "AC", "ST", "EX"):
        r = lab.client(role).post(f"/home/notifications/{n}/handle")
        if r.status_code != 404:
            bad.append(f"알림(PR) 처리 by {role} {r.status_code} (404 기대)")
    row = conn.q1("select handled_at from notification where id = %s", (n,))
    if row["handled_at"] is not None:
        bad.append("다른 부서가 알림을 처리해 버렸다")
    r = lab.client("PR").post(f"/home/notifications/{n}/handle")
    if r.status_code != 200 or conn.q1("select handled_by from notification where id = %s", (n,))["handled_by"] is None:
        bad.append(f"자기 부서(PR) 처리 {r.status_code} (200 + handled_by 기대)")
    for bogus in ("999999999999", "abc", "-1"):
        r = lab.client("PR").post(f"/home/notifications/{bogus}/handle")
        if r.status_code != 404:
            bad.append(f"알림 id {bogus!r} {r.status_code} (404 기대)")
    # 의뢰 수정 — 영업부가 만든 의뢰를 W 없는 부서가 고친다
    cu = conn.q1("insert into customer (customer_code, customer_name, created_by) values (%s, %s, 'qa3') returning id",
                 (f"{p}C", f"{p} 거래처 (예시)"))["id"]
    it = conn.q1("insert into item (item_code, item_name, item_type, created_by) values (%s, %s, '제품', 'qa3') returning id",
                 (f"{p}FG", f"{p} 제품 (예시)"))["id"]
    rq = conn.q1("insert into production_request (request_no, customer_id, item_id, qty, due_date, status, created_by) "
                 "values (%s, %s, %s, 10, current_date + 7, '확인대기', 'qa3') returning id",
                 (f"{p}REQ", cu, it))["id"]
    body = {"customer_id": str(cu), "item_id": str(it), "qty": "999", "due_date": "2027-01-31", "status": "확인대기",
            "note": "QA3 바꿔치기"}
    for role in ("DS", "AC", "ST", "EX", "SH"):
        r = lab.client(role).post(f"/wo/requests/{rq}", data=body)
        if r.status_code != 403:
            bad.append(f"의뢰 수정 by {role} {r.status_code} (403 기대)")
    if conn.q1("select qty from production_request where id = %s", (rq,))["qty"] != 10:
        bad.append("W 없는 부서가 의뢰 수량을 바꿨다")
    # API 명세 — 미로그인 401 · 시스템 관리 화면을 못 보는 부서 403 · /docs 없음
    anon = new_client()
    for path, want in (("/openapi.json", 401), ("/docs", 404), ("/redoc", 404)):
        got = anon.get(path).status_code
        if got != want:
            bad.append(f"미로그인 GET {path} {got} ({want} 기대)")
    got = lab.client("SL").get("/openapi.json").status_code
    if got != 403:
        bad.append(f"영업부 GET /openapi.json {got} (403 기대)")
    return res("G-19", "ID 바꿔치기 · 명세 노출", not bad, f"알림 다른 부서 6 · 의뢰 수정 W 없는 부서 5 · 명세 4 — 어긋남 {len(bad)}", bad)


# ── G-20 세션 ───────────────────────────────────────────────────────────
def _status(c, path: str = "/bas/items") -> int:
    return c.get(path).status_code


def check_sessions(lab: Lab) -> list[Result]:
    """요청마다 DB — 중지 · 잠금 · 비밀번호 초기화 · session_epoch · 로그아웃 쿠키 재사용 · 계정 삭제 → 401, 부서 변경은 다음 요청부터."""
    from lcomfine.db import conn

    out: list[Result] = []
    admin = lab.client("ADM")

    def target(tag: str, role: str = "PR"):
        lid, pw, uid = lab.make_user(role, tag)
        return lid, pw, uid, login(lid, pw)

    def upd(uid: int, lid: str, **kw):
        body = {"login_id": lid, "user_name": f"QA3 임시 {lid} (예시)", "role_code": kw.get("role", "PR"),
                "status": kw.get("status", "정상")}
        if "password" in kw:
            body["password"] = kw["password"]
        r = admin.post(f"/sys/users/{uid}", data=body)
        if r.status_code != 200:
            raise AssertionError(f"SYS-012 수정 {r.status_code} {r.text[:200]}")

    # 1. 중지(SYS-012 화면 경로) → 다음 요청 401 · 정상 복구해도 옛 세션은 401 · 다시 로그인하면 된다
    for st in ("중지", "잠금"):
        lid, pw, uid, c = target(f"st{'s' if st == '중지' else 'l'}")
        before = _status(c)
        upd(uid, lid, status=st)
        after, post = _status(c), c.post("/bas/items", data={}).status_code
        upd(uid, lid, status="정상")
        revived = _status(c)
        relog = _status(login(lid, pw))
        ok = (before, after, post, revived, relog) == (200, 401, 401, 401, 200)
        out.append(res("G-20", f"{st} → 401", ok,
                       f"전 {before} · {st} 뒤 GET {after} · POST {post} · 정상 복구 뒤 옛 세션 {revived} · 다시 로그인 {relog}"))
    # 2. 비활성(F-SYS-013)
    lid, pw, uid, c = target("deac")
    r = admin.post(f"/sys/users/{uid}/delete")
    got = _status(c)
    out.append(res("G-20", "비활성(F-SYS-013) → 401", r.status_code == 200 and got == 401, f"비활성 {r.status_code} · 다음 요청 {got}"))
    # 3. 부서 변경 — 끊기지 않고 다음 요청부터 새 부서 칸 (PR: BAS-030 W → SL: BAS-030 없음)
    lid, pw, uid, c = target("role")
    a = _status(c, "/bas/warehouses")
    upd(uid, lid, role="SL")
    b, home, w = _status(c, "/bas/warehouses"), _status(c, "/"), c.post("/bas/warehouses", data={}).status_code
    upd(uid, lid, role="PR")
    back = _status(c, "/bas/warehouses")
    ok = (a, b, home, w, back) == (200, 403, 200, 403, 200)
    out.append(res("G-20", "부서 변경 다음 요청 반영", ok,
                   f"PR 때 BAS-030 {a} · SL 로 바꾼 다음 요청 {b} · 같은 세션 메인 {home} · 쓰기 {w} · PR 복구 {back}"))
    # 4. 로그아웃 쿠키 재사용 · 다른 단말 세션은 그대로
    lid, pw, uid, c = target("lout")
    c2 = login(lid, pw)
    saved_cookie = dict(c.cookies)
    c.post("/logout", follow_redirects=False)
    replay = new_client()
    for k, v in saved_cookie.items():
        replay.cookies.set(k, v)
    got_r, got_c, other = _status(replay), _status(c), _status(c2)
    out.append(res("G-20", "로그아웃 쿠키 재사용 401", (got_r, got_c, other) == (401, 401, 200),
                   f"로그아웃 전 쿠키 재사용 {got_r} · 그 클라이언트 {got_c} · 같은 계정 다른 세션 {other}"))
    # 5. 비밀번호 초기화(SYS-012) → 옛 세션 401
    lid, pw, uid, c = target("pwd")
    upd(uid, lid, password=secrets.token_urlsafe(12))
    got = _status(c)
    out.append(res("G-20", "비밀번호 변경 뒤 옛 세션 401", got == 401, f"다음 요청 {got}"))
    # 6. session_epoch 를 SQL 로 올림 → 401
    lid, pw, uid, c = target("ep")
    conn.x("update sys_user set session_epoch = session_epoch + 1 where user_id = %s", (uid,))
    got = _status(c)
    out.append(res("G-20", "session_epoch 증가 → 401", got == 401, f"다음 요청 {got}"))
    # 7. 계정 삭제 → 401 (쿠키에 남은 login_id 로 통하지 않는다)
    lid, pw, uid, c = target("del")
    conn.x("delete from sys_user where user_id = %s", (uid,))
    got = _status(c)
    out.append(res("G-20", "계정 삭제 → 401", got == 401, f"다음 요청 {got}"))
    # 8. 쿠키의 부서 값을 믿지 않는다 — 쿠키는 서명돼 있어 고칠 수 없다(서명 깨진 쿠키는 미로그인)
    lid, pw, uid, c = target("tamp")
    forged = new_client()
    for k, v in dict(c.cookies).items():
        forged.cookies.set(k, v[:-4] + ("AAAA" if not v.endswith("AAAA") else "BBBB"))
    got = _status(forged)
    out.append(res("G-20", "서명 깨진 쿠키 → 401", got == 401, f"다음 요청 {got}"))
    return out


# ── G-21 비밀 ───────────────────────────────────────────────────────────
def _secret_values() -> list[str]:
    from tools import gate

    return [v for v in (os.environ.get(k) for k in gate.SECRET_KEYS) if v and len(v) >= 8]


def check_secrets() -> list[Result]:
    from tools import gate

    out = [Result("G-21", "비밀 스캔(추적 파일)", *gate.secret_scan())]
    try:
        out.append(Result("G-21", "시드 (예시)", *gate.seed_examples()))
    except Exception as exc:  # noqa: BLE001 — DB 가 없으면 그 사실이 실측이다
        out.append(Result("G-21", "시드 (예시)", FAIL, f"DB 조회 실패 — {type(exc).__name__}"))
    # 응답에 비밀 · 접속 문자열이 실리지 않는다 (인증 없이 열리는 것 + 오류 응답)
    values = _secret_values() + [get_settings().pg_dsn]
    c = new_client()
    bodies = {"/health": c.get("/health").text, "/login": c.get("/login").text, "/error?status=500": c.get("/error?status=500").text,
              "/nope (404)": c.get("/nope-qa3").text, "POST /login 실패": c.post("/login", data={"login_id": "x", "password": "y"}).text}
    leaks = [k for k, b in bodies.items() if any(v in b for v in values)]
    out.append(res("G-21", "응답에 비밀 값 없음", not leaks, f"인증 없는 응답 {len(bodies)} 중 비밀·접속 문자열 실린 것 {len(leaks)}", leaks))
    # 시드 계정 8 이 settings 의 비밀번호로 로그인된다(비밀번호는 코드·문서에 없다 — 위 스캔)
    pw = get_settings().seed_password or ""
    bad = []
    for lid in SEED_ACCOUNTS:
        r = new_client().post("/login", data={"login_id": lid, "password": pw}, follow_redirects=False)
        if r.status_code != 303:
            bad.append(f"{lid} {r.status_code}")
    out.append(res("G-21", "시드 계정 8 로그인(.env 비밀번호)", bool(pw) and not bad, f"{len(SEED_ACCOUNTS) - len(bad)}/{len(SEED_ACCOUNTS)} 로그인", bad))
    return out


def check_seed_idempotence() -> Result:
    from tools import gate

    return Result("G-21", "시드 멱등", *gate.seed_idempotence())


# ── G-22 감사 로그 ──────────────────────────────────────────────────────
def _max_change_id() -> int:
    from lcomfine.db import conn

    return conn.q1("select coalesce(max(id), 0) as m from sys_change_log")["m"]


def _changes_since(since: int, login_id: str) -> list[dict]:
    from lcomfine.db import conn

    return conn.q("select function_id, screen_id, target, method, path, login_id, role_code from sys_change_log "
                  "where id > %s and login_id = %s order by id", (since, login_id))


def check_change_log(lab: Lab) -> list[Result]:
    """쓰기 성공마다 sys_change_log 정확히 1행(기능 ID · 대상 · 사용자 · 경로) · 실패한 쓰기(422 · 403)는 0행."""
    from lcomfine.app import contracts, notify
    from lcomfine.db import conn

    p = lab.prefix
    adm_lid = lab.users["ADM"][0]
    adm = lab.client("ADM")
    cu = conn.q1("insert into customer (customer_code, customer_name, created_by) values (%s, %s, 'qa3') returning id",
                 (f"{p}c2", f"{p} 거래처 (예시)"))["id"]
    fg = conn.q1("insert into item (item_code, item_name, item_type, created_by) values (%s, %s, '제품', 'qa3') returning id",
                 (f"{p}fg2", f"{p} 제품 (예시)"))["id"]
    note_id = conn.q1("insert into notification (event_key, to_role, link, dedupe_key, message) values "
                      "('request_created', 'ADM', '/wo/requests', %s, %s) returning id",
                      (f"qa3:{p}:log", f"QA3 감사 로그 시험 {p} (예시)"))["id"]
    tmp_user = lab.make_user("SL", "logu")
    # F-HOME-031 (Ruling 24 · D-526): 자기 부서가 수신하는 이벤트를 전부 켠 채 저장하는 것만 200 — 생산부(PR)로 잰다.
    # 전산 관리자(ADM)는 수신 이벤트가 없어 어떤 본문도 422 다(아래 음성 표본).
    pr_events = [k for k, roles in notify.EVENTS.items() if "PR" in roles]
    ctx: dict = {}
    cases = [
        ("F-BAS-011", "ADM", "/bas/items", lambda: {"item_code": f"{p}it", "item_name": f"{p} 품목 (예시)", "item_type": "제품"}),
        ("F-BAS-012", "ADM", lambda: f"/bas/items/{ctx['item']}", lambda: {"item_code": f"{p}it", "item_name": f"{p} 품목 수정 (예시)",
                                                                              "item_type": "제품", "use_yn": "Y"}),
        ("F-BAS-013", "ADM", lambda: f"/bas/items/{ctx['item']}/delete", lambda: {}),
        ("F-BAS-021", "ADM", "/bas/customers", lambda: {"customer_code": f"{p}cu", "customer_name": f"{p} 거래처 (예시)"}),
        ("F-BAS-022", "ADM", f"/bas/customers/{cu}", lambda: {"customer_code": f"{p}c2", "customer_name": f"{p} 거래처 수정 (예시)",
                                                                "use_yn": "Y"}),
        ("F-BAS-031", "ADM", "/bas/warehouses", lambda: {"target": "warehouse", "warehouse_code": f"{p}wh", "warehouse_name": f"{p} 창고 (예시)"}),
        ("F-STD-021", "ADM", "/std/standards", lambda: {"standard_code": f"{p}sd", "standard_name": f"{p} 표준 (예시)"}),
        ("F-STD-022", "ADM", lambda: f"/std/standards/{ctx['std']}", lambda: {"standard_code": f"{p}sd",
                                                                                 "standard_name": f"{p} 표준 수정 (예시)", "use_yn": "Y"}),
        ("F-WO-011", "ADM", "/wo/requests", lambda: {"customer_id": str(cu), "item_id": str(fg), "qty": "5", "due_date": "2027-01-31"}),
        ("F-WO-012", "ADM", lambda: f"/wo/requests/{ctx['req']}", lambda: {"customer_id": str(cu), "item_id": str(fg), "qty": "6",
                                                                              "due_date": "2027-01-31", "status": "확인대기"}),
        ("F-HOME-022", "ADM", f"/home/notifications/{note_id}/handle", lambda: {}),
        ("F-HOME-031", "PR", "/home/notification-settings", lambda: {"events": pr_events, "enabled": pr_events}),
        ("F-SYS-011", "ADM", "/sys/users", lambda: {"login_id": f"{p}nu", "user_name": f"{p} 사용자 (예시)", "role_code": "SL",
                                                    "password": secrets.token_urlsafe(12)}),
        ("F-SYS-012", "ADM", f"/sys/users/{tmp_user[2]}", lambda: {"login_id": tmp_user[0], "user_name": "QA3 임시 수정 (예시)",
                                                                     "role_code": "SL", "status": "정상"}),
        ("F-SYS-031", "ADM", "/sys/codes", lambda: {"group_code": "설비", "code": f"{p}cd", "code_name": f"{p} 코드 (예시)", "sort_no": "1"}),
    ]
    bad, logs = [], []
    for fid, role, path, body in cases:
        path = path() if callable(path) else path
        fn = contracts.function(fid)
        since = _max_change_id()
        r = lab.client(role).post(path, data=body())
        rows = _changes_since(since, lab.users[role][0])
        if r.status_code != 200:
            bad.append(f"{fid} 쓰기 {r.status_code} {r.text[:120]}")
            continue
        data = r.json()
        if fid == "F-BAS-011":
            ctx["item"] = data.get("id")
        if fid == "F-STD-021":
            ctx["std"] = data.get("id")
        if fid == "F-WO-011":
            ctx["req"] = data.get("id")
        if len(rows) != 1:
            bad.append(f"{fid} 변경 로그 {len(rows)}행 (1 기대)")
            continue
        rw = rows[0]
        if (rw["function_id"], rw["screen_id"], rw["method"], rw["path"]) != (fid, fn.screen_id, "POST", path) or not rw["target"]:
            bad.append(f"{fid} 로그 값 {rw}")
        logs.append((lab.users[role][0], rw["target"]))
    mods = len({contracts.function(fid).screen_id.split("-")[0] for fid, *_ in cases})
    out = [res("G-22", "쓰기 성공 → sys_change_log 1행", not bad, f"쓰기 표본 {len(cases)} (모듈 {mods}) · 어긋남 {len(bad)}", bad)]
    # 실패한 쓰기는 로그 0행 — 빈 본문 422 · 칸 없는 부서 403 · F-HOME-031 의 범위 밖(ADM 은 수신 이벤트 없음) 422 ·
    # 부서 기본 끄기 422 · 메일 501(D-518). 응답 코드도 계약(function-list F-HOME-031 · D-526)과 맞아야 한다.
    bad2 = []
    neg = [
        ("ADM", "/bas/items", {}, (422,)),
        ("SL", "/bas/warehouses", {}, (403,)),
        ("ADM", "/wo/requests", {}, (422,)),
        ("ADM", "/home/notification-settings", {"events": ["material_received"], "enabled": ["material_received"]}, (422,)),
        ("PR", "/home/notification-settings", {"events": pr_events, "enabled": pr_events[1:]}, (422,)),
        ("PR", "/home/notification-settings", {"events": pr_events, "enabled": pr_events, "mail": "on"}, (501,)),
    ]
    for role, path, body, want in neg:
        lid = lab.users[role][0]
        since = _max_change_id()
        got = lab.client(role).post(path, data=body).status_code
        n = len(_changes_since(since, lid))
        if got not in want or n:
            bad2.append(f"{role} {path} → {got} (기대 {'/'.join(map(str, want))}) · 로그 {n}행")
    out.append(res("G-22", "실패한 쓰기(422·403·501) → 0행", not bad2, f"표본 {len(neg)} · 어긋남 {len(bad2)}", bad2))
    # 화면 조회 · 로그인 성공/실패 → sys_access_log
    bad3 = []
    lid = lab.users["PR"][0]

    def acc(kind: str, extra: str = "", args=()) -> int:
        return conn.q1(f"select count(*) as n from sys_access_log where login_id = %s and log_type = %s {extra}",
                       (lid, kind, *args))["n"]

    c = lab.client("PR")
    for path, sid in (("/bas/items", "BAS-010"), ("/wo/requests", "WO-010"), ("/mat/receipts", "MAT-010"), ("/lot/backward", "LOT-020")):
        before = acc("조회", "and screen_id = %s", (sid,))
        st = c.get(path).status_code
        after = acc("조회", "and screen_id = %s", (sid,))
        if st != 200 or after - before != 1:
            bad3.append(f"GET {path} {st} · 조회 로그 +{after - before}")
    ok_b, fail_b = acc("로그인", "and result = '성공'"), acc("로그인", "and result = '실패'")
    lab.fresh("PR")
    new_client().post("/login", data={"login_id": lid, "password": "wrong-" + secrets.token_hex(4)})
    ok_a, fail_a = acc("로그인", "and result = '성공'"), acc("로그인", "and result = '실패'")
    if (ok_a - ok_b, fail_a - fail_b) != (1, 1):
        bad3.append(f"로그인 성공 +{ok_a - ok_b} · 실패 +{fail_a - fail_b} (1·1 기대)")
    out.append(res("G-22", "화면 조회 · 로그인 → sys_access_log", not bad3, f"화면 4 · 로그인 2 · 어긋남 {len(bad3)}", bad3))
    # SYS-040 에서 보인다
    pages = {lid: adm.get("/sys/logs", params={"tab": "change", "login": lid}).text for lid in {lid for lid, _ in logs}}
    seen = [t for lid, t in logs if t in pages[lid]]
    out.append(res("G-22", "SYS-040 에서 조회", bool(logs) and len(seen) == len(logs),
                   f"방금 쓴 변경 로그 {len(logs)} 중 SYS-040 화면에 보인 것 {len(seen)}",
                   [f"{lid} {t}" for lid, t in logs if t not in pages[lid]]))
    # 정적 — 라우터 POST 마다 log_change 를 부르는가(같은 모듈의 도우미 한 단계까지)
    out.append(check_static_log_change())
    return out


def check_static_log_change() -> Result:
    bad, n = [], 0
    for m, meth, path, ep, ids in router_routes():
        if meth != "POST":
            continue
        n += 1
        mod = sys.modules[ep.__module__]
        src = inspect.getsource(ep)
        helpers = [getattr(mod, name) for name in set(re.findall(r"\b(_?\w+)\(", src)) if inspect.isfunction(getattr(mod, name, None))
                   and getattr(mod, name).__module__ == mod.__name__]
        text = src + "".join(inspect.getsource(h) for h in helpers)
        if "log_change(" not in text:
            bad.append(f"POST {path}")
    return res("G-22", "정적: POST 마다 log_change", not bad, f"POST {n} 중 log_change 없음 {len(bad)}", bad)


# ── 조용한 실패 (X) ──────────────────────────────────────────────────────
_DB_DOWN_SCRIPT = r"""
import json, os, sys, warnings
warnings.filterwarnings("ignore")
sys.path.insert(0, os.path.join(os.environ["QA3_ROOT"], "src"))
from fastapi.testclient import TestClient
from lcomfine.app.main import app
c = TestClient(app, raise_server_exceptions=False)
for k, v in json.loads(os.environ["QA3_COOKIE"]).items():
    c.cookies.set(k, v)
out = {}
def grab(name, r):
    out[name] = {"status": r.status_code, "body": r.text[:4000]}
grab("health", c.get("/health"))
grab("screen_json", c.get("/bas/items"))
grab("screen_html", c.get("/bas/items", headers={"accept": "text/html"}, follow_redirects=False))
grab("home_html", c.get("/", headers={"accept": "text/html"}, follow_redirects=False))
grab("write", c.post("/bas/items", data={"item_code": "x", "item_name": "y", "item_type": "제품"}))
grab("login_page", c.get("/login", headers={"accept": "text/html"}))
grab("login_post", c.post("/login", data={"login_id": "x", "password": "y"}))
import io
from openpyxl import Workbook
wb = Workbook(); ws = wb.active; ws.append(["a", "b"]); ws.append(["1", "2"]); buf = io.BytesIO(); wb.save(buf)
grab("col_upload", c.post("/col/uploads", data={"action": "apply", "period_from": "1985-01-01", "period_to": "1985-01-31"},
                          files={"file": ("q3 db-down (예시).xlsx", buf.getvalue(),
                                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}))
grab("col_screen", c.get("/col/collections"))
print("QA3JSON" + json.dumps(out, ensure_ascii=False))
"""


def db_down_probe(lab: Lab) -> dict:
    cookie = dict(lab.fresh("PR").cookies)
    env = dict(os.environ, LCOMFINE_PG_DSN=ABSENT_DB_DSN, QA3_ROOT=str(ROOT), QA3_COOKIE=json.dumps(cookie))
    p = subprocess.run([sys.executable, "-c", _DB_DOWN_SCRIPT], cwd=ROOT, env=env, capture_output=True, text=True, timeout=180)
    m = re.search(r"QA3JSON(.*)$", p.stdout, re.S)
    if not m:
        raise AssertionError(f"하위 프로세스 실패 rc={p.returncode}: {(p.stderr or p.stdout)[-400:]}")
    return json.loads(m.group(1))


def check_db_down(lab: Lab) -> list[Result]:
    """DB 를 끊으면(없는 DB 이름) /health 503 · 화면 · 쓰기 · 로그인도 503 — 멀쩡한 화면이나 「미수집」 으로 보이면 결함."""
    got = db_down_probe(lab)
    values = _secret_values() + [ABSENT_DB_DSN, "lcomfine_qa3_absent_db"]
    out = []
    h = got["health"]
    out.append(res("X-01", "DB 끊김 /health 503", h["status"] == 503 and not any(v in h["body"] for v in values),
                   f"/health {h['status']} · 본문에 접속 문자열 {'있음' if any(v in h['body'] for v in values) else '없음'}"))
    for key, label in (("screen_json", "화면 GET(JSON)"), ("screen_html", "화면 GET(브라우저)"), ("home_html", "메인 GET(브라우저)"),
                       ("write", "쓰기 POST"), ("login_page", "로그인 화면"), ("login_post", "로그인 POST"),
                       ("col_upload", "COL 업로드 POST(multipart)"), ("col_screen", "COL 화면 GET")):
        g = got[key]
        silent = g["status"] < 400 and key != "login_page"
        if key == "login_page":   # 퀵 로그인 패널이 DB 를 읽는다 — 200 이면 DB 를 읽지 않은 것이므로 「미수집」 류 거짓 값이 없어야 한다
            silent = g["status"] == 200 and "미수집" in g["body"]
        leak = any(v in g["body"] for v in values)
        out.append(res("X-02", f"DB 끊김 {label}", g["status"] == 503 and not leak and not silent,
                       f"{g['status']}{' · 미수집 표시' if '미수집' in g['body'] else ''}{' · 접속 문자열 노출' if leak else ''}"))
    return out


def malformed_cases(lab: Lab) -> list[tuple[str, str, str, dict, str]]:
    """(라벨, 부서, 경로, 본문, 기대) — 기대 '422' 는 422(또는 404 경로 키), 'reject' 는 2xx 가 아니어야 하고 500 도 아니어야 한다."""
    from lcomfine.db import conn

    cached = getattr(lab, "_malformed_cases", None)
    if cached is not None:  # 같은 Lab 에서 두 번 부르면(check_malformed 뒤 test_nul_…) 기준 행을 다시 넣지 않는다
        return cached
    p = lab.prefix
    cu = conn.q1("insert into customer (customer_code, customer_name, created_by) values (%s, %s, 'qa3') returning id",
                 (f"{p}c3", f"{p} 거래처 (예시)"))["id"]
    fg = conn.q1("insert into item (item_code, item_name, item_type, created_by) values (%s, %s, '제품', 'qa3') returning id",
                 (f"{p}fg3", f"{p} 제품 (예시)"))["id"]
    rm = conn.q1("insert into item (item_code, item_name, item_type, created_by) values (%s, %s, '원자재', 'qa3') returning id",
                 (f"{p}rm3", f"{p} 원자재 (예시)"))["id"]
    wh = conn.q1("insert into warehouse (warehouse_code, warehouse_name, created_by) values (%s, %s, 'qa3') returning id",
                 (f"{p}w3", f"{p} 창고 (예시)"))["id"]
    loc = conn.q1("insert into location (warehouse_id, location_code, location_name, created_by) values (%s, %s, %s, 'qa3') returning id",
                  (wh, f"{p}l3", f"{p} 위치 (예시)"))["id"]
    big = "가" * 100_000
    req = {"customer_id": str(cu), "item_id": str(fg), "qty": "5", "due_date": "2027-01-31"}
    rec = {"supplier_name": f"{p} 공급사 (예시)", "received_on": "2026-10-21", "item_id": [str(rm)], "qty": ["10"],
           "location_id": [str(loc)], "supplier_lot_no": [f"{p}S1"]}
    shp = {"customer_id": str(cu), "due_date": "2027-01-31", "stock_id": ["1"], "qty": ["1"]}
    lab._malformed_cases = [
        ("품목 코드 NUL", "ADM", "/bas/items", {"item_code": f"{p}n\x00ul", "item_name": "x (예시)", "item_type": "제품"}, "reject"),
        ("품목명 10만 자", "ADM", "/bas/items", {"item_code": f"{p}big", "item_name": big, "item_type": "제품"}, "no500"),
        ("거래처명 NUL", "ADM", "/bas/customers", {"customer_code": f"{p}cn", "customer_name": "a\x00b"}, "reject"),
        ("거래처 10만 자", "ADM", "/bas/customers", {"customer_code": f"{p}cb", "customer_name": big}, "no500"),
        ("의뢰 수량 음수", "SL", "/wo/requests", {**req, "qty": "-5"}, "reject"),
        ("의뢰 수량 1e30", "SL", "/wo/requests", {**req, "qty": "1e30"}, "reject"),
        ("의뢰 수량 NaN", "SL", "/wo/requests", {**req, "qty": "NaN"}, "reject"),
        ("의뢰 납기 99999년", "SL", "/wo/requests", {**req, "due_date": "99999-01-01"}, "reject"),
        ("의뢰 없는 거래처 FK", "SL", "/wo/requests", {**req, "customer_id": "999999999"}, "reject"),
        ("의뢰 거래처 id bigint 초과", "SL", "/wo/requests", {**req, "customer_id": "99999999999999999999"}, "reject"),
        ("의뢰 비고 NUL", "SL", "/wo/requests", {**req, "note": "a\x00b"}, "reject"),
        ("원자재 입고 수량 음수", "PR", "/mat/receipts", {**rec, "qty": ["-10"]}, "reject"),
        ("원자재 입고 없는 품목 FK", "PR", "/mat/receipts", {**rec, "item_id": ["999999999"]}, "reject"),
        ("원자재 입고 없는 위치 FK", "PR", "/mat/receipts", {**rec, "location_id": ["999999999"]}, "reject"),
        ("원자재 입고 공급사 NUL", "PR", "/mat/receipts", {**rec, "supplier_name": "a\x00b"}, "reject"),
        ("원자재 입고일 99999년", "PR", "/mat/receipts", {**rec, "received_on": "99999-12-31"}, "reject"),
        ("원자재 입고 수량 10^40", "PR", "/mat/receipts", {**rec, "qty": ["1" + "0" * 40]}, "reject"),
        ("출하 지시 없는 재고 FK", "SH", "/shp/orders", {**shp, "stock_id": ["999999999"]}, "reject"),
        ("출하 지시 수량 음수", "SH", "/shp/orders", {**shp, "qty": ["-1"]}, "reject"),
        ("출하 지시 납기 99999년", "SH", "/shp/orders", {**shp, "due_date": "99999-01-01"}, "reject"),
        ("재고 조정 없는 LOT 번호", "PR", "/inv/adjustments", {"lot_kind": "원자재", "lot_no": "없는LOT", "location_id": str(loc),
                                                         "actual_qty": "-3", "reason": "QA3"}, "reject"),
        ("생산일보 지시 id bigint 초과", "PR", "/prd/daily", {"work_order_id": "99999999999999999999", "work_date": "2026-10-21",
                                                         "produced_qty": "1", "good_qty": "1", "loss_qty": "0"}, "reject"),
        ("공통코드 NUL", "ADM", "/sys/codes", {"group_code": "설비", "code": "a\x00b", "code_name": "x (예시)", "sort_no": "1"}, "reject"),
        ("공통코드 정렬 bigint 초과", "ADM", "/sys/codes", {"group_code": "설비", "code": f"{p}c", "code_name": "x (예시)",
                                                       "sort_no": "99999999999999999999"}, "reject"),
        ("경로 id bigint 초과", "ADM", "/bas/items/99999999999999999999", {"item_code": "x", "item_name": "y", "item_type": "제품"}, "reject"),
        ("경로 id NUL", "ADM", "/bas/items/1%00", {"item_code": "x", "item_name": "y", "item_type": "제품"}, "reject"),
    ]
    return lab._malformed_cases


def check_malformed(lab: Lab) -> list[Result]:
    """이상한 입력 → 422(또는 경로 키 404). 500 은 결함 · 음수/없는 FK 가 저장되면(2xx) 결함 · 10만 자 저장은 계약에 길이 제한이 없어 참고."""
    out, bad, notes = [], [], []
    cases = malformed_cases(lab)
    for label, role, path, body, want in cases:
        r = lab.client(role).post(path, data=body)
        st = r.status_code
        if st >= 500:
            bad.append(f"{label} {st} {r.text[:120]}")
        elif want == "reject" and st < 400:
            bad.append(f"{label} {st} 저장됨 (422 기대)")
        elif want == "reject" and st not in (404, 422):
            bad.append(f"{label} {st} (422 기대)")
        elif want == "no500" and st < 400:
            notes.append(f"{label} {st} 저장됨")
    out.append(res("X-03", "이상한 입력 → 422", not bad, f"POST {len(cases)} (API 9) · 어긋남 {len(bad)}", bad))
    # 조회 인자에 NUL · 형식이 틀린 날짜 — 422 (500 이 아니다)
    bad_q = []
    c = lab.client("ADM")
    for path, params in (("/bas/items", {"q": "a\x00b"}), ("/wo/requests", {"q": "a\x00b"}), ("/wo/requests", {"due_from": "99999-01-01"}),
                         ("/lot/backward", {"lot_no": "a\x00b"}), ("/lot/forward", {"lot_no": "a\x00b"}), ("/shp/history", {"q": "a\x00b"}),
                         ("/mat/history", {"date_from": "2026-13-45"}), ("/sys/logs", {"tab": "change", "login": "a\x00b"})):
        st = c.get(path, params=params).status_code
        if st >= 500:
            bad_q.append(f"GET {path} {params} {st}")
    out.append(res("X-04", "조회 인자 NUL · 틀린 날짜 → 500 아님", not bad_q, f"GET 8 · 500 {len(bad_q)}", bad_q))
    if notes:
        print("참고 (X-03): " + " / ".join(notes))
    return out


# ── G-24 (정적) ─────────────────────────────────────────────────────────
def check_docker() -> list[Result]:
    out = []
    try:
        p = subprocess.run(["docker", "compose", "config", "-q"], cwd=ROOT, capture_output=True, text=True, timeout=120)
        rc, tail = p.returncode, (p.stderr or p.stdout).strip()[-160:]
    except FileNotFoundError:
        rc, tail = 127, "docker 없음"
    out.append(res("G-24", "docker compose config -q", rc == 0, f"rc={rc}{' ' + tail if rc else ''}"))
    dockerfile = (ROOT / "Dockerfile").read_text(encoding="utf-8")
    pyproject = (ROOT / "pyproject.toml").read_text(encoding="utf-8")
    lock = (ROOT / "uv.lock").read_text(encoding="utf-8")
    syncs = "uv sync --frozen --no-dev" in dockerfile
    has = 'openpyxl' in pyproject and 'name = "openpyxl"' in lock
    out.append(res("G-24", "Dockerfile 이 openpyxl 을 설치", syncs and has,
                   f"uv sync --frozen {'있음' if syncs else '없음'} · pyproject·uv.lock 의 openpyxl {'있음' if has else '없음'}"))
    ignored = [ln.strip() for ln in (ROOT / ".dockerignore").read_text(encoding="utf-8").splitlines() if ln.strip() and not ln.startswith("#")]
    missing = []
    for ln in dockerfile.splitlines():
        m = re.match(r"\s*COPY\s+(?!--from)(.+)$", ln)
        if not m:
            continue
        parts = m.group(1).split()
        for src in parts[:-1]:
            if not (ROOT / src).exists():
                missing.append(f"COPY {src} (파일 없음)")
            elif any(src.rstrip("/") == ig.rstrip("/") or src.startswith(ig.rstrip("/") + "/") for ig in ignored if not ig.startswith("!")):
                missing.append(f"COPY {src} (.dockerignore)")
    out.append(res("G-24", "Dockerfile COPY 원본이 있다", not missing, f"COPY 원본 없음 {len(missing)} — 이대로면 `docker compose build` 가 실패한다" if missing else "COPY 원본 전부 있음", missing))
    return out


# ── E2E 캡처 ────────────────────────────────────────────────────────────
E2E_TAG = "q3e2e"                                  # 브라우저 한 바퀴가 만든 회차 · 매핑의 접두 — 끝에 지운다
E2E_PERIOD = ("1985-02-01", "1985-02-28")          # 다른 테스트와 겹치지 않는 해
E2E_W2_AC = ("/col/uploads", "/col/upload-history", "/col/customer-map",
             f"/col/collections?date_from={E2E_PERIOD[0]}&date_to={E2E_PERIOD[1]}", f"/col/receivables?asof={E2E_PERIOD[1]}",
             f"/col/aging?asof={E2E_PERIOD[1]}", "/col/shipment-vs-collection?ym_from=1984-09&ym_to=1985-02")


def _e2e_xlsx(tmp_dir: Path) -> tuple[Path, str, str]:
    """지금 열 매핑의 열 이름으로 (예시) 수금현황 xlsx 를 임시 폴더에 쓴다 — (경로, 매핑할 ERP 코드, 미매핑 ERP 코드)."""
    cmap, fields, hdrs = col_map()
    code_e, code_u = f"{E2E_TAG}-E", f"{E2E_TAG}-U"
    rows = [col_row(fields, customer_code=code_e, txn_date="1985-02-05", sales_amount="1000", collected_amount="0", method="PM (예시)"),
            col_row(fields, customer_code=code_e, txn_date="1985-02-20", collected_amount="600", method="PM (예시)", bill_due_date="1985-03-15"),
            col_row(fields, customer_code=code_u, txn_date="1985-02-21", sales_amount="500")]
    path = tmp_dir / f"{E2E_TAG} 수금현황 (예시).xlsx"
    path.write_bytes(xlsx_bytes(hdrs, rows))
    return path, code_e, code_u


def _e2e_cleanup() -> None:
    from lcomfine.db import conn

    with conn.tx() as cur:
        cleanup_col(cur, E2E_TAG + "%")


def e2e(port: int = 8020) -> int:
    """앱을 port 에 띄우고(하위 프로세스) Playwright 로 1280 · 390 캡처 + 가로 스크롤 0 확인 → outputs/e2e/.
    W1 화면 5 + W2: 회계부(ac)가 COL-010 에서 (예시) 엑셀을 **실제로 올리고**(1280 = 1회차 · 390 = 2회차 교체) → COL-020 이력 →
    COL-030 매핑 등록(1280 에서 한 번) → COL-040 · 050 · 060 · 070, 영업부(sl)가 TRK-030 수금 칸. 만든 회차 · 매핑은 끝에 지운다."""
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("E2E  Playwright 없음 — `uv run --offline --with playwright python tools/check_security.py --e2e` (전역 설치 안 함)")
        return 1
    import tempfile
    import urllib.request

    from lcomfine.db import conn

    E2E_DIR.mkdir(parents=True, exist_ok=True)
    tmp_dir = Path(tempfile.mkdtemp(prefix="qa3-e2e-"))
    xlsx_path, code_e, _ = _e2e_xlsx(tmp_dir)
    cust = conn.q1("select id, customer_code from customer order by id limit 1")       # TRK-030 · 매핑 대상 (시드 거래처)
    if cust is None:
        print("E2E  거래처가 없다 — 시드를 먼저 돌린다")
        return 1
    server = subprocess.Popen([sys.executable, "-m", "uvicorn", "lcomfine.app.main:app", "--host", "127.0.0.1", "--port", str(port)],
                              cwd=ROOT, env=dict(os.environ, PYTHONPATH=str(ROOT / "src")), stdout=subprocess.DEVNULL,
                              stderr=subprocess.DEVNULL)
    base = f"http://127.0.0.1:{port}"
    results, fails = [], 0
    try:
        for _ in range(60):
            try:
                if urllib.request.urlopen(base + "/health", timeout=2).status == 200:
                    break
            except Exception:  # noqa: BLE001 — 기동 대기
                time.sleep(0.5)
        else:
            print("E2E  앱이 뜨지 않았다")
            return 1
        pw = get_settings().seed_password or ""
        plan = ([("admin", p) for p in E2E_PATHS] + [("sl", "/wo/requests"), ("sh", "/shp/orders"), ("pr", "/lot/backward")]
                + [("ac", p) for p in E2E_W2_AC] + [("sl", f"/trk/customers?customer={cust['id']}")])
        with sync_playwright() as play:
            browser = play.chromium.launch()
            for width in E2E_WIDTHS:
                for account in sorted({a for a, _ in plan}, key=SEED_ACCOUNTS.index):
                    ctx = browser.new_context(viewport={"width": width, "height": 900})
                    page = ctx.new_page()
                    page.on("dialog", lambda d: d.accept())          # data-confirm (반영 · 삭제) 는 승인
                    page.goto(base + "/login")
                    page.fill("input[name=login_id]", account)
                    page.fill("input[name=password]", pw)
                    page.locator("#login-form button.btn-primary").click()
                    page.wait_for_load_state("networkidle")
                    for acc, path in plan:
                        if acc != account:
                            continue
                        action, want_text = "", ""
                        resp = page.goto(base + path)
                        page.wait_for_load_state("networkidle")
                        if account == "ac" and path == "/col/uploads":                       # COL-010 업로드 (apply)
                            action, want_text = "upload-apply", "반영했습니다"
                            page.fill("input[name=period_from]", E2E_PERIOD[0])
                            page.fill("input[name=period_to]", E2E_PERIOD[1])
                            page.set_input_files("input[name=file]", str(xlsx_path))
                            page.click("button[name=action][value=apply]")
                            page.wait_for_load_state("networkidle")
                        elif account == "ac" and path == "/col/customer-map" and width == E2E_WIDTHS[0]:   # COL-030 매핑 등록
                            action, want_text = "map-create", "매핑을 등록했습니다"
                            form = page.locator("form[method=post]").filter(has=page.locator("input[name=erp_code]"))
                            form.locator("input[name=erp_code]").fill(code_e)
                            form.locator("input[name=erp_name]").fill(f"{E2E_TAG} ERP 거래처 (예시)")
                            form.locator("select[name=customer_id]").select_option(str(cust["id"]))
                            form.locator("button[type=submit]").click()
                            page.wait_for_load_state("networkidle")
                        sw, iw = page.evaluate("[document.documentElement.scrollWidth, window.innerWidth]")
                        clean = path.split("?")[0]
                        name = f"{account}{clean.replace('/', '_') if clean != '/' else '_home'}{'-' + action if action else ''}-{width}.png"
                        page.screenshot(path=str(E2E_DIR / name), full_page=True)
                        body = page.content()
                        ok = (resp is not None and resp.status == 200 and sw <= iw and "/login" not in page.url
                              and (not want_text or want_text in body))
                        fails += 0 if ok else 1
                        results.append({"account": account, "path": path, "action": action or None, "width": width,
                                        "status": resp.status if resp else None, "url": page.url.replace(base, ""),
                                        "scrollWidth": sw, "innerWidth": iw, "ok": ok, "file": name,
                                        **({"flash": want_text, "flash_seen": want_text in body} if want_text else {})})
                        print(f"E2E  {account} {path}{' [' + action + ']' if action else ''} @{width}  {'PASS' if ok else 'FAIL'}  "
                              f"status {resp.status if resp else '?'} · scrollWidth {sw} ≤ innerWidth {iw}"
                              f"{' · ' + want_text + (' 보임' if want_text in body else ' 안 보임') if want_text else ''} · {name}")
                    ctx.close()
            browser.close()
    finally:
        server.terminate()
        server.wait(timeout=10)
        _e2e_cleanup()
        left = conn.q1("select (select count(*) from erp_upload_batch where source_file ilike %s) as b, "
                       "(select count(*) from erp_customer_map where erp_code ilike %s) as m", (E2E_TAG + "%", E2E_TAG + "%"))
        print(f"E2E  정리 — 남은 {E2E_TAG} 회차 {left['b']} · 매핑 {left['m']}")
    (E2E_DIR / "result.json").write_text(json.dumps(results, ensure_ascii=False, indent=1), encoding="utf-8")
    return fails


# ── W2 수금 (COL) ───────────────────────────────────────────────────────
# 회차는 1985년 기준 기간(개발3 테스트 1990~1993 · QA2 는 자기 접두와 겹치지 않는다) · 파일 이름 · 거래처 · ERP 코드는 전부 `q3<난수>` 접두.
# 열 매핑(`erp_column_map`)은 전역 설정이라 **바꾸지 않는다** — 지금 저장된 매핑의 열 이름으로 (예시) xlsx 를 만든다.
COL_PERIOD = (date(1985, 1, 1), date(1985, 1, 31))          # G-22 표본
COL_PERIOD_ABUSE = (date(1985, 3, 1), date(1985, 3, 31))    # X-05~X-07
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_COL_WRITE_PATTERNS = (r"\bopen\(", r"\.write\(", r"\bPath\(", r"\bos\.", r"\bshutil\b", r"\btempfile\b", r"NamedTemporaryFile",
                       r"\.save\((?!buf\))")


def cleanup_col(cur, like: str) -> None:
    """이 Lab 이 만든 수금 회차 · 행 · 알림 · 거래처 매핑을 지운다(사용자 login_id · 파일 이름 · ERP 코드가 접두로 시작하는 것)."""
    from lcomfine.app import erp_upload

    cur.execute("select id from erp_upload_batch where created_by ilike %s or source_file ilike %s", (like, like))
    ids = [r["id"] for r in cur.fetchall()]
    if ids:
        cur.execute("delete from notification where link = any(%s)", ([erp_upload.error_link(i) for i in ids],))
        cur.execute("delete from erp_upload_row where batch_id = any(%s)", (ids,))
        cur.execute("delete from erp_upload_batch where id = any(%s)", (ids,))     # replaced_by FK 는 문장 끝에 검사된다
    cur.execute("delete from erp_customer_map where erp_code ilike %s or customer_id in (select id from customer where customer_code ilike %s)",
                (like, like))


def col_map() -> tuple[dict[str, str], list[str], list[str]]:
    """지금 저장된 열 매핑 — (항목→열 이름, 매핑된 항목 순서, 그 열 이름). 필수 매핑이 없으면 UploadError 로 검사가 FAIL 한다."""
    from lcomfine.app import erp_upload

    cmap = erp_upload.column_map()
    erp_upload.check_map(cmap)
    fields = [f for f in erp_upload.FIELDS if cmap.get(f)]
    return cmap, fields, [cmap[f] for f in fields]


def col_row(fields: list[str], **vals) -> list:
    return [vals.get(f) for f in fields]


def xlsx_bytes(headers: list, rows: list[list], *, text_formulas: bool = True) -> bytes:
    """(예시) xlsx — 저장소에 바이너리를 두지 않는다. `=` 로 시작하는 글자는 기본으로 **글자 셀**(ERP 가 값으로 내려준 꼴)로 쓴다.
    text_formulas=False 면 openpyxl 기본대로 수식 셀이 된다(캐시 값 없음)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    if text_formulas:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.data_type = "s"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def col_post(client, period: tuple[date, date], data: bytes, *, action: str = "apply", name: str = "q3 erp (예시).xlsx",
             extra: dict | None = None, headers: dict | None = None):
    form = {"action": action, "period_from": period[0].isoformat(), "period_to": period[1].isoformat(), **(extra or {})}
    return client.post("/col/uploads", data=form, files={"file": (name, data, XLSX_MIME)}, headers=headers or {})


def _counts() -> tuple[int, int, int]:
    from lcomfine.db import conn

    r = conn.q1("select (select count(*) from erp_upload_batch) as b, (select count(*) from erp_upload_row) as r, "
                "(select count(*) from notification) as n")
    return r["b"], r["r"], r["n"]


def check_col_access(lab: Lab) -> list[Result]:
    """COL 화면 7 × 부서 8 — 칸 있으면 200 · 없으면 403 · R 칸에는 쓰기 폼이 없다. COL 쓰기 4 × 부서 8 — W 아닌 부서는 403."""
    from lcomfine.app import contracts, nav, rbac

    cells = {(sid, rc): lv for sid, rc, lv in contracts.permission_cells()}
    screens = [sc for sc in nav.SCREENS if sc.screen_id.startswith("COL-")]
    fns = [f for f in write_functions_routed() if f.id.startswith("F-COL-")]
    bad, n200, n403, forms = [], 0, 0, 0
    for r in rbac.roles():
        c = lab.client(r.code)
        for sc in screens:
            lv = cells.get((sc.screen_id, r.code))
            want = 200 if lv else 403
            got = c.get(sc.path).status_code
            if want == 200:
                n200 += 1
            else:
                n403 += 1
            if got != want:
                bad.append(f"{sc.screen_id}×{r.code} {got} ({want} 기대)")
            if lv == "R" and sc.screen_id in ("COL-010", "COL-030"):
                forms += 1
                html = c.get(sc.path, headers={"accept": "text/html"}).text
                if 'enctype="multipart/form-data"' in html or 'name="action" value="map"' in html or 'name="erp_code"' in html:
                    bad.append(f"{sc.screen_id}×{r.code} R 칸인데 쓰기 폼이 보인다")
    out = [res("G-19", "COL 화면 7 × 부서 8 (200/403 · R 칸 쓰기 폼 없음)", not bad,
               f"화면 {len(screens)} × 부서 {len(rbac.roles())} — 200 기대 {n200} · 403 기대 {n403} · R 칸 폼 검사 {forms} · 어긋남 {len(bad)}", bad)]
    bad2, n403w, nw = [], 0, 0
    for f in fns:
        for r in rbac.roles():
            got = lab.client(r.code).post(concrete(f.path), data={}).status_code
            if r.code not in f.roles_w:
                n403w += 1
                if got != 403:
                    bad2.append(f"{f.id}×{r.code} {got} (403 기대)")
            else:
                nw += 1
                if got in (401, 403) or got >= 500:
                    bad2.append(f"{f.id}×{r.code} {got} (W 칸 — 403/401/5xx 아님 기대)")
    out.append(res("G-19", "COL 쓰기 4 × 부서 8 (R·없음 403)", not bad2 and len(fns) == 4,
                   f"쓰기 기능 {len(fns)} × 부서 8 — 403 기대 {n403w} · 통과 기대 {nw} (빈 본문 → 422/404) · 어긋남 {len(bad2)}", bad2))
    return out


def check_col_sessions(lab: Lab) -> Result:
    """회계부 세션 — 비활성(F-SYS-013) 다음 요청부터 COL GET · POST 401 · 로그아웃 뒤 같은 쿠키 COL 401."""
    admin = lab.client("ADM")
    lid, pw, uid = lab.make_user("AC", "cac")
    c = login(lid, pw)
    a = c.get("/col/uploads").status_code
    deact = admin.post(f"/sys/users/{uid}/delete").status_code
    b, bp = c.get("/col/collections").status_code, c.post("/col/customer-map", data={}).status_code
    lid2, pw2, _ = lab.make_user("AC", "cac2")
    c2 = login(lid2, pw2)
    a2 = c2.get("/col/customer-map").status_code
    c2.post("/logout", follow_redirects=False)
    d, dp = c2.get("/col/uploads").status_code, c2.post("/col/uploads", data={"action": "preview"}).status_code
    got = (a, deact, b, bp, a2, d, dp)
    return res("G-20", "COL 비활성 · 로그아웃 → 401", got == (200, 200, 401, 401, 200, 401, 401),
               f"AC 전 {a} · 비활성 {deact} → COL-040 GET {b} · 매핑 POST {bp} / 다른 AC 전 {a2} · 로그아웃 → COL-010 GET {d} · 업로드 POST {dp}")


def check_col_change_log(lab: Lab) -> list[Result]:
    """COL 쓰기 4(업로드 apply · 매핑 등록 · 수정 · 삭제) + 열 매핑 저장 → 쓰기 성공마다 sys_change_log 1행 · preview · 422 · 403 · 404 는 0행."""
    from lcomfine.app import contracts, erp_upload
    from lcomfine.db import conn

    p = lab.prefix
    ac, ac_lid = lab.client("AC"), lab.users["AC"][0]
    cmap, fields, hdrs = col_map()
    cu = conn.q1("insert into customer (customer_code, customer_name, created_by) values (%s, %s, 'qa3') returning id",
                 (f"{p}cc", f"{p} 거래처 (예시)"))["id"]
    code_e, code_x = f"{p}E", f"{p}X"
    rows1 = [col_row(fields, customer_code=code_e, txn_date="1985-01-05", sales_amount="1000", collected_amount="0", method="PM (예시)"),
             col_row(fields, customer_code=code_e, txn_date="1985-01-20", collected_amount="600", method="PM (예시)"),
             col_row(fields, customer_code=code_e, txn_date="1985/13/01", sales_amount="1")]          # 날짜 형식 오류
    rows2 = rows1 + [col_row(fields, customer_code=code_e, txn_date="1985-01-25", collected_amount="100")]   # 날짜 오류 행은 그대로
    ctx: dict = {}
    map_before = erp_upload.column_map()

    def c_map():
        return ac.post("/col/uploads", data={"action": "map", **{f"map_{f}": c for f, c in map_before.items()}})

    def c_apply(rows, name):
        return col_post(ac, COL_PERIOD, xlsx_bytes(hdrs, rows), name=name)

    cases = [
        ("열 매핑 저장(같은 값)", "F-COL-011", "/col/uploads", c_map),
        ("업로드 apply 1회차", "F-COL-011", "/col/uploads", lambda: c_apply(rows1, f"{p} 1회차 (예시).xlsx")),
        ("업로드 apply 2회차(교체)", "F-COL-011", "/col/uploads", lambda: c_apply(rows2, f"{p} 2회차 (예시).xlsx")),
        ("매핑 등록(참조 있음)", "F-COL-031", "/col/customer-map",
         lambda: ac.post("/col/customer-map", data={"erp_code": code_e, "erp_name": f"{p} ERP (예시)", "customer_id": str(cu)})),
        ("매핑 등록(참조 없음)", "F-COL-031", "/col/customer-map",
         lambda: ac.post("/col/customer-map", data={"erp_code": code_x, "customer_id": str(cu)})),
        ("매핑 수정", "F-COL-032", lambda: f"/col/customer-map/{ctx['map_e']}",
         lambda: ac.post(f"/col/customer-map/{ctx['map_e']}", data={"erp_name": f"{p} ERP 수정 (예시)", "customer_id": str(cu), "use_yn": "Y"})),
        ("매핑 삭제(참조 → 사용 중지)", "F-COL-033", lambda: f"/col/customer-map/{ctx['map_e']}/delete",
         lambda: ac.post(f"/col/customer-map/{ctx['map_e']}/delete")),
        ("매핑 삭제(참조 없음 → 삭제)", "F-COL-033", lambda: f"/col/customer-map/{ctx['map_x']}/delete",
         lambda: ac.post(f"/col/customer-map/{ctx['map_x']}/delete")),
    ]
    bad, logs, notes = [], [], []
    for label, fid, path, call in cases:
        path = path() if callable(path) else path
        fn = contracts.function(fid)
        since = _max_change_id()
        r = call()
        rows = _changes_since(since, ac_lid)
        if r.status_code != 200:
            bad.append(f"{label} {r.status_code} {r.text[:120]}")
            continue
        data = r.json()
        if label == "업로드 apply 1회차":
            ctx["b1"] = data["id"]
            if (data["rows"], data["errors"], data["round_no"], data["replaced"]) != (3, 3, 1, None):
                bad.append(f"{label} 응답 {data}")
            notes_ac = conn.q("select to_role, handled_at from notification where link = %s", (erp_upload.error_link(data["id"]),))
            if [n["to_role"] for n in notes_ac] != ["AC"]:
                bad.append(f"{label} erp_upload_error 알림 {notes_ac} (AC 1행 기대)")
        if label == "업로드 apply 2회차(교체)":
            ctx["b2"] = data["id"]
            b1 = conn.q1("select status, replaced_by from erp_upload_batch where id = %s", (ctx["b1"],))
            if (data["rows"], data["errors"], data["round_no"]) != (4, 4, 2) or b1["status"] != "교체" or b1["replaced_by"] != data["id"]:
                bad.append(f"{label} 교체 안 됨 — 응답 {data} · 1회차 {b1}")
        if label == "매핑 등록(참조 있음)":
            ctx["map_e"] = data["id"]
            ec = conn.q1("select error_count from erp_upload_batch where id = %s", (ctx["b2"],))["error_count"]
            if data["remapped"] < 1 or ec != 1:
                bad.append(f"{label} 재해석 {data['remapped']} · 2회차 error_count {ec} (1 기대 — 날짜 오류 행만)")
        if label == "매핑 등록(참조 없음)":
            ctx["map_x"] = data["id"]
        if label == "매핑 삭제(참조 → 사용 중지)":
            m = conn.q1("select use_yn from erp_customer_map where id = %s", (ctx["map_e"],))
            if not data.get("disabled") or not m or m["use_yn"] != "N":
                bad.append(f"{label} 사용 중지 아님 — 응답 {data} · 행 {m}")
        if label == "매핑 삭제(참조 없음 → 삭제)":
            if data.get("disabled") or conn.q1("select 1 from erp_customer_map where id = %s", (ctx["map_x"],)):
                bad.append(f"{label} 지워지지 않음 — 응답 {data}")
        if label == "열 매핑 저장(같은 값)" and erp_upload.column_map() != map_before:
            bad.append("열 매핑 저장(같은 값) 이 매핑을 바꿨다")
        if len(rows) != 1:
            bad.append(f"{label} 변경 로그 {len(rows)}행 (1 기대)")
            continue
        rw = rows[0]
        if (rw["function_id"], rw["screen_id"], rw["method"], rw["path"]) != (fid, fn.screen_id, "POST", path) or not rw["target"]:
            bad.append(f"{label} 로그 값 {rw}")
        logs.append(rw["target"])
    out = [res("G-22", "COL 쓰기 성공 → sys_change_log 1행", not bad, f"쓰기 표본 {len(cases)} (기능 4 · 업로드 apply 2 · map 1) · 어긋남 {len(bad)}", bad)]
    # 쓰지 않는 요청 — preview 200 · 422 · 403 · 404 → 로그 0행 · 회차/행/알림 수 불변
    bad2 = []
    sl = lab.client("SL")
    neg = [
        ("AC preview", lambda: col_post(ac, COL_PERIOD, xlsx_bytes(hdrs, rows2), action="preview"), (200,)),
        ("AC preview(HTML)", lambda: col_post(ac, COL_PERIOD, xlsx_bytes(hdrs, rows2), action="preview", headers={"accept": "text/html"}), (200,)),
        ("AC apply 파일 없음", lambda: ac.post("/col/uploads", data={"action": "apply", "period_from": "1985-01-01", "period_to": "1985-01-31"}), (422,)),
        ("AC apply 텍스트 파일", lambda: col_post(ac, COL_PERIOD, b"not,an,xlsx\n"), (422,)),
        ("AC apply 헤더만", lambda: col_post(ac, COL_PERIOD, xlsx_bytes(hdrs, [])), (422,)),
        ("AC map 필수 빠짐", lambda: ac.post("/col/uploads", data={"action": "map", "map_txn_date": "x"}), (422,)),
        ("AC 매핑 중복(사용 중지된 코드)", lambda: ac.post("/col/customer-map", data={"erp_code": code_e, "customer_id": str(cu)}), (422,)),
        ("AC 매핑 수정 없는 id", lambda: ac.post("/col/customer-map/999999999", data={"customer_id": str(cu)}), (404,)),
        ("SL(R) apply", lambda: col_post(sl, COL_PERIOD, xlsx_bytes(hdrs, rows2)), (403,)),
        ("SL(R) 매핑 등록", lambda: sl.post("/col/customer-map", data={"erp_code": f"{p}SL", "customer_id": str(cu)}), (403,)),
    ]
    for label, call, want in neg:
        role_lid = lab.users["SL"][0] if label.startswith("SL") else ac_lid
        since, before = _max_change_id(), _counts()
        got = call().status_code
        n, after = len(_changes_since(since, role_lid)), _counts()
        if got not in want or n or after != before:
            bad2.append(f"{label} → {got} (기대 {'/'.join(map(str, want))}) · 로그 {n}행 · 회차/행/알림 {before}→{after}")
    if erp_upload.column_map() != map_before:
        bad2.append("음성 표본이 열 매핑을 바꿨다")
    out.append(res("G-22", "COL preview · 422 · 403 · 404 → 0행 · 회차/행/알림 불변", not bad2, f"표본 {len(neg)} · 어긋남 {len(bad2)}", bad2))
    page = lab.client("ADM").get("/sys/logs", params={"tab": "change", "login": ac_lid}).text
    seen = [t for t in logs if t in page]
    out.append(res("G-22", "COL 변경 로그 SYS-040 에서 조회", bool(logs) and len(seen) == len(logs),
                   f"변경 로그 {len(logs)} 중 SYS-040 에 보인 것 {len(seen)}", [t for t in logs if t not in page]))
    return out


def _zip_not_xlsx() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("a.txt", "q3 (예시)")
    return buf.getvalue()


def check_col_upload_abuse(lab: Lab) -> list[Result]:
    """업로드 보안 — 엑셀 아닌 파일 · 빈 파일 · 20MB 초과 · 헤더 결함 → 422(500 0) · 셀의 수식/스크립트/긴 글자는 값으로만 저장되고
    화면에 이스케이프돼 나온다 · 내려받는 엑셀(COL-040)에 수식으로 되살아나지 않는다 · 파일 이름의 경로는 글자로만 저장된다."""
    from lcomfine.db import conn
    from openpyxl import load_workbook

    cached = getattr(lab, "_col_abuse", None)
    if cached is not None:        # 같은 Lab 에서 여러 테스트가 부른다 — 거래처 · 매핑 · 회차를 두 번 만들지 않는다
        return cached
    lab._col_abuse = out = _check_col_upload_abuse(lab, conn, load_workbook)
    return out


def _check_col_upload_abuse(lab: Lab, conn, load_workbook) -> list[Result]:
    p = lab.prefix
    ac = lab.client("AC")
    P = COL_PERIOD_ABUSE
    cmap, fields, hdrs = col_map()
    good = [col_row(fields, customer_code=f"{p}M", txn_date="1985-03-05", collected_amount="1")]
    out: list[Result] = []
    # 1. 파일 수준
    files = [
        ("텍스트(csv)", b"a,b\n1,2\n"), ("HTML", b"<html><body><script>alert(1)</script></body></html>"),
        ("zip(엑셀 아님)", _zip_not_xlsx()), ("0 바이트", b""), ("PNG 머리", b"\x89PNG\r\n\x1a\n" + b"\x00" * 64),
        ("20MB+1", b"\x00" * (erp_upload_max_bytes() + 1)),
        ("헤더만", xlsx_bytes(hdrs, [])), ("열 이름 중복", xlsx_bytes(hdrs + [hdrs[0]], [good[0] + ["x"]])),
        ("1행(열 이름) 비어 있음", xlsx_bytes([None] * len(hdrs), good)), ("__row_no 열", xlsx_bytes(hdrs + ["__row_no"], [good[0] + ["1"]])),
        ("매핑된 열이 파일에 없음", xlsx_bytes([f"q3없는열{i}" for i in range(len(hdrs))], good)),
    ]
    bad, statuses = [], []
    for label, data in files:
        r = col_post(ac, P, data, name=f"{p} {label} (예시).xlsx")
        statuses.append(f"{label} {r.status_code}")
        if r.status_code != 422:
            bad.append(f"{label} {r.status_code} (422 기대){' — 500!' if r.status_code >= 500 else ''} {r.text[:100]}")
        elif not r.json().get("fields"):
            bad.append(f"{label} 422 인데 fields 없음")
    for label, extra in (("기준 기간 NUL", {"period_from": "1985-03-01\x00"}), ("기준 기간 99999년", {"period_to": "99999-03-31"}),
                         ("동작 NUL", {"action": "apply\x00"})):
        r = col_post(ac, P, xlsx_bytes(hdrs, good), extra=extra)
        statuses.append(f"{label} {r.status_code}")
        if r.status_code != 422:
            bad.append(f"{label} {r.status_code} (422 기대)")
    out.append(res("X-05", "업로드: 엑셀 아닌 파일 · 빈 파일 · 20MB 초과 · 헤더 결함 → 422", not bad,
                   f"파일 {len(files)} + 폼 3 — 500 {sum(1 for s in statuses if s.endswith(('500', '502', '503')))} · 어긋남 {len(bad)}", bad))
    # 2. 셀 내용 — 수식 글자 · 스크립트 · 10만 자 · NUL 이스케이프 · 파일 이름 경로. 거래처 코드 {p}M 은 매핑해 COL-040 까지 흐르게 한다
    cu = conn.q1("insert into customer (customer_code, customer_name, created_by) values (%s, %s, 'qa3') returning id",
                 (f"{p}cx", f"{p} 거래처 X (예시)"))["id"]
    r = ac.post("/col/customer-map", data={"erp_code": f"{p}M", "erp_name": "<i>q3</i> (예시)", "customer_id": str(cu)})
    bad2 = []
    if r.status_code != 200:
        bad2.append(f"매핑 등록 {r.status_code}")
    evil_code = f"{p}<script>alert(1)</script>"
    hyper = '=HYPERLINK("http://evil.example/q3","클릭")'
    dde = "=CMD|' /C calc'!A0"
    big = "가" * 100_000
    rows = [col_row(fields, customer_code=evil_code, txn_date="1985-03-05", collected_amount="1", method=hyper),
            col_row(fields, customer_code=f"{p}M", txn_date="1985-03-06", collected_amount="2", method=dde),
            col_row(fields, customer_code=f"{p}M", txn_date="1985-03-07", collected_amount="3", method=hyper),
            col_row(fields, customer_code=f"{p}M", txn_date="1985-03-08", collected_amount="4", method=big),
            col_row(fields, customer_code=f"{p}M", txn_date="1985-03-09", collected_amount="5", method="a_x0000_b"),
            col_row(fields, customer_code=f"{p}M", txn_date="1985-03-10", collected_amount="6", method="<img src=x onerror=alert(1)>")]
    fname = f"../../../etc/{p}<b>passwd</b>.xlsx"
    r = col_post(ac, P, xlsx_bytes(hdrs, rows), name=fname)
    if r.status_code != 200:
        bad2.append(f"셀 내용 업로드 {r.status_code} {r.text[:160]}")
        out.append(res("X-06", "셀의 수식·스크립트·긴 글자 → 값으로 저장 · 화면 이스케이프", False, "업로드 실패", bad2))
        return out
    bid = r.json()["id"]
    b = conn.q1("select source_file, row_count, error_count from erp_upload_batch where id = %s", (bid,))
    if b["source_file"] != fname[:200]:
        bad2.append(f"source_file 저장 값 {b['source_file']!r} (올린 이름 그대로 기대)")
    if (b["row_count"], b["error_count"]) != (6, 1):
        bad2.append(f"행 {b['row_count']} · 오류 {b['error_count']} (6 · 1 기대 — 미매핑 코드 1행만)")
    stored = {r_["row_no"]: r_ for r_ in conn.q("select row_no, raw, parsed, error from erp_upload_row where batch_id = %s", (bid,))}
    mcol = cmap.get("method")
    if mcol:
        if stored[3]["raw"].get(mcol) != dde or stored[3]["parsed"].get("method") != dde:
            bad2.append(f"DDE 글자가 그대로 저장되지 않음 raw={stored[3]['raw'].get(mcol)!r}")
        # 엑셀 셀 상한은 32,767자 — openpyxl 이 쓴 파일에 든 길이 그대로 raw 에 보존되고 parsed 는 200자로 잘린다
        raw_len = len(stored[5]["raw"].get(mcol) or "")
        if raw_len < 32_767 or len(stored[5]["parsed"].get("method") or "") != 200:
            bad2.append(f"10만 자 — raw {raw_len} · parsed {len(stored[5]['parsed'].get('method') or '')} (≥32767 · 200 기대)")
    if stored[2]["error"] is None or evil_code not in (stored[2]["error"] or ""):
        bad2.append(f"<script> 거래처 코드 행이 미매핑 오류가 아님 {stored[2]['error']!r}")
    pages = {
        "COL-010 ?batch": ac.get("/col/uploads", params={"batch": str(bid)}, headers={"accept": "text/html"}),
        "COL-020 ?batch": ac.get("/col/upload-history", params={"batch": str(bid)}, headers={"accept": "text/html"}),
        "COL-030": ac.get("/col/customer-map", headers={"accept": "text/html"}),
        "COL-040": ac.get("/col/collections", params={"date_from": "1985-03-01", "date_to": "1985-03-31", "customer": str(cu)},
                          headers={"accept": "text/html"}),
        "HOME": ac.get("/", headers={"accept": "text/html"}),
    }
    escaped_seen = 0
    for name, pg in pages.items():
        if pg.status_code != 200:
            bad2.append(f"{name} {pg.status_code}")
            continue
        html = pg.text
        if "<script>alert(1)</script>" in html or "<img src=x onerror" in html or 'href="http://evil.example' in html or "<b>passwd</b>" in html:
            bad2.append(f"{name}: 저장된 글자가 이스케이프 없이 렌더됨")
        if "&lt;script&gt;alert(1)&lt;/script&gt;" in html or "&lt;img src=x" in html or "&lt;b&gt;passwd" in html:
            escaped_seen += 1
    if not escaped_seen:
        bad2.append("어느 화면에도 이스케이프된 글자가 보이지 않는다(저장 글자가 화면에 안 나왔다 — 검사가 비었다)")
    out.append(res("X-06", "셀의 수식·스크립트·10만 자·파일 이름 경로 → 값으로만 저장 · 화면 이스케이프", not bad2,
                   f"행 6 · 화면 {len(pages)} · 이스케이프 확인 화면 {escaped_seen} · source_file 글자 보존 · 10만 자 셀 raw {len(stored[5]['raw'].get(mcol) or '') if mcol else '-'}자(엑셀 상한 32767) · 어긋남 {len(bad2)}", bad2))
    # 3. 내려받는 엑셀(COL-040 ?format=xlsx) — 저장된 `=…` 글자가 수식 셀로 되살아나면 수식 주입(DDE · HYPERLINK)
    bad3 = []
    x = ac.get("/col/collections", params={"date_from": "1985-03-01", "date_to": "1985-03-31", "customer": str(cu), "format": "xlsx"})
    if x.status_code != 200 or not x.headers.get("content-type", "").startswith(XLSX_MIME):
        bad3.append(f"xlsx 내려받기 {x.status_code} {x.headers.get('content-type')}")
    else:
        wb = load_workbook(io.BytesIO(x.content))
        formulas = [(c.coordinate, str(c.value)[:40]) for row in wb.active.iter_rows() for c in row if c.data_type == "f"]
        if formulas:
            bad3.append(f"수식 셀 {len(formulas)}: " + " · ".join(f"{k}={v}" for k, v in formulas[:3]))
    out.append(res("X-07", "COL-040 엑셀 내려받기 — 저장된 `=…` 글자가 수식 셀로 되살아나지 않는다", not bad3,
                   "수식 셀 0" if not bad3 else "DDE · HYPERLINK 글자가 수식 셀(data_type f)로 쓰였다 — 여는 사람의 엑셀에서 실행될 수 있다", bad3))
    # 4. 진짜 수식 셀(캐시 값 없음)은 빈 칸으로 읽힌다 — preview 로 쓰지 않고 확인 (참고)
    pv = col_post(ac, P, xlsx_bytes(hdrs, [col_row(fields, customer_code=f"{p}M", txn_date="1985-03-11", collected_amount="=1+1")],
                                     text_formulas=False), action="preview")
    if pv.status_code == 200:
        item = pv.json()["items"][0]
        print(f"참고 (X-06): 진짜 수식 셀 `=1+1`(캐시 값 없음)은 수금액 {item['parsed'].get('collected_amount')!r} 로 읽힌다 · error {item['error']!r}"
              " — 값이 아닌 수식은 빈 칸(조용히 0 이 되지는 않는다)")
    return out


def erp_upload_max_bytes() -> int:
    from lcomfine.app import erp_upload

    return erp_upload.MAX_BYTES


def check_col_static() -> list[Result]:
    """정적 — 업로드 모듈 · 라우터가 파일을 디스크에 쓰지 않는다(open · write · Path · tempfile 0) · 쓰기 4 전부 require_fn + log_change."""
    hits = []
    for rel in ("src/lcomfine/app/erp_upload.py", "src/lcomfine/app/routers/col.py"):
        for i, ln in enumerate((ROOT / rel).read_text(encoding="utf-8").splitlines(), 1):
            if ln.strip().startswith("#"):
                continue
            for pat in _COL_WRITE_PATTERNS:
                if re.search(pat, ln):
                    hits.append(f"{rel}:{i} {ln.strip()[:60]}")
    return [res("X-08", "정적: 업로드 파일을 디스크에 쓰지 않는다", not hits,
                f"erp_upload.py · col.py 의 open/write/Path/os./tempfile {len(hits)} (wb.save(buf) BytesIO 만 허용)", hits)]


def check_seed_dev3_static() -> Result:
    """seed_dev3.py — `on conflict … do nothing`(멱등) · 지우는 문장 없음 · created_by 'seed'. 실측(2회 실행 diff 0)은 gate-full 이 잰다."""
    from lcomfine.db import conn

    src = (ROOT / "src/lcomfine/db/seed_dev3.py").read_text(encoding="utf-8")
    body = "\n".join(ln for ln in src.splitlines() if not ln.strip().startswith("#"))
    idem = "on conflict (field) do nothing" in body
    destructive = bool(re.search(r"\b(delete|truncate|drop)\b", body, re.I))
    seeded = conn.q1("select count(*) as n from erp_column_map where created_by = 'seed'")["n"]
    ex = "(예시)" in src
    ok = idem and not destructive and "SEEDED_BY = \"seed\"" in body
    return res("G-21", "seed_dev3 멱등 · (예시) (정적)", ok,
               f"on conflict do nothing {'있음' if idem else '없음'} · delete/truncate {'있음' if destructive else '없음'} · "
               f"erp_column_map(created_by=seed) {seeded}행 · 열 이름은 정본 낱말이라 (예시) 표식 없음(문서에만 — D-802) · 실측 멱등은 gate-full")


# ── 판정표 ──────────────────────────────────────────────────────────────
def w3_rows() -> list[Result]:
    return [Result("G-16", "데이터 질의 검증", FAIL, "W3/W4 — RPT-130 데이터 질의 미구현 (라우터 없음 · 판정 대상 없음)"),
            Result("G-17", "키 없음", FAIL, "W3/W4 — AI 호출(ai.py · sys_ai_call_log) 미구현"),
            Result("G-18", "SSO 토큰", FAIL, "W3/W4 — DOC-010 SSO 미구현")]


def run_all(run_seeds: bool = False) -> list[Result]:
    rows: list[Result] = w3_rows()

    def guard(gid: str, item: str, fn, *args):
        try:
            got = fn(*args)
            return got if isinstance(got, list) else [got]
        except Exception as exc:  # noqa: BLE001 — 검사 자체가 깨진 것도 FAIL 로 보인다(조용히 넘기지 않는다)
            return [Result(gid, item, FAIL, f"검사 실패 — {type(exc).__name__}: {str(exc)[:200]}")]

    rows += guard("G-19", "쓰기 라우트 require_fn", check_static_require_fn)
    rows += guard("G-19", "세션 없음 401", check_no_session_401)
    with Lab() as lab:
        rows += guard("G-19", "R·없음 부서 쓰기 403", check_write_403, lab)
        rows += guard("G-19", "칸 밖 화면 GET 403", check_screen_403, lab)
        rows += guard("G-19", "ID 바꿔치기", check_idor, lab)
        rows += guard("G-19", "권한 칸 변경", check_permission_change, lab)
        rows += guard("G-20", "세션 상태", check_sessions, lab)
        rows += guard("G-21", "비밀", check_secrets)
        if run_seeds:
            rows += guard("G-21", "시드 멱등", check_seed_idempotence)
        rows += guard("G-22", "감사 로그", check_change_log, lab)
        rows += guard("X-01", "DB 끊김", check_db_down, lab)
        rows += guard("X-03", "이상한 입력", check_malformed, lab)
        # W2 수금 (COL)
        rows += guard("G-19", "COL 화면·쓰기 × 부서 8", check_col_access, lab)
        rows += guard("G-20", "COL 비활성·로그아웃 401", check_col_sessions, lab)
        rows += guard("G-22", "COL 쓰기 4 변경 로그", check_col_change_log, lab)
        rows += guard("X-05", "COL 업로드 보안", check_col_upload_abuse, lab)
    rows += guard("X-08", "COL 정적", check_col_static)
    rows += guard("G-21", "seed_dev3 정적", check_seed_dev3_static)
    rows += guard("G-24", "docker", check_docker)
    return rows


def main() -> int:
    if "--e2e" in sys.argv:
        return e2e()
    get_settings()
    t0 = time.time()
    rows = run_all("--run-seeds" in sys.argv)
    print("G-16~G-22 보안 · 운영 (tools/check_security.py · QA3) — 임시 계정으로 실제 요청")
    for r in rows:
        print(r.line())
    if "--run-seeds" not in sys.argv:
        print("참고 G-21 시드 멱등: 이 실행에서는 시드를 돌리지 않았다 — `make gate-full`(gate.py --run-seeds) 또는 `--run-seeds` 가 잰다")
    fails = sum(1 for r in rows if r.status == FAIL)
    print(f"판정: FAIL {fails} · PASS {sum(1 for r in rows if r.status == PASS)} · WARN {sum(1 for r in rows if r.status == WARN)} "
          f"({time.time() - t0:.0f}초)")
    return fails


if __name__ == "__main__":
    sys.exit(main())
