"""DOC 문서관리 — 바로가기(DOC-010 · SSO 진입) · 계정 매핑(DOC-020). 담당 개발1 (W4).

- DOC-010 `GET /doc/go`(F-DOC-012) 안내 화면 · `POST /doc/go`(F-DOC-011 · 권한 예외 = 화면 R 이상) 토큰 발급 → 자동 POST 폼.
  토큰은 `sso.issue` 한 곳에서 만든다 · 폼의 hidden 칸에만 싣는다(주소창 · JSON · 로그 · 변경 로그에 없음).
  `LCOMFINE_DOC_SSO_URL` 이 없으면 `미확정 (D-517)` — 토큰을 발급하지 않는다(422 · 비밀이 없어도 같다 · D-610 —
  RPT 의 「AI 키 없음 422」와 같은 규칙 · 브라우저는 303 → 안내 화면).
  호출마다 `sys_access_log` 구분 SSO 1행(성공 · 실패 · 사유) — 실패는 트랜잭션 밖 자체 연결로 남아 롤백되지 않는다.
- DOC-020 `doc_account_map` — MES 사용자 ↔ 문서관리 계정 ID 1:1(사용 중지 매핑 포함 · DB 제약 2 개와 같다 · D-611).
  엑셀 일괄 등록은 첫 시트 · 1행 헤더 `사번` · `문서관리 계정 ID` · 디스크에 쓰지 않는다 · 오류 행이 하나라도 있으면 아무것도 쓰지 않는다(D-611).
  삭제는 그 매핑으로 발급한 `sso_token` 이 있으면 사용 중지(`active = false`), 없으면 지운다.
"""

from __future__ import annotations

import io
import zipfile

from fastapi import APIRouter, File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse

from ...db import conn
from .. import nav, rbac, sso, templating
from ..settings import get_settings
from ..util import audit, http
from .bas import bad, exists_or_404, like, need, pk, ref_id, s, yn

router = APIRouter()

P_GO = nav.path_of("DOC-010")          # /doc/go
P_MAP = nav.path_of("DOC-020")         # /doc/account-map
ACCOUNT_MAX = 200
XLSX_MAX_BYTES = 5 * 1024 * 1024
XLSX_MAX_ROWS = 5000
XLSX_HEADERS = {"사번": "login_id", "문서관리 계정 ID": "account_id"}


# ── DOC-010 문서관리 바로가기 ───────────────────────────────────────────
def _sso_log(request: Request, user: rbac.User, ok: bool, *, target: str | None = None, detail: str = "") -> None:
    audit.write_log(log_type=audit.SSO, login_id=user.login_id, role_code=user.role_code, method=request.method,
                    path=request.url.path, screen_id="DOC-010", function_id="F-DOC-011", target=target, ok=ok,
                    detail=detail, ip=audit.client_ip(request))


def _undecided_parts() -> list[str]:
    st = get_settings()
    out = []
    if not st.doc_sso_url:
        out.append("문서관리 SSO 주소(LCOMFINE_DOC_SSO_URL)")
    if not st.sso_secret:
        out.append("SSO 서명 비밀(LCOMFINE_SSO_SECRET)")
    return out


@router.get(P_GO, response_class=HTMLResponse)
def go_page(request: Request, user: rbac.User = rbac.require_screen("DOC-010")):
    """F-DOC-012 문서관리 바로가기 — 안내 · 주소 없음 `미확정 (D-517)` · 매핑 없음 「관리자에게 계정 연결을 요청하세요」."""
    with conn.tx() as cur:
        reason, mapping = sso.refusal(cur, user.id)
    return templating.render(request, "doc/go.html",
                             {"missing": _undecided_parts(), "reason": reason, "mapping": mapping,
                              "request_mapping": sso.REQUEST_MAPPING, "ttl": sso.TTL_SECONDS},
                             screen_id="DOC-010")


@router.post(P_GO)
def go(request: Request, user: rbac.User = rbac.require_fn("F-DOC-011")):
    """F-DOC-011 문서관리 진입 — 일회용 토큰 발급 → 새 창의 자동 POST 폼(토큰은 hidden 칸에만)."""
    missing = _undecided_parts()
    if missing:
        what = " · ".join(missing)
        _sso_log(request, user, False, detail=f"{what} 미확정 (D-517) — 발급하지 않음")
        raise http.validation_error(f"{what} 미확정 (D-517) — 토큰을 발급하지 않았습니다",
                                    fields=[{"name": "SSO 설정", "reason": f"{m} 미확정 (D-517)"} for m in missing])
    try:
        with conn.tx() as cur:
            token, jti = sso.issue(cur, user)
    except sso.NotMapped as exc:
        _sso_log(request, user, False, detail=f"발급 거부 — {exc.reason}")
        raise bad("계정 매핑", exc.reason, f"{sso.REQUEST_MAPPING} — {exc.reason}") from None
    _sso_log(request, user, True, target=f"sso_token:{jti}", detail=f"토큰 발급 · 유효 {sso.TTL_SECONDS}초")
    action = get_settings().doc_sso_url
    if http.wants_html(request):
        resp = templating.render(request, "doc/go_post.html",
                                 {"action": action, "token": token, "ttl": sso.TTL_SECONDS}, screen_id="DOC-010")
    else:                       # JSON(테스트 · 스크립트) — 토큰은 싣지 않는다(폼 hidden 만이 토큰이 나가는 길)
        resp = JSONResponse({"ok": True, "message": "문서관리 진입 토큰을 발급했습니다", "jti": jti,
                             "ttl_seconds": sso.TTL_SECONDS, "method": "POST", "field": "token"})
    audit.log_change(request, user, "F-DOC-011", f"sso_token:{jti}", "문서관리 진입 · 토큰 발급")
    return resp


# ── DOC-020 문서관리 계정 매핑 ──────────────────────────────────────────
def _account(v: str | None) -> str:
    a = need(v, "문서관리 계정 ID")
    if len(a) > ACCOUNT_MAX:
        raise bad("문서관리 계정 ID", f"{ACCOUNT_MAX}자 이하로 입력해 주세요")
    return a


def _one_to_one(cur, user_id: int, account_id: str, *, self_id: int | None = None) -> list[dict]:
    """1:1 이 깨지는 사유 목록 — 한 사용자 ↔ 한 계정 ID (사용 중지 매핑도 자리를 차지한다 · D-611)."""
    out = []
    cur.execute("""select m.id, m.active, u.login_id from doc_account_map m join sys_user u on u.user_id = m.user_id
                    where m.user_id = %s and m.id is distinct from %s""", (user_id, self_id))
    r = cur.fetchone()
    if r:
        out.append({"name": "MES 사용자",
                    "reason": f"이미 계정이 연결된 사용자입니다 ({r['login_id']}"
                              f"{' · 사용 중지 매핑 — 그 행을 수정해 다시 쓰세요' if not r['active'] else ''})"})
    cur.execute("""select m.id, m.active, u.login_id from doc_account_map m join sys_user u on u.user_id = m.user_id
                    where m.account_id = %s and m.id is distinct from %s""", (account_id, self_id))
    r = cur.fetchone()
    if r:
        out.append({"name": "문서관리 계정 ID",
                    "reason": f"다른 사용자({r['login_id']})에 이미 연결된 계정 ID 입니다 ({account_id}"
                              f"{' · 사용 중지 매핑' if not r['active'] else ''})"})
    return out


@router.get(P_MAP, response_class=HTMLResponse)
def account_map(request: Request, q: str | None = None, edit: str | None = None,
                user: rbac.User = rbac.require_screen("DOC-020")):
    """F-DOC-024 계정 매핑 조회 — 매핑 목록 · 미연결 사용자 목록 · 사번 동일 방식 `미확정 (D-517)`."""
    pat = like(q, "사번·이름·계정 ID")
    rows = conn.q(
        """select m.id, m.user_id, u.login_id, u.user_name, u.role_code, u.status, m.account_id, m.active,
                  m.created_at, m.created_by, m.updated_at, m.updated_by,
                  (select count(*) from sso_token t where t.user_id = m.user_id and t.account_id = m.account_id) as token_n
             from doc_account_map m join sys_user u on u.user_id = m.user_id
            where (%(q)s::text is null or u.login_id ilike %(q)s or u.user_name ilike %(q)s or m.account_id ilike %(q)s)
            order by u.login_id limit 1000""", {"q": pat})
    unmapped = conn.q(
        """select u.user_id, u.login_id, u.user_name, u.role_code, u.status,
                  (select m.account_id from doc_account_map m where m.user_id = u.user_id) as inactive_account
             from sys_user u
            where not exists (select 1 from doc_account_map m where m.user_id = u.user_id and m.active)
              and (%(q)s::text is null or u.login_id ilike %(q)s or u.user_name ilike %(q)s)
            order by u.login_id limit 1000""", {"q": pat})
    free_users = conn.q("""select u.user_id, u.login_id, u.user_name from sys_user u
                            where not exists (select 1 from doc_account_map m where m.user_id = u.user_id)
                            order by u.login_id""")
    ed = None
    if edit not in (None, ""):
        ed = conn.q1("""select m.*, u.login_id, u.user_name from doc_account_map m join sys_user u on u.user_id = m.user_id
                         where m.id = %s""", (pk(edit),))
        if ed is None:
            raise http.not_found()
    return templating.render(request, "doc/account_map.html",
                             {"rows": rows, "unmapped": unmapped, "free_users": free_users, "edit": ed,
                              "headers": list(XLSX_HEADERS), "form": {"q": q or ""}}, screen_id="DOC-020")


def _cell(v) -> str:
    """엑셀 칸 → 글자. 숫자 사번(1001.0)은 정수 글자로."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


async def _xlsx_rows(file: UploadFile) -> list[dict]:
    """첫 시트 · 1행 헤더 · 메모리에서만 읽는다(디스크에 쓰지 않는다). 반환 [{row_no, login_id, account_id}] (빈 줄 제외)."""
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    data = await file.read()
    if not data:
        raise bad("엑셀 파일", "빈 파일입니다")
    if len(data) > XLSX_MAX_BYTES:
        raise bad("엑셀 파일", f"{XLSX_MAX_BYTES // (1024 * 1024)}MB 이하만 올릴 수 있습니다")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise bad("엑셀 파일", f"엑셀(.xlsx) 파일이 아닙니다 [{type(exc).__name__}]") from None
    try:
        it = wb.worksheets[0].iter_rows(values_only=True)
        head = next(it, None)
        names = [("" if h is None else str(h).strip()) for h in (head or ())]
        missing = [h for h in XLSX_HEADERS if h not in names]
        if missing:
            raise bad("엑셀 파일", f"1행에 열 이름이 없습니다 — {' · '.join(missing)} (필요한 열: {' · '.join(XLSX_HEADERS)})")
        idx = {XLSX_HEADERS[h]: names.index(h) for h in XLSX_HEADERS}
        out = []
        for row_no, values in enumerate(it, start=2):
            vals = list(values or ())
            cell = {k: _cell(vals[i] if i < len(vals) else None) for k, i in idx.items()}
            if not any(cell.values()):
                continue
            out.append({"row_no": row_no, **cell})
            if len(out) > XLSX_MAX_ROWS:
                raise bad("엑셀 파일", f"{XLSX_MAX_ROWS}행 이하만 올릴 수 있습니다")
    except (SyntaxError, zipfile.BadZipFile, EOFError) as exc:   # 시트 XML 이 깨짐 — read_only 는 행을 읽을 때 파싱한다 (DEF-QA3-008)
        raise bad("엑셀 파일", f"엑셀 시트를 읽을 수 없습니다 — 파일이 깨졌습니다 [{type(exc).__name__}]") from None
    finally:
        wb.close()
    if not out:
        raise bad("엑셀 파일", "등록할 행이 없습니다 (1행은 열 이름)")
    return out


def _bulk(cur, rows: list[dict], user: rbac.User) -> int:
    """엑셀 행 검증 — 오류 행이 하나라도 있으면 422(행마다 사유) · 아무것도 쓰지 않는다. 반환 = 등록 행 수."""
    errors: list[dict] = []
    seen_login: dict[str, int] = {}
    seen_acc: dict[str, int] = {}
    todo = []
    for r in rows:
        n = r["row_no"]
        reasons = []
        lid, acc = r["login_id"], r["account_id"]
        if "\x00" in lid or "\x00" in acc:
            reasons.append("글자에 NUL(0x00)을 넣을 수 없습니다")
        if not lid:
            reasons.append("사번이 비었습니다")
        if not acc:
            reasons.append("문서관리 계정 ID 가 비었습니다")
        elif len(acc) > ACCOUNT_MAX:
            reasons.append(f"문서관리 계정 ID 는 {ACCOUNT_MAX}자 이하")
        if not reasons:
            if lid in seen_login:
                reasons.append(f"파일 안에서 사번이 겹칩니다 ({seen_login[lid]}행)")
            if acc in seen_acc:
                reasons.append(f"파일 안에서 계정 ID 가 겹칩니다 ({seen_acc[acc]}행)")
            seen_login.setdefault(lid, n)
            seen_acc.setdefault(acc, n)
            cur.execute("select user_id from sys_user where login_id = %s", (lid,))
            u = cur.fetchone()
            if u is None:
                reasons.append(f"없는 사번입니다 ({lid})")
            else:
                reasons += [f["reason"] for f in _one_to_one(cur, u["user_id"], acc)]
                todo.append((u["user_id"], acc))
        if reasons:
            errors.append({"name": f"{n}행", "reason": " · ".join(reasons)})
    if errors:
        raise http.validation_error(f"엑셀 일괄 등록 — 오류 행 {len(errors)} / {len(rows)} (아무것도 등록하지 않았습니다)",
                                    fields=errors[:200])
    for uid, acc in todo:
        cur.execute("insert into doc_account_map (user_id, account_id, created_by) values (%s, %s, %s)",
                    (uid, acc, user.login_id))
    return len(todo)


@router.post(P_MAP)
async def account_map_create(request: Request, user_id: str | None = Form(None), account_id: str | None = Form(None),
                             file: UploadFile | None = File(None), user: rbac.User = rbac.require_fn("F-DOC-021")):
    """F-DOC-021 계정 매핑 등록 — 한 건(`user_id` · `account_id`) 또는 엑셀 일괄(`file`)."""
    if file is not None and (file.filename or "").strip():
        rows = await _xlsx_rows(file)
        with conn.tx() as cur:
            n = _bulk(cur, rows, user)
        resp = http.saved(request, f"계정 매핑 {n}건을 일괄 등록했습니다", back=P_MAP, data={"count": n})
        audit.log_change(request, user, "F-DOC-021", f"doc_account_map:엑셀 {n}건", f"계정 매핑 엑셀 일괄 등록 {n}건")
        return resp
    uid = ref_id(user_id, "MES 사용자")
    acc = _account(account_id)
    with conn.tx() as cur:
        cur.execute("select login_id from sys_user where user_id = %s", (uid,))
        u = cur.fetchone()
        if u is None:
            raise bad("MES 사용자", "목록에서 고른 값이 아닙니다")
        errs = _one_to_one(cur, uid, acc)
        if errs:
            raise http.validation_error("1:1 연결이 깨집니다 — 한 MES 사용자에 한 문서관리 계정 ID", fields=errs)
        cur.execute("insert into doc_account_map (user_id, account_id, created_by) values (%s, %s, %s) returning id",
                    (uid, acc, user.login_id))
        new_id = cur.fetchone()["id"]
    resp = http.saved(request, "계정 매핑을 등록했습니다", back=P_MAP, data={"id": new_id})
    audit.log_change(request, user, "F-DOC-021", f"doc_account_map:{u['login_id']}", f"계정 매핑 등록 → {acc}")
    return resp


@router.post(P_MAP + "/{id}")
def account_map_update(request: Request, id: str, user_id: str | None = Form(None), account_id: str | None = Form(None),
                       use_yn: str | None = Form(None), user: rbac.User = rbac.require_fn("F-DOC-022")):
    """F-DOC-022 계정 매핑 수정 — 계정 ID · 사용 여부. MES 사용자는 바꾸지 않는다(지우고 새로 등록 · D-611)."""
    mid = pk(id)
    acc = _account(account_id)
    active = yn(use_yn) == "Y"
    with conn.tx() as cur:
        row = exists_or_404(cur, "doc_account_map", mid, "id, user_id, account_id, active")
        if s(user_id, "MES 사용자") is not None and s(user_id, "MES 사용자") != str(row["user_id"]):
            raise bad("MES 사용자", "매핑의 MES 사용자는 바꾸지 않습니다 — 지우고 새로 등록하세요", "MES 사용자는 바꿀 수 없습니다")
        errs = _one_to_one(cur, row["user_id"], acc, self_id=mid)
        if errs:
            raise http.validation_error("1:1 연결이 깨집니다 — 한 MES 사용자에 한 문서관리 계정 ID", fields=errs)
        cur.execute("""update doc_account_map set account_id = %s, active = %s, updated_at = now(), updated_by = %s
                        where id = %s""", (acc, active, user.login_id, mid))
        cur.execute("select login_id from sys_user where user_id = %s", (row["user_id"],))
        lid = cur.fetchone()["login_id"]
    resp = http.saved(request, "계정 매핑을 수정했습니다", back=P_MAP, data={"id": mid})
    audit.log_change(request, user, "F-DOC-022", f"doc_account_map:{lid}",
                     f"계정 매핑 수정 · {row['account_id']} → {acc} · 사용 {'Y' if active else 'N'}")
    return resp


@router.post(P_MAP + "/{id}/delete")
def account_map_delete(request: Request, id: str, user: rbac.User = rbac.require_fn("F-DOC-023")):
    """F-DOC-023 계정 매핑 삭제 — 그 매핑으로 발급한 토큰(`sso_token`)이 있으면 사용 중지, 없으면 지운다."""
    mid = pk(id)
    with conn.tx() as cur:
        row = exists_or_404(cur, "doc_account_map", mid, "id, user_id, account_id")
        cur.execute("select 1 from sso_token where user_id = %s and account_id = %s limit 1",
                    (row["user_id"], row["account_id"]))
        if cur.fetchone():
            cur.execute("update doc_account_map set active = false, updated_at = now(), updated_by = %s where id = %s",
                        (user.login_id, mid))
            done = "사용 중지"
        else:
            cur.execute("delete from doc_account_map where id = %s", (mid,))
            done = "삭제"
        cur.execute("select login_id from sys_user where user_id = %s", (row["user_id"],))
        lid = cur.fetchone()["login_id"]
    msg = "발급 기록이 있어 지우지 않고 사용 중지로 바꿨습니다" if done == "사용 중지" else "계정 매핑을 삭제했습니다"
    resp = http.saved(request, msg, back=P_MAP, data={"id": mid, "result": done})
    audit.log_change(request, user, "F-DOC-023", f"doc_account_map:{lid}", f"계정 매핑 {done} · {row['account_id']}")
    return resp

