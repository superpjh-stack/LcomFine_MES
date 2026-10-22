# tests/test_dev3_col.py — COL 수금리포트 (개발3 · W2) + erp_upload + stats.customer_receivable · due_compliance.
# `(예시)` 엑셀은 테스트 안에서 openpyxl 로 만든다(저장소에 바이너리 없음). 기준 기간은 1990~1993년(다른 데이터와 겹치지 않는 해).
# 열 매핑은 전역 설정이라 모듈 시작에 스냅샷을 떠 테스트 매핑으로 바꾸고 끝에 되돌린다. 거래처 · 출하는 test_dev3_world.World.
from __future__ import annotations

import io
from datetime import date, datetime, timedelta
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from lcomfine.app import erp_upload, stats
from lcomfine.db import conn
from test_dev3_world import World, login

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
#: 테스트 매핑 — MES 항목 → 엑셀 열 이름 (전부 테스트 전용 이름)
H = {"customer_code": "T3거래처코드", "txn_date": "T3일자", "sales_amount": "T3매출액", "collected_amount": "T3수금액",
     "balance_amount": "T3잔액", "method": "T3수단", "bill_due_date": "T3어음만기"}
COLS = list(H)


def xlsx(rows: list[list], headers: list[str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(headers or [H[c] for c in COLS])
    for i, r in enumerate(rows, start=2):
        for j, v in enumerate(r, start=1):
            cell = ws.cell(row=i, column=j, value=v)
            if isinstance(v, str):
                cell.data_type = "s"      # ERP 내려받기처럼 글자 셀 — `=…` 도 수식이 아니라 글자로 올라간다
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def row(code, day, sales=None, collected=None, balance=None, method=None, bill_due=None) -> list:
    return [code, day, sales, collected, balance, method, bill_due]


@pytest.fixture(scope="module")
def c():
    return {lid: login(lid) for lid in ("ac", "sl", "sh", "admin", "ex")}


@pytest.fixture(scope="module", autouse=True)
def column_map(c):
    """열 매핑을 테스트 매핑으로 바꾸고 모듈 끝에 원래 매핑으로 되돌린다."""
    before = conn.q("select field, column_name, created_by from erp_column_map order by id")
    r = c["ac"].post("/col/uploads", data={"action": "map", **{f"map_{k}": v for k, v in H.items()}})
    assert r.status_code == 200, r.text
    assert erp_upload.column_map() == H
    yield H
    with conn.tx() as cur:
        cur.execute("delete from erp_column_map")
        for b in before:
            cur.execute("insert into erp_column_map (field, column_name, created_by) values (%s, %s, %s)",
                        (b["field"], b["column_name"], b["created_by"]))
    assert len(conn.q("select 1 from erp_column_map")) == len(before)


class ColWorld(World):
    """World + ERP 쪽: 이 테스트가 만든 회차 · 매핑을 끝에 지운다."""

    def __init__(self) -> None:
        super().__init__()
        self.code = f"{self.tag}-E"              # 매핑할 ERP 코드
        self.unmapped = f"{self.tag}-U"          # 매핑하지 않는 ERP 코드
        self.batches: list[int] = []

    def upload(self, client, f: date, t: date, rows: list[list], *, action: str = "apply", name: str = "erp (예시).xlsx",
               headers: list[str] | None = None):
        r = client.post("/col/uploads", data={"action": action, "period_from": f.isoformat(), "period_to": t.isoformat()},
                        files={"file": (name, xlsx(rows, headers), XLSX)})
        if r.status_code == 200 and action == "apply":
            self.batches.append(r.json()["id"])
        return r

    def map_code(self, client, code: str | None = None, **extra):
        r = client.post("/col/customer-map", data={"erp_code": code or self.code, "erp_name": f"ERP 거래처 {self.tag} (예시)",
                                                   "customer_id": str(self.customer_id), **extra})
        return r

    def batch(self, bid: int) -> dict:
        return conn.q1("select * from erp_upload_batch where id = %s", (bid,))

    def rows(self, bid: int) -> list[dict]:
        return conn.q("select row_no, raw, parsed, error from erp_upload_row where batch_id = %s order by row_no", (bid,))

    def notes(self, bid: int) -> list[dict]:
        return conn.q("select to_role, handled_at, link, message from notification where link = %s order by to_role",
                      (erp_upload.error_link(bid),))

    def close(self) -> None:
        ids = conn.q("""select id from erp_upload_batch where id = any(%s) or created_by = 'test'
                        or period_from between '1990-01-01' and '1993-12-31'""", (self.batches,))
        ids = [r["id"] for r in ids]
        with conn.tx() as cur:
            if ids:
                cur.execute("delete from notification where dedupe_key like 'erp_upload_error:%%' and link = any(%s)",
                            ([erp_upload.error_link(i) for i in ids],))
                cur.execute("delete from erp_upload_row where batch_id = any(%s)", (ids,))
                cur.execute("delete from erp_upload_batch where id = any(%s)", (ids,))   # replaced_by FK 는 문장 끝에 검사된다
            cur.execute("delete from erp_customer_map where customer_id = %s or erp_code like %s", (self.customer_id, f"{self.tag}%"))
        super().close()


@pytest.fixture
def w():
    world = ColWorld()
    yield world
    world.close()


P1 = (date(1990, 1, 1), date(1990, 1, 31))
P2 = (date(1991, 6, 1), date(1991, 6, 30))


# ── F-COL-011 업로드 · 회차 교체 · 오류 행 · 알림 ────────────────────────────

@pytest.mark.fn("F-COL-011")
def test_upload_applies_replaces_and_keeps_history(c, w):
    assert w.map_code(c["ac"]).status_code == 200
    rows = [row(w.code, "1990-01-05", "1000", "0", "1000", "PM01", None),
            row(w.code, "1990-01-20", "0", "600", "400", "PM02", "1990-02-15"),
            row(w.code, "1990/13/01", "1", "1", None, None, None),          # 날짜 형식 오류
            row(w.unmapped, "1990-01-21", "500", None, None, None, None)]   # 거래처 코드 미매핑
    before_notes = conn.q1("select count(*) as n from notification where event_key = 'erp_upload_error'")["n"]
    r = w.upload(c["ac"], *P1, rows)
    assert r.status_code == 200, r.text
    body = r.json()
    assert (body["rows"], body["errors"], body["round_no"], body["replaced"]) == (4, 2, 1, None)
    b1 = w.batch(body["id"])
    assert (b1["status"], b1["row_count"], b1["error_count"], b1["replaced_by"]) == ("반영", 4, 2, None)
    assert b1["batch_no"].startswith("ERP-")
    stored = w.rows(b1["id"])
    assert [x["error"] is None for x in stored] == [True, True, False, False]
    assert "일자" in stored[2]["error"] and erp_upload.UNMAPPED_PREFIX in stored[3]["error"]
    assert stored[0]["raw"][H["customer_code"]] == w.code and stored[0]["parsed"]["sales_amount"] == "1000.00"
    # 알림 — 회계부 1행 · 수신 부서는 AC 만 (D-511) · link 는 업로드 이력의 그 회차
    notes = w.notes(b1["id"])
    assert [n["to_role"] for n in notes] == ["AC"] and notes[0]["handled_at"] is None
    assert "미매핑" in notes[0]["message"] and w.unmapped in notes[0]["message"]
    assert conn.q1("select count(*) as n from notification where event_key = 'erp_upload_error'")["n"] == before_notes + 1
    # 반영 집계에는 오류 행이 없다 (G-13)
    col = stats.collections(*P1, w.customer_id)
    assert col["total"] == {"rows": 2, "sales_amount": Decimal("1000.00"), "collected_amount": Decimal("600.00")}
    assert [x["bill_due"] for x in col["rows"]] == [None, "도래"]
    assert conn.q1("select count(*) as n from v_dq_collections where batch_no = %s", (b1["batch_no"],))["n"] == 2
    # 같은 기간 재업로드 = 새 회차 반영 · 이전 회차 교체 · 이전 행 보존 (G-12)
    r2 = w.upload(c["ac"], *P1, [row(w.code, "1990-01-05", "1000", "0"), row(w.code, "1990-01-20", "0", "700"),
                                  row(w.code, "1990-01-25", "0", "100")])
    assert r2.status_code == 200, r2.text
    body2 = r2.json()
    assert (body2["round_no"], body2["replaced"], body2["errors"]) == (2, b1["batch_no"], 0)
    b1_after, b2 = w.batch(b1["id"]), w.batch(body2["id"])
    assert (b1_after["status"], b1_after["replaced_by"]) == ("교체", b2["id"])
    assert (b2["status"], b2["replaced_by"], b2["round_no"]) == ("반영", None, 2)
    assert len(w.rows(b1["id"])) == 4                                     # 이전 회차 행은 지우지 않는다
    assert w.notes(b1["id"])[0]["handled_at"] is not None                  # 교체된 회차의 오류 알림은 닫힌다
    assert w.notes(b2["id"]) == []                                          # 오류 0 → 알림 없음
    col2 = stats.collections(*P1, w.customer_id)
    assert col2["total"]["collected_amount"] == Decimal("800.00") and col2["total"]["rows"] == 3
    assert conn.q1("select count(*) as n from v_dq_collections where batch_no = %s", (b1["batch_no"],))["n"] == 0
    assert conn.q1("select count(*) as n from erp_upload_batch where period_from = %s and period_to = %s and status = '반영'",
                   P1)["n"] == 1
    assert conn.q1("select count(*) as n from sys_change_log where function_id = 'F-COL-011' and target = %s",
                   (f"erp_upload_batch:{b2['batch_no']}",))["n"] == 1
    # 누락 · 형식 · 권한 — 422 는 항목을 짚는다
    r = c["ac"].post("/col/uploads", data={"action": "apply", "period_from": "1990-01-01", "period_to": "1990-01-31"})
    assert r.status_code == 422 and r.json()["fields"][0]["name"] == "파일"
    r = c["ac"].post("/col/uploads", data={"action": "apply"}, files={"file": ("a.xlsx", xlsx([row(w.code, "1990-01-05", "1")]), XLSX)})
    assert r.status_code == 422 and "기준 기간" in r.json()["fields"][0]["name"]
    r = c["ac"].post("/col/uploads", data={"action": "apply", "period_from": "1990-02-01", "period_to": "1990-01-31"},
                     files={"file": ("a.xlsx", xlsx([row(w.code, "1990-01-05", "1")]), XLSX)})
    assert r.status_code == 422 and r.json()["fields"][0]["name"] == "기준 기간"
    r = c["ac"].post("/col/uploads", data={"action": "apply", "period_from": "1990-01-01", "period_to": "1990-01-31"},
                     files={"file": ("a.xlsx", b"not an xlsx file", XLSX)})
    assert r.status_code == 422 and r.json()["fields"][0]["name"] == "파일"
    r = w.upload(c["ac"], *P1, [row(w.code, "1990-01-05", "1")], headers=["엉뚱한열", "b", "c", "d", "e", "f", "g"])
    assert r.status_code == 422 and r.json()["fields"][0]["name"] == "거래처코드"      # 매핑된 열이 파일에 없다
    r = w.upload(c["ac"], *P1, [], action="apply")
    assert r.status_code == 422 and r.json()["fields"][0]["name"] == "파일"              # 데이터 행 없음
    r = w.upload(c["ac"], *P1, [row(w.code, "1990-01-05", "1")], action="bogus")
    assert r.status_code == 422 and r.json()["fields"][0]["name"] == "동작"
    assert w.upload(c["sl"], *P1, [row(w.code, "1990-01-05", "1")]).status_code == 403    # 영업부 R
    assert w.upload(c["sh"], *P1, [row(w.code, "1990-01-05", "1")]).status_code == 403    # 출하실 칸 없음
    # 열 매핑 저장 — 필수 빠지면 422 · 같은 열 두 항목 422 · NUL 422
    r = c["ac"].post("/col/uploads", data={"action": "map", "map_txn_date": "x"})
    assert r.status_code == 422 and r.json()["fields"][0]["name"] == "거래처코드"
    r = c["ac"].post("/col/uploads", data={"action": "map", **{f"map_{k}": v for k, v in H.items()}, "map_balance_amount": H["sales_amount"]})
    assert r.status_code == 422 and r.json()["fields"][0]["name"] == "잔액"
    r = c["ac"].post("/col/uploads", data={"action": "map", **{f"map_{k}": v for k, v in H.items()}, "map_method": "a\x00b"})
    assert r.status_code == 422 and "수단" in r.json()["fields"][0]["name"]
    assert erp_upload.column_map() == H


@pytest.mark.fn("F-COL-012")
def test_upload_screen_shows_live_batch_errors_and_preview_writes_nothing(c, w):
    assert w.map_code(c["ac"]).status_code == 200
    r = w.upload(c["ac"], *P2, [row(w.code, "1991-06-03", "100", "0"), row(w.unmapped, "1991-06-04", "10", "0"),
                                 row(w.code, "bad-date", "1", "1")])
    bid = r.json()["id"]
    b = w.batch(bid)
    page = c["ac"].get("/col/uploads")
    assert page.status_code == 200 and b["batch_no"] in page.text and 'data-row="3"' in page.text and 'data-row="4"' in page.text
    assert H["customer_code"] in page.text and f'data-unmapped' not in page.text    # 열 매핑 보임 · 미매핑 목록은 COL-030
    assert c["ac"].get("/col/uploads", params={"batch": str(bid)}).status_code == 200
    assert c["ac"].get("/col/uploads", params={"batch": "999999999"}).status_code == 404
    assert c["ac"].get("/col/uploads", params={"batch": "abc"}).status_code == 404
    assert c["sl"].get("/col/uploads").status_code == 200                           # 영업부 R — 화면은 보되 폼 없음
    assert "F-COL-011" in c["sl"].get("/col/uploads").text
    assert c["sh"].get("/col/uploads").status_code == 403
    # 미리보기 — 아무것도 쓰지 않는다
    before = conn.q1("select (select count(*) from erp_upload_batch) as b, (select count(*) from erp_upload_row) as r, "
                     "(select count(*) from notification) as n")
    pv = w.upload(c["ac"], *P2, [row(w.code, "1991-06-10", "5", "0"), row(w.unmapped, "x", None, None)], action="preview")
    assert pv.status_code == 200 and pv.json()["applied"] is False and (pv.json()["rows"], pv.json()["errors"]) == (2, 1)
    assert conn.q1("select (select count(*) from erp_upload_batch) as b, (select count(*) from erp_upload_row) as r, "
                   "(select count(*) from notification) as n") == before
    html = c["ac"].post("/col/uploads", data={"action": "preview", "period_from": "1991-06-01", "period_to": "1991-06-30"},
                        files={"file": ("p.xlsx", xlsx([row(w.code, "1991-06-10", "5", "0"), row(w.unmapped, "1991-06-11", "5", "0")]), XLSX)},
                        headers={"Accept": "text/html"})
    assert html.status_code == 200 and "반영하지 않았다" in html.text and 'data-row="3"' in html.text and w.unmapped in html.text


# ── F-COL-021 업로드 이력 ─────────────────────────────────────────────────

@pytest.mark.fn("F-COL-021")
def test_history_distinguishes_live_and_replaced(c, w):
    assert w.map_code(c["ac"]).status_code == 200
    r1 = w.upload(c["ac"], *P2, [row(w.code, "1991-06-03", "100", "0"), row(w.code, "nope", "1", "1")]).json()
    r2 = w.upload(c["ac"], *P2, [row(w.code, "1991-06-03", "100", "0")]).json()
    page = c["ac"].get("/col/upload-history")
    assert page.status_code == 200
    seg1 = page.text[page.text.index(f'data-batch="{r1["id"]}"'):][:200]
    seg2 = page.text[page.text.index(f'data-batch="{r2["id"]}"'):][:200]
    assert 'data-status="교체"' in seg1 and "교체됨" in seg1 and 'data-status="반영"' in seg2
    detail = c["ac"].get("/col/upload-history", params={"batch": str(r1["id"])})
    assert detail.status_code == 200 and 'data-row="3"' in detail.text and "교체됨" in detail.text and r2["batch_no"] in detail.text
    only_live = c["ac"].get("/col/upload-history", params={"status": "반영"})
    assert f'data-batch="{r2["id"]}"' in only_live.text and f'data-batch="{r1["id"]}"' not in only_live.text
    assert c["ac"].get("/col/upload-history", params={"status": "뭐"}).status_code == 422
    assert c["ac"].get("/col/upload-history", params={"batch": "999999999"}).status_code == 404
    assert c["sh"].get("/col/upload-history").status_code == 403 and c["ex"].get("/col/upload-history").status_code == 403


# ── COL-030 거래처 코드 매핑 ──────────────────────────────────────────────

@pytest.mark.fn("F-COL-031")
def test_customer_map_create_remaps_unmapped_rows(c, w):
    r = w.upload(c["ac"], *P2, [row(w.code, "1991-06-03", "100", "0"), row(w.code, "1991-06-04", "0", "40")])
    bid = r.json()["id"]
    assert w.batch(bid)["error_count"] == 2 and len(w.notes(bid)) == 1
    page = c["ac"].get("/col/customer-map")
    assert page.status_code == 200 and f'data-unmapped="{w.code}"' in page.text
    assert [u["erp_code"] for u in erp_upload.unmapped_codes() if u["erp_code"] == w.code] == [w.code]
    assert stats.customer_receivable(w.customer_id)["status"] == "미수집"
    r = w.map_code(c["ac"])
    assert r.status_code == 200 and r.json()["remapped"] == 2 and (r.json()["reinterpreted"], r.json()["applied"]) == (2, 2)
    assert "다시 해석한 행 2 · 그중 반영 2" in r.json()["message"] and "여전히 제외" not in r.json()["message"]
    assert w.batch(bid)["error_count"] == 0 and all(x["error"] is None for x in w.rows(bid))
    assert w.notes(bid)[0]["handled_at"] is not None                         # 오류 0 → 회계부 알림 닫힘
    assert f'data-unmapped="{w.code}"' not in c["ac"].get("/col/customer-map").text
    got = stats.customer_receivable(w.customer_id)
    assert (got["status"], got["sales_amount"], got["collected_amount"], got["receivable_amount"], got["erp_codes"]) == \
        ("ok", Decimal("100.00"), Decimal("40.00"), Decimal("60.00"), [w.code])
    assert got["asof"] is not None and got["rows"] == 2
    assert conn.q1("select count(*) as n from sys_change_log where function_id = 'F-COL-031' and target = %s",
                   (f"erp_customer_map:{w.code}",))["n"] == 1
    # 중복 · 누락 · NUL · 없는 거래처 · 권한
    assert w.map_code(c["ac"]).status_code == 422
    r = c["ac"].post("/col/customer-map", data={"erp_code": "", "customer_id": str(w.customer_id)})
    assert r.status_code == 422 and r.json()["fields"][0]["name"] == "ERP 거래처코드"
    r = c["ac"].post("/col/customer-map", data={"erp_code": f"{w.tag}-X", "customer_id": ""})
    assert r.status_code == 422 and r.json()["fields"][0]["name"] == "MES 거래처"
    r = c["ac"].post("/col/customer-map", data={"erp_code": f"{w.tag}-X", "erp_name": "a\x00b", "customer_id": str(w.customer_id)})
    assert r.status_code == 422 and r.json()["fields"][0]["name"] == "ERP 거래처명"
    r = c["ac"].post("/col/customer-map", data={"erp_code": f"{w.tag}-X", "customer_id": "999999999"})
    assert r.status_code == 422 and r.json()["fields"][0]["name"] == "MES 거래처"
    assert w.map_code(c["sl"], f"{w.tag}-X").status_code == 403


@pytest.mark.fn("F-COL-031")
def test_customer_map_create_counts_only_rows_that_become_applied(c, w):
    """DEF-QA2-004 — 미매핑 + 다른 오류(날짜 · 금액)가 함께 있던 행은 다시 해석하지만 반영에 들지 않는다. 문장의 수가 DB 와 같다."""
    bid = w.upload(c["ac"], *P2, [row(w.code, "1991-06-03", "100", "0"), row(w.code, "1991-06-04", "0", "40"),
                                   row(w.code, "1991/13/01", "1", "1"), row(w.code, "1991-06-06", "abc", "0")]).json()["id"]
    assert w.batch(bid)["error_count"] == 4
    r = w.map_code(c["ac"])
    assert r.status_code == 200
    body = r.json()
    assert (body["reinterpreted"], body["applied"], body["still_error"], body["remapped"]) == (4, 2, 2, 2)
    assert "다시 해석한 행 4 · 그중 반영 2 · 다른 오류로 여전히 제외 2" in body["message"]
    rows = w.rows(bid)
    assert [x["error"] is None for x in rows] == [True, True, False, False]
    assert all(erp_upload.UNMAPPED_PREFIX not in (x["error"] or "") for x in rows)      # 미매핑 글자는 전부 사라졌다
    assert "일자" in rows[2]["error"] and "매출액" in rows[3]["error"]
    assert w.batch(bid)["error_count"] == 2 and w.notes(bid)[0]["handled_at"] is None   # 오류가 남아 알림은 열린 채
    assert conn.q1("select count(*) as n from v_dq_collections where erp_customer_code = %s", (w.code,))["n"] == 2
    assert conn.q1("select detail from sys_change_log where function_id = 'F-COL-031' and target = %s order by id desc limit 1",
                   (f"erp_customer_map:{w.code}",))["detail"].endswith("다시 해석한 행 4 · 그중 반영 2 · 다른 오류로 여전히 제외 2")


@pytest.mark.fn("F-COL-032")
def test_customer_map_update_disable_and_reenable(c, w):
    mid = w.map_code(c["ac"]).json()["id"]
    bid = w.upload(c["ac"], *P2, [row(w.code, "1991-06-03", "100", "0")]).json()["id"]
    assert w.batch(bid)["error_count"] == 0
    r = c["ac"].post(f"/col/customer-map/{mid}", data={"erp_name": "바꾼 이름 (예시)", "customer_id": str(w.customer_id), "use_yn": "N"})
    assert r.status_code == 200 and r.json()["use_yn"] == "N"
    assert [u["erp_code"] for u in erp_upload.unmapped_codes() if u["erp_code"] == w.code] == [w.code]   # 다시 미매핑 목록에
    assert w.batch(bid)["error_count"] == 0                                 # 사용 중지는 매핑만 바꾼다 — 행은 그대로
    assert stats.customer_receivable(w.customer_id)["status"] == "미수집"
    r = c["ac"].post(f"/col/customer-map/{mid}", data={"customer_id": str(w.customer_id), "use_yn": "Y"})
    assert r.status_code == 200 and (r.json()["reinterpreted"], r.json()["applied"], r.json()["remapped"]) == (1, 0, 0)
    assert conn.q1("select erp_name, use_yn from erp_customer_map where id = %s", (mid,)) == {"erp_name": None, "use_yn": "Y"}
    assert c["ac"].post("/col/customer-map/999999999", data={"customer_id": str(w.customer_id)}).status_code == 404
    assert c["ac"].post("/col/customer-map/abc", data={"customer_id": str(w.customer_id)}).status_code == 404
    r = c["ac"].post(f"/col/customer-map/{mid}", data={"customer_id": str(w.customer_id), "use_yn": "Z"})
    assert r.status_code == 422 and r.json()["fields"][0]["name"] == "사용 여부"
    r = c["ac"].post(f"/col/customer-map/{mid}", data={"customer_id": ""})
    assert r.status_code == 422 and r.json()["fields"][0]["name"] == "MES 거래처"
    assert c["sl"].post(f"/col/customer-map/{mid}", data={"customer_id": str(w.customer_id)}).status_code == 403
    assert c["ac"].get("/col/customer-map", params={"edit": str(mid)}).status_code == 200


@pytest.mark.fn("F-COL-033")
def test_customer_map_delete_disables_when_referenced(c, w):
    mid = w.map_code(c["ac"]).json()["id"]
    other = w.map_code(c["ac"], f"{w.tag}-Z").json()["id"]
    w.upload(c["ac"], *P2, [row(w.code, "1991-06-03", "100", "0")])
    r = c["ac"].post(f"/col/customer-map/{mid}/delete")
    assert r.status_code == 200 and r.json()["disabled"] is True and r.json()["refs"] == 1
    assert conn.q1("select use_yn from erp_customer_map where id = %s", (mid,))["use_yn"] == "N"
    r = c["ac"].post(f"/col/customer-map/{other}/delete")
    assert r.status_code == 200 and r.json()["disabled"] is False
    assert conn.q1("select 1 from erp_customer_map where id = %s", (other,)) is None
    assert c["ac"].post(f"/col/customer-map/{other}/delete").status_code == 404
    assert c["sl"].post(f"/col/customer-map/{mid}/delete").status_code == 403
    assert conn.q1("select count(*) as n from sys_change_log where function_id = 'F-COL-033'")["n"] >= 2


@pytest.mark.fn("F-COL-034")
def test_customer_map_screen(c, w):
    mid = w.map_code(c["ac"]).json()["id"]
    page = c["ac"].get("/col/customer-map", params={"q": w.tag})
    assert page.status_code == 200 and f'data-map="{mid}"' in page.text
    assert f'data-map="{mid}"' not in c["ac"].get("/col/customer-map", params={"use": "N"}).text
    assert c["ac"].get("/col/customer-map", params={"use": "X"}).status_code == 422
    assert c["ac"].get("/col/customer-map", params={"edit": "999999999"}).status_code == 404
    assert c["ac"].get("/col/customer-map", params={"code": w.unmapped}).status_code == 200
    assert c["sl"].get("/col/customer-map").status_code == 200 and c["sh"].get("/col/customer-map").status_code == 403


# ── COL-040 ~ COL-070 조회 ────────────────────────────────────────────────

@pytest.mark.fn("F-COL-041")
def test_collections_screen_and_xlsx(c, w):
    w.map_code(c["ac"])
    far = (date.today() + timedelta(days=30)).isoformat()
    w.upload(c["ac"], *P2, [row(w.code, "1991-06-03", "100", "0", None, "PM01"), row(w.code, "1991-06-10", "0", "70", "30", "PM02", "1991-07-01"),
                             row(w.code, "1991-06-20", "0", "10", "20", "PM02", far), row(w.code, "1991-07-01", "999", "0")])
    r = c["ac"].get("/col/collections", params={"date_from": "1991-06-01", "date_to": "1991-06-30", "customer": str(w.customer_id)})
    assert r.status_code == 200 and "999" not in r.text.split("data-erp-code")[1] if "data-erp-code" in r.text else False
    rows = [seg for seg in r.text.split('<tr data-erp-code="')[1:]]
    assert len(rows) == 3 and "도래" in rows[1] and "예정" in rows[2]
    r = c["ac"].get("/col/collections", params={"date_from": "1991-06-01", "date_to": "1991-06-30", "customer": str(w.customer_id), "method": "PM02"})
    assert len(r.text.split('<tr data-erp-code="')) - 1 == 2
    x = c["ex"].get("/col/collections", params={"date_from": "1991-06-01", "date_to": "1991-06-30", "customer": str(w.customer_id), "format": "xlsx"})
    assert x.status_code == 200 and x.headers["content-type"].startswith(XLSX)
    ws = load_workbook(io.BytesIO(x.content), read_only=True).worksheets[0]
    data = list(ws.iter_rows(values_only=True))
    assert data[0][0] == "일자" and len(data) == 4 and data[1][1] == w.code
    assert c["ac"].get("/col/collections", params={"date_from": "bad"}).status_code == 422
    assert c["ac"].get("/col/collections", params={"date_from": "1991-06-30", "date_to": "1991-06-01"}).status_code == 422
    assert c["ac"].get("/col/collections", params={"format": "csv"}).status_code == 422
    assert c["sh"].get("/col/collections").status_code == 403 and c["sl"].get("/col/collections").status_code == 200


@pytest.mark.fn("F-COL-041")
def test_collections_xlsx_has_no_formula_cells(c, w):
    """DEF-QA3-004 — 업로드 원문이 `=…` · `+…` · `@…` 라도 내려받은 엑셀에서 수식 셀(data_type f)이 되지 않는다."""
    w.map_code(c["ac"])
    evil = "=CMD|' /C calc'!A0"
    r = w.upload(c["ac"], *P2, [row(w.code, "1991-06-03", "100", "0", None, evil), row(w.code, "1991-06-04", "0", "5", None, "=HYPERLINK(\"http://x\",\"a\")"),
                                 row(w.code, "1991-06-05", "0", "5", None, "+1+1"), row(w.code, "1991-06-06", "0", "5", None, "@SUM(1)")])
    assert r.status_code == 200 and r.json()["errors"] == 0
    x = c["ac"].get("/col/collections", params={"date_from": "1991-06-01", "date_to": "1991-06-30", "customer": str(w.customer_id), "format": "xlsx"})
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).active
    cells = [cl for rw in ws.iter_rows() for cl in rw]
    assert [cl.coordinate for cl in cells if cl.data_type == "f"] == []
    methods = [rw[7].value for rw in ws.iter_rows(min_row=2)]                  # 수단 열 — 글자는 그대로, 수식은 아니다
    assert methods == [evil, '=HYPERLINK("http://x","a")', "+1+1", "@SUM(1)"]
    assert all(cl.data_type == "s" for cl in cells if isinstance(cl.value, str))
    assert ws["A1"].value == "일자" and ws["A1"].data_type == "s"


@pytest.mark.fn("F-COL-051")
def test_receivables_and_undecided_when_sales_unmapped(c, w):
    w.map_code(c["ac"])
    w.upload(c["ac"], *P2, [row(w.code, "1991-06-03", "100", "0"), row(w.code, "1991-06-10", "0", "70", "30"),
                             row(w.code, "1991-07-05", "50", "0", "80")])
    data = stats.receivables(date(1991, 6, 30))
    mine = [r for r in data["rows"] if r["customer_id"] == w.customer_id]
    assert len(mine) == 1 and (mine[0]["sales_amount"], mine[0]["collected_amount"], mine[0]["receivable_amount"], mine[0]["erp_balance"]) == \
        (Decimal("100.00"), Decimal("70.00"), Decimal("30.00"), Decimal("30.00"))
    page = c["ac"].get("/col/receivables", params={"asof": "1991-06-30"})
    assert page.status_code == 200 and f'data-customer="{w.customer_id}"' in page.text
    seg = page.text[page.text.index(f'data-customer="{w.customer_id}"'):][:900]
    assert "30.00" not in seg and "30" in seg                                 # 소수 0자리 표시
    assert stats.receivables(date(1991, 7, 31))["rows"][0]["receivable_amount"] is not None
    # 미수 0 숨김
    w.upload(c["ac"], date(1991, 8, 1), date(1991, 8, 31), [row(w.code, "1991-08-01", "0", "80")])
    zero = stats.receivables(date(1991, 8, 31))
    assert [r for r in zero["rows"] if r["customer_id"] == w.customer_id] == []
    assert [r for r in stats.receivables(date(1991, 8, 31), include_zero=True)["rows"] if r["customer_id"] == w.customer_id][0]["receivable_amount"] == 0
    assert c["ac"].get("/col/receivables", params={"asof": "1991-08-31", "show_zero": "1"}).status_code == 200
    # 매출액 열이 매핑되지 않으면 미확정 (D-516) — 수금 합은 그대로
    nosales = {k: v for k, v in H.items() if k != "sales_amount"}
    assert c["ac"].post("/col/uploads", data={"action": "map", **{f"map_{k}": v for k, v in nosales.items()}}).status_code == 200
    try:
        und = stats.receivables(date(1991, 6, 30))
        assert und["undecided"] == "D-516"
        m = [r for r in und["rows"] if r["customer_id"] == w.customer_id][0]
        assert (m["sales_amount"], m["receivable_amount"], m["collected_amount"]) == (None, None, Decimal("70.00"))
        page = c["ac"].get("/col/receivables", params={"asof": "1991-06-30"})
        assert page.status_code == 200 and "미확정 (D-516)" in page.text
        cr = stats.customer_receivable(w.customer_id)
        assert cr["status"] == "미확정 (D-516)" and cr["decision"] == "D-516" and cr["collected_amount"] == Decimal("150.00")
        assert "미확정 (D-516)" in c["ac"].get("/col/aging", params={"asof": "1991-06-30"}).text
    finally:
        assert c["ac"].post("/col/uploads", data={"action": "map", **{f"map_{k}": v for k, v in H.items()}}).status_code == 200
    assert c["ac"].get("/col/receivables", params={"asof": "x"}).status_code == 422
    assert c["sh"].get("/col/receivables").status_code == 403


@pytest.mark.fn("F-COL-061")
def test_aging_buckets_by_sales_date(c, w):
    w.map_code(c["ac"])
    asof = date(1991, 6, 30)
    days = lambda n: (asof - timedelta(days=n)).isoformat()  # noqa: E731
    w.upload(c["ac"], date(1991, 1, 1), date(1991, 6, 30),
             [row(w.code, days(10), "100", "0"), row(w.code, days(45), "200", "50"), row(w.code, days(75), "300", "0"),
              row(w.code, days(120), "400", "100"), row(w.code, days(0), "0", "20"), row(w.code, "1991-07-02", "999", "0")])
    data = stats.aging(asof)
    me = [r for r in data["rows"] if r["customer_id"] == w.customer_id][0]
    nets = {b: me["cells"][b]["net"] for b in data["buckets"]}
    assert nets == {"0~30일": Decimal("80"), "31~60일": Decimal("150"), "61~90일": Decimal("300"), "90일 초과": Decimal("300")}
    assert me["total_net"] == Decimal("830") and w.customer_id in [r["customer_id"] for r in data["over90"]]
    page = c["ac"].get("/col/aging", params={"asof": asof.isoformat()})
    assert page.status_code == 200 and f'data-customer="{w.customer_id}" data-over90="1"' in page.text and "D-803" in page.text
    assert c["ac"].get("/col/aging", params={"asof": "1991-13-01"}).status_code == 422
    assert c["sh"].get("/col/aging").status_code == 403 and c["ex"].get("/col/aging").status_code == 200


@pytest.mark.fn("F-COL-071")
def test_shipment_vs_collection_independent_aggregates(c, w):
    w.map_code(c["ac"])
    lot = w.lot("100")
    with conn.tx() as cur:
        for no, confirmed, q in (("A", "1992-03-05 10:00+09", 10), ("B", "1992-03-20 10:00+09", 5), ("C", "1992-04-02 10:00+09", 7)):
            cur.execute("""insert into shipment_order (shipment_no, request_id, customer_id, due_date, status, confirmed_at, created_by)
                           values (%s, %s, %s, '1992-03-31', '확정', %s::timestamptz, 'test') returning id""",
                        (f"{w.tag}-S{no}", w.request_id, w.customer_id, confirmed))
            oid = cur.fetchone()["id"]
            cur.execute("insert into shipment_line (shipment_order_id, product_lot_id, location_id, qty) values (%s, %s, %s, %s)",
                        (oid, lot["id"], w.location_id, q))
        cur.execute("""insert into shipment_order (shipment_no, customer_id, due_date, status, created_by)
                       values (%s, %s, '1992-03-31', '지시', 'test')""", (f"{w.tag}-SX", w.customer_id))   # 미확정 — 세지 않는다
    w.upload(c["ac"], date(1992, 3, 1), date(1992, 4, 30),
             [row(w.code, "1992-03-03", "1000", "0"), row(w.code, "1992-03-28", "0", "400"), row(w.code, "1992-04-10", "0", "600"),
              row(w.unmapped, "1992-03-15", "77", "77")])
    data = stats.shipment_vs_collection("1992-03", "1992-04", w.customer_id)
    assert data["months"] == ["1992-03", "1992-04"]
    by = {r["ym"]: r for r in data["rows"]}
    assert (by["1992-03"]["shipments"], by["1992-03"]["shipped_qty"], by["1992-03"]["sales_amount"], by["1992-03"]["collected_amount"]) == \
        (2, Decimal("15"), Decimal("1000.00"), Decimal("400.00"))
    assert (by["1992-04"]["shipments"], by["1992-04"]["shipped_qty"], by["1992-04"]["sales_amount"], by["1992-04"]["collected_amount"]) == \
        (1, Decimal("7"), Decimal("0.00"), Decimal("600.00"))
    # 독립 SQL 로 다시 센다 (G-14)
    ship = conn.q1("""select count(distinct o.id) as n, sum(l.qty) as q from shipment_order o join shipment_line l on l.shipment_order_id = o.id
                       where o.customer_id = %s and o.status = '확정' and o.confirmed_at >= '1992-03-01' and o.confirmed_at < '1992-04-01'""",
                   (w.customer_id,))
    erp = conn.q1("""select sum(sales_amount) as s, sum(collected_amount) as c from v_dq_collections
                      where erp_customer_code = %s and txn_date >= '1992-03-01' and txn_date < '1992-04-01'""", (w.code,))
    assert (ship["n"], ship["q"], erp["s"], erp["c"]) == (2, Decimal("15"), Decimal("1000.00"), Decimal("400.00"))
    page = c["sh"].get("/col/shipment-vs-collection", params={"ym_from": "1992-03", "ym_to": "1992-04", "customer": str(w.customer_id)})
    assert page.status_code == 200                                              # 출하실 R
    seg = page.text[page.text.index(f'data-month="1992-03" data-customer="{w.customer_id}"'):][:700]
    assert "15.000" in seg and "1,000" in seg and "400" in seg
    assert c["ac"].get("/col/shipment-vs-collection", params={"ym_from": "1992-3"}).status_code == 422
    assert c["ac"].get("/col/shipment-vs-collection", params={"ym_from": "1992-04", "ym_to": "1992-03"}).status_code == 422
    assert c["ac"].get("/col/shipment-vs-collection", params={"ym_from": "1990-01", "ym_to": "1993-12"}).status_code == 422
    assert c["ac"].get("/col/shipment-vs-collection").status_code == 200
    assert c["sl"].get("/col/shipment-vs-collection").status_code == 200
    assert login("pr").get("/col/shipment-vs-collection").status_code == 403


def test_stats_due_compliance(w):
    """TRK-020 납기 준수율 (공표 · D-804) — 납기 안 완료 · 지연 완료 · 지연 · 진행."""
    lot = w.lot("100")
    due = date(1993, 5, 15)
    r_ok, _ = w.request(w.item_id, Decimal("10"), due)
    r_late, _ = w.request(w.item_id, Decimal("10"), due)
    r_open, _ = w.request(w.item_id, Decimal("10"), due)
    with conn.tx() as cur:
        for no, rid, confirmed, q in (("A", r_ok, "1993-05-10", 10), ("B", r_late, "1993-05-20", 10), ("C", r_open, "1993-05-10", 4)):
            cur.execute("""insert into shipment_order (shipment_no, request_id, customer_id, due_date, status, confirmed_at, created_by)
                           values (%s, %s, %s, %s, '확정', %s::date + time '09:00', 'test') returning id""",
                        (f"{w.tag}-D{no}", rid, w.customer_id, due, confirmed))
            cur.execute("insert into shipment_line (shipment_order_id, product_lot_id, location_id, qty) values (%s, %s, %s, %s)",
                        (cur.fetchone()["id"], lot["id"], w.location_id, q))
    got = stats.due_compliance("1993-05")
    assert got["period"] == (date(1993, 5, 1), date(1993, 5, 31))
    mine = {r["request_id"]: r["bucket"] for r in got["rows"] if r["request_id"] in (r_ok, r_late, r_open)}
    assert mine == {r_ok: "납기 안 완료", r_late: "지연 완료", r_open: "지연"}
    assert (got["on_time"], got["late_done"], got["late_open"], got["in_progress"]) == (1, 1, 1, 0)
    assert got["rate"] == Decimal("33.3") and "D-804" in got["formula"]
    today = date.today()
    prog = stats.due_compliance((today + timedelta(days=10), today + timedelta(days=10)))   # World 기본 의뢰 — 출하 없음
    assert {r["bucket"] for r in prog["rows"] if r["request_id"] == w.request_id} == {"진행"}
    assert stats.due_compliance("1980-01")["rate"] is None                       # 분모 0 → 미수집
    with pytest.raises(ValueError):
        stats.due_compliance("1993/05")
    with pytest.raises(ValueError):
        stats.due_compliance((date(1993, 6, 1), date(1993, 5, 1)))
